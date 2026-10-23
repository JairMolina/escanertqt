"""Módulo Caiquén (v1.4.37). Clon del módulo SpeedLocker (app/services/speedlocker.py) con tablas propias
`caiquen_pcb` y `caiquen_eventos` (se crean solas, migración idempotente). Misma forma de datos y operaciones.
No escribe en `escaneos`, ni en el inventario TQT, ni usa log_evento. El QR se interpreta con `models.parse_caiquen`
(QR PROVISIONAL 'CQ-0001' → 'CQ-V1-0001'; el mismo parse hace que Escáner TQT, TPMS, AntiJammer y SpeedLocker las rechacen).
"""
import re
import sqlite3
from typing import Any, Dict, Iterable, List, Optional

from app.database import db
from app.database.inventario import normalizar_serie
from app.database.models import VERSION_CAIQUEN, parse_antijammer, parse_caiquen, parse_sigfox, parse_speedlocker, parse_tarjeta_code, parse_tpms

ESTADOS = ("RECIBIDA", "PROGRAMADA", "FALLA", "BAJA")
ACCIONES_MASA = ("firmware", "version", "lote", "estado", "eliminar")
SIN_FIRMWARE = "__sin"
_AHORA = "datetime('now','localtime')"
_SOLO_SERIE = re.compile(r"^\d{1,4}$", re.A)
_VERSION = re.compile(r"^(\d{1,2})(?:\.(\d))?$", re.A)


# ============================================================================
# Errores de negocio (el router los traduce a 404 / 409 / 422)
# ============================================================================
class CaiquenError(Exception):
    pass


class NoEncontrada(CaiquenError):
    pass


class Duplicada(CaiquenError):
    pass


class Invalida(CaiquenError):
    pass


# ============================================================================
# Esquema
# ============================================================================
_COLUMNAS_NUEVAS = {   # ALTER TABLE ADD COLUMN solo admite valores por omisión constantes
    "firmware": "TEXT",
    "estado": "TEXT NOT NULL DEFAULT 'RECIBIDA'",
    "nota": "TEXT NOT NULL DEFAULT ''",
    "programada_en": "TEXT",
    "actualizado_en": "TEXT",
}


def _asegurar() -> None:
    """Crea o migra las tablas (idempotente; mismo esquema que tpms_pcb)."""
    with db.transaction() as c:   # BEGIN IMMEDIATE (fix v1.4.29 de TPMS): espera su turno, sin "database is locked"
        c.execute(
            """
            CREATE TABLE IF NOT EXISTS caiquen_pcb(
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              nombre TEXT NOT NULL UNIQUE,          -- 'CQ-V1-0001'
              version TEXT NOT NULL,
              serie TEXT NOT NULL,
              codigo_qr TEXT NOT NULL DEFAULT '',   -- texto leído tal cual
              lote TEXT NOT NULL DEFAULT '',
              operador TEXT,
              recibida_en TEXT NOT NULL DEFAULT (datetime('now','localtime'))
            )
            """
        )
        existentes = {r[1] for r in c.execute("PRAGMA table_info(caiquen_pcb)")}
        for col, tipo in _COLUMNAS_NUEVAS.items():
            if col not in existentes:
                c.execute(f"ALTER TABLE caiquen_pcb ADD COLUMN {col} {tipo}")
        c.execute("UPDATE caiquen_pcb SET actualizado_en = recibida_en WHERE actualizado_en IS NULL")
        c.execute(
            """
            CREATE TABLE IF NOT EXISTS caiquen_eventos(
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              pcb_id INTEGER,
              nombre TEXT NOT NULL DEFAULT '',
              accion TEXT NOT NULL,
              detalle TEXT NOT NULL DEFAULT '',
              operador TEXT,
              creado_en TEXT NOT NULL DEFAULT (datetime('now','localtime'))
            )
            """
        )
        for sql in ("CREATE INDEX IF NOT EXISTS ix_caiquen_pcb_lote ON caiquen_pcb(lote)",
                    "CREATE INDEX IF NOT EXISTS ix_caiquen_pcb_estado ON caiquen_pcb(estado)",
                    "CREATE INDEX IF NOT EXISTS ix_caiquen_pcb_version ON caiquen_pcb(version)",
                    "CREATE INDEX IF NOT EXISTS ix_caiquen_pcb_firmware ON caiquen_pcb(firmware)",
                    "CREATE INDEX IF NOT EXISTS ix_caiquen_eventos_pcb ON caiquen_eventos(pcb_id, id)",
                    "CREATE INDEX IF NOT EXISTS ix_caiquen_eventos_nombre ON caiquen_eventos(nombre)"):
            c.execute(sql)


# ============================================================================
# Utilidades
# ============================================================================
def normalizar_version(valor: Optional[str]) -> str:
    """'V1' / '1' / '1.0' -> '1'; '1.1' / 'v1_1' -> '1.1' (mismo criterio que parse_caiquen)."""
    s = str(valor or "").strip().upper().replace("_", ".")
    if s.startswith("V"):
        s = s[1:]
    m = _VERSION.fullmatch(s)
    if not m or int(m.group(1)) == 0:
        raise Invalida(f"Versión inválida: '{valor}'. Usa un número como 1 o 1.1.")
    menor = m.group(2)
    return f"{int(m.group(1))}.{menor}" if menor and menor != "0" else str(int(m.group(1)))


def _serie(valor: Optional[str]) -> str:
    try:
        return normalizar_serie(valor)
    except ValueError as e:
        raise Invalida(str(e))


def nombre_canonico(version: str, serie: str) -> str:
    return f"CQ-V{version}-{serie}"


def _texto(valor: Optional[str], largo: int) -> str:
    return (valor or "").strip()[:largo]


def _estado(valor: Optional[str]) -> str:
    e = str(valor or "").strip().upper()
    if e not in ESTADOS:
        raise Invalida(f"Estado inválido: '{valor}'. Usa {', '.join(ESTADOS)}.")
    return e


def _evento(c: sqlite3.Connection, pcb_id: Optional[int], nombre: str, accion: str, detalle: str,
            operador: Optional[str]) -> None:
    c.execute("INSERT INTO caiquen_eventos(pcb_id, nombre, accion, detalle, operador) VALUES (?,?,?,?,?)",
              (pcb_id, nombre, accion, detalle[:300], operador))


def _fila(c: sqlite3.Connection, pcb_id: int) -> Dict[str, Any]:
    r = c.execute("SELECT * FROM caiquen_pcb WHERE id = ?", (pcb_id,)).fetchone()
    if not r:
        raise NoEncontrada(f"Placa Caiquén con ID {pcb_id} no encontrada.")
    return dict(r)


def _nombre_ocupado(c: sqlite3.Connection, nombre: str, excluir: Iterable[int] = ()) -> bool:
    excl = list(excluir)
    sql = "SELECT 1 FROM caiquen_pcb WHERE nombre = ?"
    if excl:
        sql += f" AND id NOT IN ({','.join('?' * len(excl))})"
    return c.execute(sql, (nombre, *excl)).fetchone() is not None


def _aplicar_firmware(c: sqlite3.Connection, fila: Dict[str, Any], firmware: Optional[str]) -> None:
    """Con firmware y RECIBIDA => PROGRAMADA; sin firmware y PROGRAMADA => RECIBIDA."""
    estado = fila["estado"]
    if firmware:
        if estado == "RECIBIDA":
            c.execute(f"UPDATE caiquen_pcb SET firmware = ?, estado = 'PROGRAMADA', programada_en = {_AHORA}, "
                      f"actualizado_en = {_AHORA} WHERE id = ?", (firmware, fila["id"]))
            return
        c.execute(f"UPDATE caiquen_pcb SET firmware = ?, programada_en = COALESCE(programada_en, {_AHORA}), "
                  f"actualizado_en = {_AHORA} WHERE id = ?", (firmware, fila["id"]))
        return
    if estado == "PROGRAMADA":
        c.execute(f"UPDATE caiquen_pcb SET firmware = NULL, estado = 'RECIBIDA', programada_en = NULL, "
                  f"actualizado_en = {_AHORA} WHERE id = ?", (fila["id"],))
    else:
        c.execute(f"UPDATE caiquen_pcb SET firmware = NULL, actualizado_en = {_AHORA} WHERE id = ?", (fila["id"],))


# ============================================================================
# Operaciones
# ============================================================================
def escanear(codigo: str, lote: str = "", operador: Optional[str] = None) -> Dict[str, Any]:
    """Alta idempotente: AGREGADA | REPETIDA | INVALIDA (nunca lanza por negocio)."""
    p = parse_caiquen(codigo)
    if not p and parse_sigfox(codigo):
        return {"resultado": "INVALIDA", "mensaje": f"{parse_sigfox(codigo)['nombre']} es una placa Skudo Sigfox: regístrala en el módulo Skudo Sigfox.", "pcb": None}
    if not p and parse_speedlocker(codigo):
        return {"resultado": "INVALIDA", "mensaje": f"{parse_speedlocker(codigo)['nombre']} es una placa SpeedLocker: regístrala en el módulo SpeedLocker.", "pcb": None}
    if not p and parse_antijammer(codigo):
        return {"resultado": "INVALIDA", "mensaje": f"{parse_antijammer(codigo)['nombre']} es una placa Skudo AntiJammer: regístrala en el módulo Skudo AntiJammer.", "pcb": None}
    if not p and parse_tpms(codigo):
        return {"resultado": "INVALIDA", "mensaje": f"{parse_tpms(codigo)['nombre']} es una placa TPMS: regístrala en el módulo TPMS.", "pcb": None}
    if not p and parse_tarjeta_code(codigo):
        return {"resultado": "INVALIDA", "mensaje": f"{parse_tarjeta_code(codigo)['nombre']} es una placa de Escáner TQT.", "pcb": None}
    if not p:
        return {"resultado": "INVALIDA", "mensaje": f"'{(codigo or '')[:40]}' no es un QR de placa Caiquén (ej. CQ-0001).", "pcb": None}
    _asegurar()
    with db.transaction() as c:
        fila = c.execute("SELECT * FROM caiquen_pcb WHERE nombre = ?", (p["nombre"],)).fetchone()
        if fila:
            return {"resultado": "REPETIDA", "mensaje": f"{p['nombre']} ya estaba registrada.", "pcb": dict(fila)}
        cur = c.execute("INSERT INTO caiquen_pcb(nombre, version, serie, codigo_qr, lote, operador) VALUES (?,?,?,?,?,?)",
                        (p["nombre"], p["version"], p["serie"], (codigo or "").strip()[:120], _texto(lote, 60), operador))
        _evento(c, cur.lastrowid, p["nombre"], "RECIBIDA", "Alta por QR", operador)
        fila = _fila(c, cur.lastrowid)
    return {"resultado": "AGREGADA", "mensaje": f"{p['nombre']} registrada.", "pcb": fila}


def alta_manual(version: str, serie: str, lote: str = "", operador: Optional[str] = None) -> Dict[str, Any]:
    v, s = normalizar_version(version), _serie(serie)
    nombre = nombre_canonico(v, s)
    _asegurar()
    with db.transaction() as c:
        if _nombre_ocupado(c, nombre):
            raise Duplicada(f"{nombre} ya está registrada.")
        cur = c.execute("INSERT INTO caiquen_pcb(nombre, version, serie, codigo_qr, lote, operador) VALUES (?,?,?,'',?,?)",
                        (nombre, v, s, _texto(lote, 60), operador))
        _evento(c, cur.lastrowid, nombre, "RECIBIDA", "Alta manual", operador)
        return _fila(c, cur.lastrowid)


def listar(lote: Optional[str] = None, limit: int = 500, q: Optional[str] = None, estado: Optional[str] = None,
           version: Optional[str] = None, firmware: Optional[str] = None) -> Dict[str, Any]:
    _asegurar()
    where: List[str] = []
    args: List[Any] = []
    if lote is not None:
        where.append("lote = ?"); args.append(lote)
    if estado:
        where.append("estado = ?"); args.append(_estado(estado))
    if version:
        where.append("version = ?"); args.append(normalizar_version(version))
    if firmware == SIN_FIRMWARE:
        where.append("(firmware IS NULL OR firmware = '')")
    elif firmware:
        where.append("firmware = ?"); args.append(firmware)
    if q and q.strip():
        patron = f"%{q.strip()}%"
        where.append("(nombre LIKE ? OR serie LIKE ? OR IFNULL(firmware,'') LIKE ? OR nota LIKE ?)")
        args += [patron] * 4
    filtro = (" WHERE " + " AND ".join(where)) if where else ""
    with db.get_db() as c:
        items = [dict(r) for r in c.execute(f"SELECT * FROM caiquen_pcb{filtro} ORDER BY id DESC LIMIT ?", (*args, limit))]
        total = c.execute(f"SELECT COUNT(*) FROM caiquen_pcb{filtro}", args).fetchone()[0]
    return {"items": items, "total": total}


def _nombre_de_codigo(codigo: str) -> Optional[str]:
    raw = (codigo or "").strip()
    if _SOLO_SERIE.fullmatch(raw):
        return nombre_canonico(VERSION_CAIQUEN, raw.zfill(4)) if int(raw) else None
    p = parse_caiquen(raw)
    return p["nombre"] if p else None


def por_codigo(codigo: str) -> Dict[str, Any]:
    """QR, 'CQ-V1-0001' o solo la serie ('1' => V1, VERSION_CAIQUEN) -> {pcb, eventos}."""
    nombre = _nombre_de_codigo(codigo)
    if not nombre:
        raise NoEncontrada(f"'{(codigo or '')[:40]}' no corresponde a una placa Caiquén.")
    _asegurar()
    with db.get_db() as c:
        r = c.execute("SELECT * FROM caiquen_pcb WHERE nombre = ?", (nombre,)).fetchone()
        if not r:
            raise NoEncontrada(f"{nombre} no está registrada.")
        eventos = [dict(e) for e in c.execute(
            "SELECT id, accion, detalle, operador, creado_en FROM caiquen_eventos WHERE pcb_id = ? ORDER BY id DESC", (r["id"],))]
    return {"pcb": dict(r), "eventos": eventos}


def poner_firmware(pcb_id: int, firmware: Optional[str], operador: Optional[str] = None) -> Dict[str, Any]:
    fw = _texto(firmware, 60) or None
    _asegurar()
    with db.transaction() as c:
        fila = _fila(c, pcb_id)
        _aplicar_firmware(c, fila, fw)
        if fw:
            _evento(c, pcb_id, fila["nombre"], "PROGRAMADA", f"Firmware {fw}", operador)
        else:
            _evento(c, pcb_id, fila["nombre"], "EDITADA", "Firmware quitado", operador)
        return _fila(c, pcb_id)


def editar(pcb_id: int, cambios: Dict[str, Any], operador: Optional[str] = None) -> Dict[str, Any]:
    """Cambia version/serie/lote/estado/nota (las claves ausentes o None no se tocan)."""
    _asegurar()
    with db.transaction() as c:
        fila = _fila(c, pcb_id)
        nuevo: Dict[str, Any] = {}
        if cambios.get("version") is not None:
            nuevo["version"] = normalizar_version(cambios["version"])
        if cambios.get("serie") is not None:
            nuevo["serie"] = _serie(cambios["serie"])
        if cambios.get("lote") is not None:
            nuevo["lote"] = _texto(cambios["lote"], 60)
        if cambios.get("estado") is not None:
            nuevo["estado"] = _estado(cambios["estado"])
        if cambios.get("nota") is not None:
            nuevo["nota"] = _texto(cambios["nota"], 500)
        nuevo = {k: v for k, v in nuevo.items() if v != fila[k]}
        if not nuevo:
            return fila
        if "version" in nuevo or "serie" in nuevo:
            nombre = nombre_canonico(nuevo.get("version", fila["version"]), nuevo.get("serie", fila["serie"]))
            if _nombre_ocupado(c, nombre, [pcb_id]):
                raise Duplicada(f"{nombre} ya está registrada.")
            nuevo["nombre"] = nombre
        sets = ", ".join(f"{k} = ?" for k in nuevo)
        extra = ", programada_en = COALESCE(programada_en, " + _AHORA + ")" if nuevo.get("estado") == "PROGRAMADA" else ""
        try:
            c.execute(f"UPDATE caiquen_pcb SET {sets}{extra}, actualizado_en = {_AHORA} WHERE id = ?", (*nuevo.values(), pcb_id))
        except sqlite3.IntegrityError:
            raise Duplicada(f"{nuevo.get('nombre', fila['nombre'])} ya está registrada.")
        detalle = "; ".join(f"{k}: {fila[k]!s} -> {v}" for k, v in nuevo.items() if k != "nombre")
        _evento(c, pcb_id, nuevo.get("nombre", fila["nombre"]), "EDITADA", detalle, operador)
        return _fila(c, pcb_id)


def eliminar(pcb_id: int, operador: Optional[str] = None) -> Dict[str, Any]:
    _asegurar()
    with db.transaction() as c:
        fila = _fila(c, pcb_id)
        c.execute("DELETE FROM caiquen_pcb WHERE id = ?", (pcb_id,))
        _evento(c, pcb_id, fila["nombre"], "ELIMINADA", "Placa eliminada", operador)
    return {"ok": True}


def masa(ids: List[int], accion: str, valor: Optional[str] = None, operador: Optional[str] = None) -> Dict[str, Any]:
    """Aplica una acción a varias placas: todo o nada (si una falla, no cambia ninguna)."""
    if accion not in ACCIONES_MASA:
        raise Invalida(f"Acción inválida: '{accion}'. Usa {', '.join(ACCIONES_MASA)}.")
    ids = list(dict.fromkeys(int(i) for i in ids))
    if not ids:
        raise Invalida("Selecciona al menos una placa.")
    if accion == "version":
        valor = normalizar_version(valor)
    elif accion == "estado":
        valor = _estado(valor)
    elif accion == "firmware":
        valor = _texto(valor, 60) or None
    elif accion == "lote":
        valor = _texto(valor, 60)
    _asegurar()
    with db.transaction() as c:
        filas = [dict(r) for r in c.execute(f"SELECT * FROM caiquen_pcb WHERE id IN ({','.join('?' * len(ids))})", ids)]
        faltan = sorted(set(ids) - {f["id"] for f in filas})
        if faltan:
            raise NoEncontrada(f"Placas Caiquén no encontradas: {', '.join(map(str, faltan[:20]))}.")
        if accion == "version":
            nombres = [nombre_canonico(valor, f["serie"]) for f in filas]
            repetidos = {n for n in nombres if nombres.count(n) > 1}
            ocupados = [n for n in nombres if _nombre_ocupado(c, n, ids)]
            if repetidos or ocupados:
                raise Duplicada(f"Nombre duplicado: {sorted(repetidos | set(ocupados))[0]}.")
        try:
            for f in filas:
                if accion == "firmware":
                    _aplicar_firmware(c, f, valor)
                elif accion == "version":
                    c.execute(f"UPDATE caiquen_pcb SET version = ?, nombre = ?, actualizado_en = {_AHORA} WHERE id = ?",
                              (valor, nombre_canonico(valor, f["serie"]), f["id"]))
                elif accion == "lote":
                    c.execute(f"UPDATE caiquen_pcb SET lote = ?, actualizado_en = {_AHORA} WHERE id = ?", (valor, f["id"]))
                elif accion == "estado":
                    extra = f", programada_en = COALESCE(programada_en, {_AHORA})" if valor == "PROGRAMADA" else ""
                    c.execute(f"UPDATE caiquen_pcb SET estado = ?{extra}, actualizado_en = {_AHORA} WHERE id = ?", (valor, f["id"]))
                else:
                    c.execute("DELETE FROM caiquen_pcb WHERE id = ?", (f["id"],))
                nombre = nombre_canonico(valor, f["serie"]) if accion == "version" else f["nombre"]
                detalle = "eliminada" if accion == "eliminar" else f"{accion}: {'' if valor is None else valor}"
                _evento(c, f["id"], nombre, "MASA", detalle, operador)
        except sqlite3.IntegrityError as e:
            raise Duplicada(f"Nombre duplicado: {e}")
    return {"afectadas": len(filas)}


def resumen() -> Dict[str, Any]:
    _asegurar()
    with db.get_db() as c:
        total = c.execute("SELECT COUNT(*) FROM caiquen_pcb").fetchone()[0]
        por_estado = {e: 0 for e in ESTADOS}
        for r in c.execute("SELECT estado, COUNT(*) FROM caiquen_pcb GROUP BY estado"):
            por_estado[r[0]] = r[1]
        versiones = [{"version": r[0], "n": r[1]} for r in
                     c.execute("SELECT version, COUNT(*) FROM caiquen_pcb GROUP BY version ORDER BY version")]
        firmwares = [{"firmware": r[0], "n": r[1]} for r in c.execute(   # firmware None = placas sin firmware (filtro «__sin»)
            "SELECT NULLIF(firmware, ''), COUNT(*) FROM caiquen_pcb GROUP BY NULLIF(firmware, '') ORDER BY firmware IS NULL, firmware")]
        lotes = [{"lote": r[0], "n": r[1]} for r in
                 c.execute("SELECT lote, COUNT(*) FROM caiquen_pcb GROUP BY lote ORDER BY lote")]
        hoy = c.execute("SELECT COUNT(*) FROM caiquen_pcb WHERE date(recibida_en) = date('now','localtime')").fetchone()[0]
        actividad = _actividad(c, 14)
        ultimos = [dict(r) for r in c.execute(f"SELECT {_COLS_EVENTO} FROM caiquen_eventos ORDER BY id DESC LIMIT 8")]
    return {"total": total, "por_estado": por_estado, "versiones": versiones, "firmwares": firmwares,
            "lotes": lotes, "hoy": hoy, "actividad": actividad, "ultimos": ultimos}


# ============================================================================
# v1.4.28b: consola (actividad, eventos, lotes, exportar)
# ============================================================================
_COLS_EVENTO = "id, pcb_id, nombre, accion, detalle, operador, creado_en"
_FECHA = re.compile(r"^\d{4}-\d{2}-\d{2}$", re.A)


def _actividad(c: sqlite3.Connection, dias: int) -> List[Dict[str, Any]]:
    """Recibidas y programadas por día de los últimos `dias` días (incluye días en 0), del más antiguo al de hoy."""
    fechas = [r[0] for r in c.execute(
        "WITH RECURSIVE d(n) AS (SELECT 0 UNION ALL SELECT n + 1 FROM d WHERE n < ?) "
        "SELECT date('now','localtime', '-' || n || ' days') FROM d ORDER BY n DESC", (dias - 1,))]
    desde = fechas[0]
    rec = dict(c.execute("SELECT date(recibida_en), COUNT(*) FROM caiquen_pcb WHERE date(recibida_en) >= ? "
                         "GROUP BY date(recibida_en)", (desde,)).fetchall())
    prog = dict(c.execute("SELECT date(programada_en), COUNT(*) FROM caiquen_pcb WHERE programada_en IS NOT NULL "
                          "AND date(programada_en) >= ? GROUP BY date(programada_en)", (desde,)).fetchall())
    return [{"dia": f, "recibidas": rec.get(f, 0), "programadas": prog.get(f, 0)} for f in fechas]


def _fecha(valor: Optional[str], campo: str) -> Optional[str]:
    if not valor:
        return None
    if not _FECHA.fullmatch(valor.strip()):
        raise Invalida(f"Fecha inválida en «{campo}»: '{valor}'. Usa AAAA-MM-DD.")
    return valor.strip()


def eventos(accion: Optional[str] = None, q: Optional[str] = None, desde: Optional[str] = None,
            hasta: Optional[str] = None, limit: int = 200, offset: int = 0) -> Dict[str, Any]:
    _asegurar()
    where: List[str] = []
    args: List[Any] = []
    if accion:
        where.append("accion = ?"); args.append(accion.strip().upper())
    if q and q.strip():
        patron = f"%{q.strip()}%"
        where.append("(nombre LIKE ? OR detalle LIKE ? OR IFNULL(operador,'') LIKE ?)"); args += [patron] * 3
    d, h = _fecha(desde, "desde"), _fecha(hasta, "hasta")
    if d:
        where.append("date(creado_en) >= ?"); args.append(d)
    if h:
        where.append("date(creado_en) <= ?"); args.append(h)
    filtro = (" WHERE " + " AND ".join(where)) if where else ""
    with db.get_db() as c:
        items = [dict(r) for r in c.execute(
            f"SELECT {_COLS_EVENTO} FROM caiquen_eventos{filtro} ORDER BY id DESC LIMIT ? OFFSET ?", (*args, limit, offset))]
        total = c.execute(f"SELECT COUNT(*) FROM caiquen_eventos{filtro}", args).fetchone()[0]
    return {"items": items, "total": total}


def lotes() -> Dict[str, Any]:
    _asegurar()
    with db.get_db() as c:
        items = [dict(r) for r in c.execute(
            "SELECT lote, COUNT(*) AS placas, "
            "SUM(estado = 'RECIBIDA') AS recibidas, SUM(estado = 'PROGRAMADA') AS programadas, "
            "SUM(estado = 'FALLA') AS falla, SUM(estado = 'BAJA') AS baja, "
            "MIN(recibida_en) AS primera, MAX(recibida_en) AS ultima "
            "FROM caiquen_pcb GROUP BY lote ORDER BY MAX(recibida_en) DESC")]
    return {"items": items}


def renombrar_lote(de: str, a: str, operador: Optional[str] = None) -> Dict[str, Any]:
    origen, destino = _texto(de, 60), _texto(a, 60)
    if origen == destino:
        raise Invalida("El nombre nuevo del lote es igual al actual.")
    _asegurar()
    with db.transaction() as c:
        filas = [dict(r) for r in c.execute("SELECT id, nombre FROM caiquen_pcb WHERE lote = ?", (origen,))]
        c.execute(f"UPDATE caiquen_pcb SET lote = ?, actualizado_en = {_AHORA} WHERE lote = ?", (destino, origen))
        for f in filas:
            _evento(c, f["id"], f["nombre"], "EDITADA", f"lote: {origen} -> {destino}", operador)
    return {"afectadas": len(filas)}


COLUMNAS_XLSX = [("Nombre", 18), ("Hardware", 10), ("Serie", 8), ("Firmware", 16), ("Estado", 12), ("Lote", 20),
                 ("Nota", 30), ("Operador", 26), ("Recibida", 20), ("Programada", 20)]
COLUMNAS_MOV = [("Fecha", 20), ("Placa", 18), ("Acción", 12), ("Detalle", 40), ("Operador", 26)]


def exportar_xlsx(q: Optional[str] = None, estado: Optional[str] = None, version: Optional[str] = None,
                  firmware: Optional[str] = None, lote: Optional[str] = None) -> bytes:
    """Libro con hoja «Placas» (mismos filtros que la lista) y hoja «Movimientos» (últimos 2000 eventos)."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    placas = listar(lote, 1_000_000, q, estado, version, firmware)["items"]
    movs = eventos(limit=2000)["items"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Placas"
    hojas = [(ws, COLUMNAS_XLSX, [[p["nombre"], f"V{p['version']}", p["serie"], p["firmware"] or "", p["estado"], p["lote"],
                                  p["nota"], p["operador"] or "", p["recibida_en"], p["programada_en"] or ""] for p in placas]),
             (wb.create_sheet("Movimientos"), COLUMNAS_MOV,
              [[e["creado_en"], e["nombre"], e["accion"], e["detalle"], e["operador"] or ""] for e in movs])]
    for hoja, columnas, filas in hojas:
        hoja.append([t for t, _ in columnas])
        for i, (_, ancho) in enumerate(columnas, start=1):
            hoja.cell(row=1, column=i).font = Font(bold=True)
            hoja.column_dimensions[get_column_letter(i)].width = ancho
        for f in filas:
            hoja.append(f)
        hoja.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
