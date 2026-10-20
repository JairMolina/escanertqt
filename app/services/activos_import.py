"""v1.4.15: importación de activos desde el Excel de inventario (una hoja por departamento).

`analizar` lee el libro (sin escribir nada) y devuelve, por hoja, lo que se crearía, se actualizaría u omitiría;
`aplicar` escribe esos items en una sola transacción. Solo se leen las filas con relleno verde (`FFE2EFD9` en la columna A),
que son las ya revisadas, salvo que se pida `solo_verdes=False`.

El libro llega con texto dañado (mojibake) y con tres variantes de columnas, así que:
  * los encabezados se normalizan (sin acentos, mayúsculas ni signos) y se mapean por nombre;
  * si las columnas 12 y 13 dicen «Descripción del item» y «Colaborador» la hoja sigue el formato estándar de 26 columnas
    por posición (SKYCENTER trae mal los encabezados 1-10, pero sus datos son estándar).
"""
import io
import logging
import re
import unicodedata
from datetime import date, datetime, timedelta
from typing import Any, Dict, List, Optional, Set, Tuple

from app.database import db
from app.services import activos as act

log = logging.getLogger(__name__)

VERDE = "FFE2EFD9"
FORMATO_ESTANDAR = "estandar"
FORMATO_POR_NOMBRE = "por_nombre"

# Posiciones (1-based) del formato estándar de 26 columnas.
ESTANDAR = ["nombre", "status", "departamento", "numero", "fecha_revision", "valor", "asignado", "modificado_por", "ubicacion",
            "serie", "foto", "descripcion", "colaborador", "estado_fisico", "notas", "baja_fecha", "baja_razon", "baja_motivo",
            "baja_comentarios", "ultima_mod", "origen", "creacion", "procesador", "ram", "disco", "pantalla"]

# Encabezado normalizado -> campo (para las hojas de otro orden: VIPE, MKT, ALMACEN y las de encabezados limpios).
ALIAS = {
    "nombreactivo": "nombre", "nombredelactivo": "nombre", "status": "status", "estatus": "status", "departamento": "departamento",
    "activonumber": "numero", "numerodeactivo": "numero", "fechaderevision": "fecha_revision", "valordelactivo": "valor",
    "asignadoa": "asignado", "ultimamodificacionhechapor": "modificado_por", "ubicaciondelactivo": "ubicacion",
    "ubicacion": "ubicacion", "numerodeserie": "serie", "fotografiadelitem": "foto", "descripciondelitem": "descripcion",
    "descripciondelarticulo": "descripcion", "descripcion": "descripcion", "colaborador": "colaborador",
    "comentariosestadofisico": "estado_fisico", "notas": "notas", "fechadebaja": "baja_fecha", "razondebaja": "baja_razon",
    "motivodebaja": "baja_motivo", "comentariosdebaja": "baja_comentarios", "ultimamodificacion": "ultima_mod",
    "origen": "origen", "fechadecreacion": "creacion", "procesador": "procesador", "memoriaram": "ram", "discoduro": "disco",
}
VACIOS = {"", "s/n", "sn", "n/a", "na", "undefined", "none", "null", "-", "--", ".", "..", "#n/a", "#ref!"}

# (categoría, palabras) en orden de prioridad: «MONITOR PARA COMPUTADORA» es Monitores, «SILLA PARA LAPTOP» es Mobiliario.
CATEGORIAS = (
    ("Monitores", ("MONITOR",)),
    ("Telefonía", ("CELULAR", "TELEFONO", "DIADEMA")),
    ("Mobiliario", ("SILLA", "ESCRITORIO", "MESA", "ARCHIVERO", "MUEBLE", "CAJONERA", "PIZARRON")),
    ("Seguridad", ("CAMARA",)),
    ("Periféricos", ("TECLADO", "MOUSE")),
    ("Equipo de oficina", ("AIRE", "FRIGOBAR", "CAFETERA")),
    ("Cómputo", ("LAPTOP", "COMPUTADORA", "CPU")),
)
CATEGORIA_DEFECTO = "General"

_RE_FFFD = re.compile("�(?:_x00([0-9A-Fa-f]{2})_|([\x80-\xbf]))")
_RE_PREFIJO = re.compile(r"^[A-Za-z_]+::::")
_RE_ACT = re.compile(r"^\s*ACT[\s\-_]*(\d{1,6})\s*$", re.IGNORECASE)
_RE_FECHA = re.compile(r"(\d{1,2})[-/](\d{1,2})[-/](\d{4})(?:\s+(\d{1,2}):(\d{2})(?::\d{2})?\s*([ap])?\.?\s*m?\.?)?", re.IGNORECASE)
_RE_ISO = re.compile(r"(\d{4})-(\d{2})-(\d{2})(?:[T\s](\d{2}):(\d{2}))?")


# ------------------------------------------------------------------ limpieza de texto
def reparar(texto: Any) -> str:
    """Repara el mojibake del libro: «\\ufffd»+byte 0x80-0xBF (o «\\ufffd_x00NN_») = byte UTF-8 con prefijo C3, «Ã³» = «ó», «ﾑ» = «Ñ»."""
    s = "" if texto is None else str(texto)
    if not s:
        return s
    s = _RE_FFFD.sub(lambda m: bytes([0xC3, int(m.group(1), 16) if m.group(1) else ord(m.group(2))]).decode("utf-8", "replace"), s)
    s = re.sub(r"_x00([0-9A-Fa-f]{2})_", lambda m: chr(int(m.group(1), 16)), s)   # escape de Excel «_x008D_» (Ã_x008D_ = Í)
    if re.search("[ÂÃ]", s):   # UTF-8 leído como cp1252 / latin-1
        for cod in ("cp1252", "latin-1"):
            try:
                s = s.encode(cod).decode("utf-8")
                break
            except (UnicodeEncodeError, UnicodeDecodeError):
                continue
    return s.replace("ﾑ", "Ñ")


def normalizar(texto: Any) -> str:
    """Minúsculas, sin acentos y solo letras/dígitos (para comparar encabezados)."""
    s = unicodedata.normalize("NFKD", reparar(texto))
    return re.sub(r"[^a-z0-9]", "", "".join(ch for ch in s if not unicodedata.combining(ch)).lower())


def texto(v: Any) -> str:
    """Valor de celda como texto limpio: reparado, sin prefijos «Employees::::», y «S/N», «N/A», «undefined», «.» → vacío."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (datetime, date)):
        return v.date().isoformat() if isinstance(v, datetime) else v.isoformat()
    s = " ".join(reparar(v).split()) if not isinstance(v, str) or "\n" not in v else reparar(v).strip()
    s = _RE_PREFIJO.sub("", s).strip()
    if s.lower() in VACIOS or normalizar(s).startswith("soloelectronicos"):   # «Solo electronicos campos en adelante» = basura de plantilla
        return ""
    return s


def fecha_hora(v: Any) -> Optional[datetime]:
    """datetime, número de serie de Excel (44861.0) o texto «16-11-2018 06:27 p.m.» → datetime (None si no se entiende)."""
    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 1 < v < 2958466:
            return datetime(1899, 12, 30) + timedelta(days=float(v))
        return None
    s = texto(v)
    if not s:
        return None
    try:
        m = _RE_FECHA.search(s)
        if m:
            d, mes, a, h, mi, ap = m.groups()
            hh = int(h or 0)
            if ap:
                hh = hh % 12 + (12 if ap.lower() == "p" else 0)
            return datetime(int(a), int(mes), int(d), hh, int(mi or 0))
        m = _RE_ISO.search(s)
        if m:
            a, mes, d, h, mi = m.groups()
            return datetime(int(a), int(mes), int(d), int(h or 0), int(mi or 0))
        if re.fullmatch(r"\d+(\.\d+)?", s):
            return fecha_hora(float(s))
    except ValueError:
        return None
    return None


def valor_num(v: Any) -> Optional[float]:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        n = float(v)
    else:
        s = re.sub(r"[$\s,]", "", texto(v))
        try:
            n = float(s)
        except ValueError:
            return None
    return round(n, 2) if 0 <= n <= 1e12 else None


def categoria_de(nombre: str) -> str:
    up = "".join(ch for ch in unicodedata.normalize("NFKD", nombre or "") if not unicodedata.combining(ch)).upper()
    for cat, palabras in CATEGORIAS:
        if re.search(r"\b(?:" + "|".join(palabras) + ")", up):
            return cat
    return CATEGORIA_DEFECTO


def estado_de(status: str, tiene_baja: bool) -> str:
    n = normalizar(status)
    if n.startswith("baja"):
        return "baja"
    if n.startswith("resguardo"):
        return "resguardo"
    if n.startswith("reparacion"):
        return "reparacion"
    return "baja" if (not n and tiene_baja) else "activo"   # sin status pero con fecha de baja: ya salió


# ------------------------------------------------------------------ lectura de hojas
def _mapa_columnas(encabezados: List[Any]) -> Tuple[str, Dict[str, int]]:
    """(formato, campo -> índice 0-based) de una hoja a partir de su fila 1."""
    n = [normalizar(h) for h in encabezados]
    if len(n) >= 13 and n[11] == "descripciondelitem" and n[12] == "colaborador":
        return FORMATO_ESTANDAR, {campo: i for i, campo in enumerate(ESTANDAR)}
    mapa: Dict[str, int] = {}
    for i, h in enumerate(n):
        campo = "pantalla" if h.startswith("tamano") else ALIAS.get(h)
        if h == "comentarios":   # ALMACEN: la 1.ª es el estado físico, la 2.ª los comentarios de baja
            campo = "estado_fisico" if "estado_fisico" not in mapa else "baja_comentarios"
        if campo and campo not in mapa:
            mapa[campo] = i
    return FORMATO_POR_NOMBRE, mapa


def _es_verde(celda: Any) -> bool:
    f = celda.fill
    rgb = getattr(f.fgColor, "rgb", None) if f is not None else None
    return bool(f is not None and f.fill_type == "solid" and isinstance(rgb, str) and rgb.upper() == VERDE)


def _item(hoja: str, fila: int, num: int, g) -> Dict[str, Any]:
    """Item normalizado de una fila. `g(campo)` devuelve el valor crudo de la celda."""
    nombre = texto(g("nombre"))[:120]
    f_baja = fecha_hora(g("baja_fecha"))
    baja = [f_baja.date().isoformat() if f_baja else texto(g("baja_fecha")), texto(g("baja_razon")), texto(g("baja_motivo")),
            texto(g("baja_comentarios"))]
    colaborador, asignado = texto(g("colaborador")), texto(g("asignado"))
    ub = texto(g("ubicacion"))
    creacion, revision = fecha_hora(g("creacion")), fecha_hora(g("fecha_revision"))
    desc = texto(g("descripcion"))
    pares = [
        ("Asignado a", asignado), ("Colaborador", colaborador), ("Descripción", desc), ("Procesador", texto(g("procesador"))),
        ("Memoria RAM", texto(g("ram"))), ("Disco duro", texto(g("disco"))), ("Tamaño de pantalla", texto(g("pantalla"))),
        ("Fecha de revisión", revision.date().isoformat() if revision else ""),
        ("Comentarios de estado físico", texto(g("estado_fisico"))), ("Notas", texto(g("notas"))), ("Origen", texto(g("origen"))),
        ("Hoja de origen", hoja), ("Fecha de baja", baja[0]), ("Razón de baja", baja[1]), ("Motivo de baja", baja[2]),
        ("Comentarios de baja", baja[3]),
    ]
    ultima = fecha_hora(g("ultima_mod"))
    return {
        "numero": num, "codigo": act.codigo_de(num), "hoja": hoja, "fila": fila, "nombre": nombre,
        "estado": estado_de(texto(g("status")), bool(baja[0])), "departamento": texto(g("departamento"))[:80],
        "categoria": categoria_de(nombre), "serie": texto(g("serie"))[:80], "ubicacion": ub[:120],
        "responsable": (colaborador or asignado)[:120], "valor": valor_num(g("valor")),
        "fecha_adquisicion": creacion.date().isoformat() if creacion else None, "notas": desc[:2000],
        "caracteristicas": [{"etiqueta": e, "valor": v[:500]} for e, v in pares if v],
        "ultima_mod": ultima.isoformat(timespec="minutes") if ultima else None,
    }


def _leer_hoja(ws, solo_verdes: bool) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
    nombre = ws.title
    info: Dict[str, Any] = {"hoja": nombre, "formato": "", "filas_verdes": 0, "filas_con_numero": 0, "omitidos": []}
    filas = ws.iter_rows(min_row=1)
    try:
        encabezados = [c.value for c in next(filas)]
    except StopIteration:
        info["formato"] = "vacía"
        return info, []
    formato, mapa = _mapa_columnas(encabezados)
    info["formato"] = formato
    if "numero" not in mapa or "nombre" not in mapa:
        info["formato"] = "desconocido"
        info["omitidos"].append({"fila": 1, "motivo": "No se encontraron las columnas «Nombre Activo» y «Activo Number»."})
        return info, []
    items: List[Dict[str, Any]] = []
    for fila in filas:
        celdas = list(fila)
        verde = bool(celdas) and _es_verde(celdas[0])
        info["filas_verdes"] += verde
        if solo_verdes and not verde:
            continue
        g = lambda campo, _c=celdas: (_c[mapa[campo]].value if campo in mapa and mapa[campo] < len(_c) else None)
        bruto_num, bruto_nombre = texto(g("numero")), texto(g("nombre"))
        if not bruto_num and not bruto_nombre:
            continue   # fila vacía (la hoja trae miles con formato)
        info["filas_con_numero"] += 1
        m = _RE_ACT.match(bruto_num)
        n = int(m.group(1)) if m else 0
        if not 1 <= n <= act.NUMERO_MAX:
            info["omitidos"].append({"fila": fila[0].row, "motivo": f"Número de activo inválido o vacío («{bruto_num[:20]}»)."})
            continue
        if not bruto_nombre:
            info["omitidos"].append({"fila": fila[0].row, "motivo": f"{act.codigo_de(n)} no tiene nombre."})
            continue
        items.append(_item(nombre, fila[0].row, n, g))
    return info, items


def _unificar_departamentos(items: List[Dict[str, Any]]) -> None:
    """Una sola forma por departamento: «Administracion»/«OPERACIONES» pasan a «Administración»/«Operaciones».
    Gana la escritura con acentos y sin todo en mayúsculas; a igualdad, la más frecuente."""
    formas: Dict[str, Dict[str, int]] = {}
    for it in items:
        if it["departamento"]:
            f = formas.setdefault(normalizar(it["departamento"]), {})
            f[it["departamento"]] = f.get(it["departamento"], 0) + 1
    def puntaje(forma: str, n: int):
        return (forma != forma.upper(), sum(ord(ch) > 127 for ch in forma), n)
    canon = {k: max(f.items(), key=lambda x: puntaje(*x))[0] for k, f in formas.items()}
    for it in items:
        if it["departamento"]:
            it["departamento"] = canon[normalizar(it["departamento"])]


def analizar(datos: bytes, solo_verdes: bool = True, existentes: Optional[Set[int]] = None) -> Dict[str, Any]:
    """Analiza el .xlsx sin escribir en la BD. ValueError si el archivo no es un libro legible.

    Devuelve {hojas:[{hoja, formato, filas_verdes, nuevos, actualizar, omitidos}], items, avisos, totales}. Si un número
    aparece en varias hojas gana la fila con «Ultima Modificación» más reciente (y se avisa)."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(datos), data_only=True)
    except Exception as e:   # zip dañado, no es xlsx, etc.
        log.warning("Importación de activos: libro ilegible (%s)", e)
        raise ValueError("No se pudo leer el archivo: no es un libro .xlsx válido.")
    if existentes is None:
        act._asegurar()
        with db.get_db() as c:
            existentes = {r[0] for r in c.execute("SELECT numero FROM activos")}
    hojas: List[Dict[str, Any]] = []
    candidatos: Dict[int, Dict[str, Any]] = {}
    avisos: List[Dict[str, Any]] = []
    por_hoja: Dict[str, Dict[str, Any]] = {}
    for ws in wb.worksheets:
        info, items = _leer_hoja(ws, solo_verdes)
        hojas.append(info)
        por_hoja[info["hoja"]] = info
        for it in items:
            previo = candidatos.get(it["numero"])
            if previo is None:
                candidatos[it["numero"]] = it
                continue
            gana, pierde = (it, previo) if (it["ultima_mod"] or "") > (previo["ultima_mod"] or "") else (previo, it)
            candidatos[it["numero"]] = gana
            por_hoja[pierde["hoja"]]["omitidos"].append({
                "fila": pierde["fila"],
                "motivo": f"{pierde['codigo']} repetido: se usa la fila {gana['fila']} de «{gana['hoja']}» (modificación más reciente)."})
            avisos.append({"codigo": gana["codigo"], "gana": f"{gana['hoja']}!{gana['fila']}", "descartada": f"{pierde['hoja']}!{pierde['fila']}"})
    items = sorted(candidatos.values(), key=lambda i: i["numero"])
    _unificar_departamentos(items)
    for it in items:
        it["accion"] = "actualizar" if it["numero"] in existentes else "nuevo"
    for h in hojas:
        mios = [i for i in items if i["hoja"] == h["hoja"]]
        h["nuevos"] = sum(i["accion"] == "nuevo" for i in mios)
        h["actualizar"] = len(mios) - h["nuevos"]
    totales = {"nuevos": sum(h["nuevos"] for h in hojas), "actualizar": sum(h["actualizar"] for h in hojas),
               "omitidos": sum(len(h["omitidos"]) for h in hojas)}
    return {"hojas": hojas, "items": items, "avisos": avisos, "totales": totales}


# ------------------------------------------------------------------ escritura
def aplicar(items: List[Dict[str, Any]], por: str) -> Dict[str, Any]:
    """Crea los items nuevos y actualiza los existentes (por número) en UNA transacción; cada cambio deja el evento
    IMPORTACION con la hoja y la fila. En una actualización los campos vacíos del Excel no borran lo ya registrado y las
    características se mezclan con las existentes."""
    act._asegurar()
    token = datetime.now().strftime("%Y%m%d%H%M%S")
    creados = actualizados = sin_cambios = 0
    codigos: List[str] = []
    errores: List[Dict[str, Any]] = []
    with db.transaction() as c:
        for it in items:
            origen = f"Importación {token} · hoja {it['hoja']}, fila {it['fila']}"
            try:
                v = act._limpiar({k: it.get(k) for k in act.EDITABLES if k != "numero"}, parcial=False)
                r = c.execute("SELECT * FROM activos WHERE numero=?", (it["numero"],)).fetchone()
                if r is None:
                    act._insertar(c, it["numero"], v, por, "IMPORTACION", f"{origen}: alta de {it['codigo']}")
                    creados += 1
                else:
                    nuevas = {k: x for k, x in v.items() if k == "estado" or x not in ("", None)}
                    if "caracteristicas" in nuevas:
                        mezcla = dict(act._caracteristicas_dict(r["caracteristicas"]))
                        mezcla.update(act._caracteristicas_dict(nuevas["caracteristicas"]))
                        nuevas["caracteristicas"] = act._caracteristicas_json(mezcla)
                    if act._modificar(c, r, nuevas, por, "IMPORTACION", origen + ": "):
                        actualizados += 1
                    else:
                        sin_cambios += 1
                codigos.append(it["codigo"])
            except ValueError as e:
                errores.append({"codigo": it.get("codigo"), "detail": str(e)})
    if creados or actualizados:
        db.log_evento("ACTIVO_IMPORTACION", valor=token, detalle=f"{creados} nuevos, {actualizados} actualizados", operador=por)
    return {"creados": creados, "actualizados": actualizados, "sin_cambios": sin_cambios, "codigos": codigos,
            "errores": errores, "lote": token}


def codigos_ultima_importacion() -> List[str]:
    """Códigos tocados por la última importación (los eventos IMPORTACION comparten el mismo token de lote)."""
    act._asegurar()
    with db.get_db() as c:
        r = c.execute("SELECT detalle FROM activos_eventos WHERE evento='IMPORTACION' ORDER BY id DESC LIMIT 1").fetchone()
        m = re.search(r"Importación (\d{14})", r["detalle"]) if r else None
        if not m:
            return []
        filas = c.execute("SELECT DISTINCT a.numero FROM activos_eventos e JOIN activos a ON a.id=e.activo_id "
                          "WHERE e.evento='IMPORTACION' AND e.detalle LIKE ? ORDER BY a.numero", (f"Importación {m.group(1)}%",))
        return [act.codigo_de(x[0]) for x in filas]
