"""Costeo por módulo (v1.4.34): el «Cotizador» de cada proyecto (lista de materiales + costo por tarjeta) dentro de su
módulo, con el mismo criterio que Costos de Inventario TQTR. Un solo motor para todos los módulos: tablas
`costeo_componentes` y `costeo_meta` con columna `modulo`; cada módulo tiene SUS componentes y SU tipo de cambio.

- Cada QR es una sola tarjeta: el costo se calcula por tarjeta (AntiJammer: RX y TX; Sigfox: Sigfox).
- `tipo='recurrente'` (p. ej. la licencia Sigfox anual) se muestra aparte y no suma al costo de producción.
- `importar_xlsx` lee el formato de los cotizadores (hoja «Componentes»: tipo de tarjeta, categoría, descripción,
  cantidad por lote, unidad, cantidad por ensamble, …, moneda, costo unitario). Una fila «Tarjeta RX y TX» trae la
  cantidad de las dos tarjetas: se reparte a partes iguales (igual que el Dashboard del Excel).
"""
import io
import re
from typing import Any, Dict, List, Optional

from app.database import db

# slug de la URL → id del módulo (app/modulos.py) y sus tarjetas
MODULOS_COSTEO: Dict[str, Dict[str, Any]] = {
    "antijammer": {"modulo": "skudo-antijammer", "nombre": "Skudo AntiJammer", "tarjetas": ("RX", "TX")},
    "sigfox": {"modulo": "skudo-sigfox", "nombre": "Skudo Sigfox", "tarjetas": ("Sigfox",)},
}
MONEDAS = ("MXN", "USD")
TIPOS = ("produccion", "recurrente")
TIPO_CAMBIO = 18.5
_CAMPOS = ("tarjeta", "categoria", "descripcion", "proveedor", "cantidad_lote", "unidad", "cantidad", "moneda",
           "costo_unitario", "tipo", "notas", "orden")
_listo = False


class CosteoError(Exception):
    pass


class NoEncontrado(CosteoError):
    pass


def _conf(slug: str) -> Dict[str, Any]:
    if slug not in MODULOS_COSTEO:
        raise NoEncontrado(f"El módulo «{slug}» no tiene costeo.")
    return MODULOS_COSTEO[slug]


def _asegurar() -> None:
    global _listo
    if _listo:
        return
    with db.transaction() as c:
        c.executescript("""
        CREATE TABLE IF NOT EXISTS costeo_componentes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            modulo TEXT NOT NULL,
            tarjeta TEXT NOT NULL,
            categoria TEXT NOT NULL DEFAULT '',
            descripcion TEXT NOT NULL,
            proveedor TEXT NOT NULL DEFAULT '',
            cantidad_lote REAL NOT NULL DEFAULT 0,
            unidad TEXT NOT NULL DEFAULT '',
            cantidad REAL NOT NULL DEFAULT 0,
            moneda TEXT NOT NULL DEFAULT 'MXN',
            costo_unitario REAL NOT NULL DEFAULT 0,
            tipo TEXT NOT NULL DEFAULT 'produccion',
            notas TEXT NOT NULL DEFAULT '',
            orden INTEGER NOT NULL DEFAULT 0,
            actualizado_en TEXT NOT NULL DEFAULT (datetime('now','localtime')),
            UNIQUE (modulo, tarjeta, descripcion)
        );
        CREATE TABLE IF NOT EXISTS costeo_meta (modulo TEXT NOT NULL, clave TEXT NOT NULL, valor TEXT, PRIMARY KEY (modulo, clave));
        """)
    _listo = True


def _tc(c, modulo: str) -> float:
    r = c.execute("SELECT valor FROM costeo_meta WHERE modulo=? AND clave='tipo_cambio'", (modulo,)).fetchone()
    try:
        return float(r["valor"]) if r else TIPO_CAMBIO
    except (TypeError, ValueError):
        return TIPO_CAMBIO


def _num(v: Any, campo: str) -> float:
    try:
        x = float(v if v not in (None, "") else 0)
    except (TypeError, ValueError):
        raise CosteoError(f"«{campo}» debe ser un número.")
    if x < 0 or x != x:
        raise CosteoError(f"«{campo}» no puede ser negativo.")
    return x


def _validar(conf: Dict[str, Any], d: Dict[str, Any], parcial: bool) -> Dict[str, Any]:
    out: Dict[str, Any] = {}
    for k in _CAMPOS:
        if k not in d or d[k] is None:
            continue
        v = d[k]
        if k in ("cantidad_lote", "cantidad", "costo_unitario"):
            out[k] = _num(v, k)
        elif k == "orden":
            out[k] = int(v or 0)
        else:
            out[k] = str(v).strip()[:300]
    if not parcial:
        out.setdefault("tarjeta", conf["tarjetas"][0])
        if not out.get("descripcion"):
            raise CosteoError("Falta la descripción del material.")
    if "descripcion" in out and not out["descripcion"]:
        raise CosteoError("Falta la descripción del material.")
    if "tarjeta" in out and out["tarjeta"] not in conf["tarjetas"]:
        raise CosteoError(f"Tarjeta inválida: usa {', '.join(conf['tarjetas'])}.")
    if "moneda" in out and out["moneda"].upper() not in MONEDAS:
        raise CosteoError("Moneda inválida: MXN o USD.")
    if "moneda" in out:
        out["moneda"] = out["moneda"].upper()
    if "tipo" in out and out["tipo"] not in TIPOS:
        raise CosteoError("Tipo inválido: produccion o recurrente.")
    return out


def _resumen(c, conf: Dict[str, Any]) -> Dict[str, Any]:
    tc = _tc(c, conf["modulo"])
    items: List[Dict[str, Any]] = []
    por_tarjeta = {t: 0.0 for t in conf["tarjetas"]}
    recurrente = 0.0
    for r in c.execute("SELECT * FROM costeo_componentes WHERE modulo=? ORDER BY tarjeta, orden, id", (conf["modulo"],)):
        d = dict(r)
        d["subtotal"] = round(d["costo_unitario"] * d["cantidad"], 4)
        d["subtotal_mxn"] = round(d["subtotal"] * (tc if d["moneda"] == "USD" else 1.0), 4)
        d["rendimiento"] = round(d["cantidad_lote"] / d["cantidad"], 2) if d["cantidad"] else None
        if d["tipo"] == "recurrente":
            recurrente += d["subtotal_mxn"]
        elif d["tarjeta"] in por_tarjeta:
            por_tarjeta[d["tarjeta"]] += d["subtotal_mxn"]
        items.append(d)
    return {
        "modulo": conf["modulo"], "nombre": conf["nombre"], "tarjetas": list(conf["tarjetas"]), "monedas": list(MONEDAS),
        "tipo_cambio": tc, "items": items,
        "por_tarjeta": {k: round(v, 2) for k, v in por_tarjeta.items()},
        "total_produccion": round(sum(por_tarjeta.values()), 2),
        "recurrente": round(recurrente, 2),
    }


def resumen(slug: str) -> Dict[str, Any]:
    conf = _conf(slug); _asegurar()
    with db.get_db() as c:
        return _resumen(c, conf)


def _uno(c, modulo: str, cid: int) -> Dict[str, Any]:
    r = c.execute("SELECT * FROM costeo_componentes WHERE id=? AND modulo=?", (cid, modulo)).fetchone()
    if not r:
        raise NoEncontrado("Ese componente no existe en este módulo.")
    return dict(r)


def _duplicado(e: Exception) -> CosteoError:
    return CosteoError("Ya existe ese material en esa tarjeta.") if "UNIQUE" in str(e) else e  # type: ignore[return-value]


def crear(slug: str, datos: Dict[str, Any]) -> Dict[str, Any]:
    conf = _conf(slug); _asegurar()
    d = _validar(conf, datos, parcial=False)
    with db.transaction() as c:
        if "orden" not in d:
            d["orden"] = (c.execute("SELECT COALESCE(MAX(orden),0)+1 FROM costeo_componentes WHERE modulo=?", (conf["modulo"],)).fetchone()[0])
        cols = ["modulo"] + list(d)
        try:
            cur = c.execute(f"INSERT INTO costeo_componentes ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})",
                            [conf["modulo"]] + list(d.values()))
        except Exception as e:
            raise _duplicado(e)
        return _uno(c, conf["modulo"], cur.lastrowid)


def editar(slug: str, cid: int, datos: Dict[str, Any]) -> Dict[str, Any]:
    conf = _conf(slug); _asegurar()
    d = _validar(conf, datos, parcial=True)
    with db.transaction() as c:
        _uno(c, conf["modulo"], cid)
        if d:
            sets = ", ".join(f"{k}=?" for k in d) + ", actualizado_en=datetime('now','localtime')"
            try:
                c.execute(f"UPDATE costeo_componentes SET {sets} WHERE id=?", list(d.values()) + [cid])
            except Exception as e:
                raise _duplicado(e)
        return _uno(c, conf["modulo"], cid)


def eliminar(slug: str, cid: int) -> None:
    conf = _conf(slug); _asegurar()
    with db.transaction() as c:
        _uno(c, conf["modulo"], cid)
        c.execute("DELETE FROM costeo_componentes WHERE id=?", (cid,))


def fijar_tipo_cambio(slug: str, valor: Any) -> Dict[str, Any]:
    conf = _conf(slug); _asegurar()
    tc = _num(valor, "tipo de cambio")
    if not 1 <= tc <= 1000:
        raise CosteoError("El tipo de cambio debe estar entre 1 y 1000.")
    with db.transaction() as c:
        c.execute("INSERT INTO costeo_meta (modulo, clave, valor) VALUES (?, 'tipo_cambio', ?) "
                  "ON CONFLICT(modulo, clave) DO UPDATE SET valor=excluded.valor", (conf["modulo"], str(tc)))
        return _resumen(c, conf)


# ----------------------------------------------------------------------------- Excel
def _texto(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def _filas_cotizador(contenido: bytes, conf: Dict[str, Any]) -> Dict[str, Any]:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    except Exception:
        raise CosteoError("No se pudo leer el archivo: sube un .xlsx del cotizador.")
    hoja = wb["Componentes"] if "Componentes" in wb.sheetnames else wb.worksheets[0]
    filas = list(hoja.iter_rows(values_only=True))
    cab = next((i for i, f in enumerate(filas) if any(_texto(x).lower().startswith("descripción del material") for x in f)), None)
    if cab is None:
        raise CosteoError("No encontré la tabla: la hoja «Componentes» debe tener la columna «Descripción del Material».")
    idx = {_texto(x).lower(): j for j, x in enumerate(filas[cab]) if x is not None}

    def col(*nombres: str) -> Optional[int]:
        return next((j for k, j in idx.items() for n in nombres if k.startswith(n)), None)
    c_tar, c_cat, c_desc = col("tipo de tarjeta"), col("categor"), col("descripción del material")
    c_lote, c_uni, c_cant = col("cantidad por caja"), col("unidad"), col("cantidad por ensamble")
    c_mon, c_costo = col("moneda"), col("costo unitario")
    if None in (c_desc, c_cant, c_costo):
        raise CosteoError("Faltan columnas: se necesitan «Descripción del Material», «Cantidad por ensamble» y «Costo unitario».")
    # proveedores y tipo de cambio desde «Dashboard de costos» (si existe)
    proveedores: Dict[str, str] = {}
    tc = None
    if "Dashboard de costos" in wb.sheetnames:
        for f in wb["Dashboard de costos"].iter_rows(values_only=True):
            f = list(f) + [None] * 8
            if isinstance(f[4], (int, float)) and _texto(f[3]).lower().startswith("tipo de cambio"):
                tc = float(f[4])
            if _texto(f[2]) and _texto(f[3]) and _texto(f[3]).lower() != "proveedor":
                proveedores.setdefault(_texto(f[2]).lower(), _texto(f[3]))
    out: List[Dict[str, Any]] = []
    get = lambda f, j: f[j] if j is not None and j < len(f) else None   # noqa: E731
    for n, f in enumerate(filas[cab + 1:], start=1):
        desc = _texto(get(f, c_desc))
        if not desc or not isinstance(get(f, c_costo), (int, float)):
            continue
        tipo_txt = _texto(get(f, c_tar))
        recurrente = "licencia" in desc.lower() or "anual" in tipo_txt.lower()
        tarjetas = [t for t in conf["tarjetas"] if re.search(rf"\b{re.escape(t)}\b", tipo_txt, re.I)] or \
            ([conf["tarjetas"][0]] if recurrente or len(conf["tarjetas"]) == 1 else list(conf["tarjetas"]))
        cant = float(get(f, c_cant) or 0) / len(tarjetas)
        lote = get(f, c_lote)
        for t in tarjetas:
            out.append({
                "tarjeta": t, "categoria": _texto(get(f, c_cat)), "descripcion": desc,
                "proveedor": proveedores.get(desc.lower(), ""),
                "cantidad_lote": float(lote) / len(tarjetas) if isinstance(lote, (int, float)) else 0.0,
                "unidad": _texto(get(f, c_uni)), "cantidad": round(cant, 6),
                "moneda": (_texto(get(f, c_mon)) or "MXN").upper(), "costo_unitario": float(get(f, c_costo)),
                "tipo": "recurrente" if recurrente else "produccion", "orden": n,
            })
    if not out:
        raise CosteoError("El archivo no trae materiales con costo.")
    return {"filas": out, "tipo_cambio": tc}


def importar_xlsx(slug: str, contenido: bytes, reemplazar: bool = True) -> Dict[str, Any]:
    """Carga el cotizador del módulo. reemplazar=True borra antes los componentes del módulo (no toca otros módulos)."""
    conf = _conf(slug); _asegurar()
    leido = _filas_cotizador(contenido, conf)
    with db.transaction() as c:
        if reemplazar:
            c.execute("DELETE FROM costeo_componentes WHERE modulo=?", (conf["modulo"],))
        for d in leido["filas"]:
            d = _validar(conf, d, parcial=False)
            cols = ["modulo"] + list(d)
            c.execute(f"INSERT INTO costeo_componentes ({','.join(cols)}) VALUES ({','.join('?' * len(cols))}) "
                      "ON CONFLICT(modulo, tarjeta, descripcion) DO UPDATE SET "
                      + ", ".join(f"{k}=excluded.{k}" for k in d) + ", actualizado_en=datetime('now','localtime')",
                      [conf["modulo"]] + list(d.values()))
        if leido["tipo_cambio"] and 1 <= leido["tipo_cambio"] <= 1000:
            c.execute("INSERT INTO costeo_meta (modulo, clave, valor) VALUES (?, 'tipo_cambio', ?) "
                      "ON CONFLICT(modulo, clave) DO UPDATE SET valor=excluded.valor", (conf["modulo"], str(leido["tipo_cambio"])))
        r = _resumen(c, conf)
    r["importados"] = len(leido["filas"])
    return r


def exportar_xlsx(slug: str) -> bytes:
    """Cotizador del módulo con fórmulas vivas: Componentes (subtotal y MXN con el tipo de cambio de Resumen!C3) y Resumen."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    r = resumen(slug)
    wb = openpyxl.Workbook()
    hc = wb.active; hc.title = "Componentes"
    hr = wb.create_sheet("Resumen")
    neg, fondo = Font(bold=True), PatternFill("solid", fgColor="FFE5E7EB")
    hc.append(["Tarjeta", "Categoría", "Descripción del Material", "Proveedor", "Cantidad por caja/bobina/lote", "Unidad",
               "Cantidad por tarjeta", "Moneda", "Costo unitario", "Subtotal", "Subtotal (MXN)", "Tipo", "Notas"])
    for celda in hc[1]:
        celda.font, celda.fill = neg, fondo
    for i, it in enumerate(r["items"], start=2):
        hc.append([it["tarjeta"], it["categoria"], it["descripcion"], it["proveedor"], it["cantidad_lote"], it["unidad"],
                   it["cantidad"], it["moneda"], it["costo_unitario"], f"=I{i}*G{i}", f'=IF(H{i}="USD",J{i}*Resumen!$C$3,J{i})',
                   "Recurrente" if it["tipo"] == "recurrente" else "Producción", it["notas"]])
    ult = max(2, len(r["items"]) + 1)
    for col, ancho in zip("ABCDEFGHIJKLM", (10, 24, 44, 20, 12, 8, 12, 8, 12, 12, 14, 12, 30)):
        hc.column_dimensions[col].width = ancho
    hr["B1"] = f"COSTOS DE PRODUCCIÓN | {r['nombre'].upper()}"; hr["B1"].font = Font(bold=True, size=14)
    hr["B3"] = "Tipo de cambio (MXN por USD)"; hr["C3"] = r["tipo_cambio"]
    fila = 5
    for t in r["tarjetas"]:
        hr[f"B{fila}"] = f"Costo tarjeta {t} (MXN)"
        hr[f"C{fila}"] = f'=SUMIFS(Componentes!K2:K{ult},Componentes!A2:A{ult},"{t}",Componentes!L2:L{ult},"Producción")'
        fila += 1
    hr[f"B{fila}"] = "COSTO DE PRODUCCIÓN TOTAL (MXN)"; hr[f"C{fila}"] = f"=SUM(C5:C{fila - 1})"
    hr[f"B{fila}"].font = hr[f"C{fila}"].font = neg
    hr[f"B{fila + 2}"] = "Costos recurrentes (aparte, MXN)"
    hr[f"C{fila + 2}"] = f'=SUMIFS(Componentes!K2:K{ult},Componentes!L2:L{ult},"Recurrente")'
    hr.column_dimensions["B"].width = 38; hr.column_dimensions["C"].width = 16
    for celda in ("C3",) + tuple(f"C{x}" for x in range(5, fila + 3)):
        hr[celda].number_format = "#,##0.00"
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()
