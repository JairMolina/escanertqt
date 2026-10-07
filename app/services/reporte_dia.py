"""Reporte de un día (v1.3.35): tarjetas COMPLETADAS (`fecha_finalizado`) y ENTREGADAS (`fecha_real`) en una fecha,
de todos los lotes, y su exportación a un Excel de fácil lectura (resumen + detalle con filtros)."""
import io
import re
from datetime import date
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.database import db
from app.database.models import MESES_ES

_FECHA = re.compile(r"^\d{4}-\d{2}-\d{2}$")
DIAS_ES = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]


MAX_DIAS = 366


def _validar(fecha: str) -> date:
    if not _FECHA.match(fecha or ""):
        raise ValueError("Fecha no válida (AAAA-MM-DD).")
    try:
        return date.fromisoformat(fecha)
    except ValueError:
        raise ValueError("Fecha no válida (AAAA-MM-DD).")


def validar_rango(desde: str, hasta: Optional[str] = None) -> Tuple[date, date]:
    """Rango de fechas (v1.3.41): `hasta` vacío = un solo día. Máximo un año."""
    d1 = _validar(desde)
    d2 = _validar(hasta) if hasta else d1
    if d2 < d1:
        raise ValueError("La fecha final no puede ser anterior a la inicial.")
    if (d2 - d1).days >= MAX_DIAS:
        raise ValueError(f"El rango máximo es de {MAX_DIAS} días.")
    return d1, d2


def nombre_archivo(desde: str, hasta: Optional[str] = None) -> str:
    d1, d2 = validar_rango(desde, hasta)
    return f"Tarjetas_{d1}.xlsx" if d1 == d2 else f"Tarjetas_{d1}_a_{d2}.xlsx"


def _corta(d: date) -> str:
    return f"{d.day} de {MESES_ES[d.month - 1].lower()} de {d.year}"


def texto_rango(desde: str, hasta: Optional[str] = None) -> str:
    """'7 de octubre de 2026' o 'del 1 al 6 de octubre de 2026' (para títulos y correos)."""
    d1, d2 = validar_rango(desde, hasta)
    if d1 == d2:
        return _corta(d1)
    if (d1.year, d1.month) == (d2.year, d2.month):
        return f"del {d1.day} al {_corta(d2)}"
    if d1.year == d2.year:
        return f"del {d1.day} de {MESES_ES[d1.month - 1].lower()} al {_corta(d2)}"
    return f"del {_corta(d1)} al {_corta(d2)}"


def tarjetas_del_dia(fecha: str, db_path: Optional[Path] = None) -> Dict[str, Any]:
    return tarjetas_del_rango(fecha, fecha, db_path)


def tarjetas_del_rango(desde: str, hasta: Optional[str] = None, db_path: Optional[Path] = None) -> Dict[str, Any]:
    d1, d2 = validar_rango(desde, hasta)
    a, b = d1.isoformat(), d2.isoformat()
    sql = """SELECT t.id, t.id_tarjeta_num, p1.nombre AS nombre_r1, p1.mac AS mac_r1, p2.nombre AS nombre_r2, p2.mac AS mac_r2,
                    p3.nombre AS nombre_r3, t.fecha_finalizado, t.fecha_real, t.gabinete, l.mes, l.anio
             FROM tarjetas_produccion t JOIN lotes_mensuales l ON l.id = t.lote_id
             LEFT JOIN pcb_inventario p1 ON p1.id = t.pcb_r1_id
             LEFT JOIN pcb_inventario p2 ON p2.id = t.pcb_r2_id
             LEFT JOIN pcb_inventario p3 ON p3.id = t.pcb_r3_id
             WHERE t.fecha_finalizado BETWEEN ? AND ? OR t.fecha_real BETWEEN ? AND ?
             ORDER BY l.anio, l.mes, t.id_tarjeta_num"""
    with db.get_db(db_path) as c:
        filas = [dict(r) for r in c.execute(sql, (a, b, a, b))]
    for f in filas:
        f["lote"] = f"{MESES_ES[f.pop('mes') - 1]} {f.pop('anio')}"
        f["completada"] = bool(f["fecha_finalizado"]) and a <= f["fecha_finalizado"] <= b
        f["entregada"] = bool(f["fecha_real"]) and a <= f["fecha_real"] <= b
    return {"fecha": a, "desde": a, "hasta": b, "dias": (d2 - d1).days + 1,
            "completadas": sum(f["completada"] for f in filas), "entregadas": sum(f["entregada"] for f in filas), "items": filas}


# ---------------------------------------------------------------- Excel
AZUL, AZUL_CLARO, VERDE, VERDE_CLARO, GRIS, BORDE = "1F3A5F", "E8EEF6", "1E7B4F", "E6F4EC", "667085", "D0D7E2"
F = "Arial"


def _fecha_larga(d: date) -> str:
    return f"{DIAS_ES[d.weekday()].capitalize()} {d.day} de {MESES_ES[d.month - 1].lower()} de {d.year}"


def _dia(iso: Optional[str]) -> Any:
    try:
        return date.fromisoformat(iso) if iso else "—"
    except ValueError:
        return iso


def _formato_hoja(ws, anchos: List[int], congelar: str, titulos: str) -> None:
    for i, a in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = a
    ws.freeze_panes = congelar
    ws.print_title_rows = titulos
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _hoja_por_dia(wb: Workbook, d1: date, d2: date, filas: List[Dict[str, Any]]) -> None:
    """Segunda hoja de un rango: cuántas se completaron y entregaron cada día (fórmulas sobre la hoja de detalle)."""
    ws = wb.create_sheet("Por día")
    ws.sheet_view.showGridLines = False
    fino = Side(style="thin", color=BORDE)
    ws.merge_cells("A1:D1")
    ws["A1"] = "Tarjetas por día"
    ws["A1"].font = Font(name=F, size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=AZUL)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 28
    for i, t in enumerate(["Fecha", "Día", "Completadas", "Entregadas"], 1):
        c = ws.cell(row=3, column=i, value=t)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center")
    ultima = 8 + max(len(filas), 1)
    det = "'Reporte del periodo'"
    r, d = 4, d1
    while d <= d2:
        ws.cell(row=r, column=1, value=d).number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=2, value=DIAS_ES[d.weekday()].capitalize())
        ws.cell(row=r, column=3, value=f"=COUNTIF({det}!K9:K{ultima},A{r})")
        ws.cell(row=r, column=4, value=f"=COUNTIF({det}!L9:L{ultima},A{r})")
        for col in range(1, 5):
            c = ws.cell(row=r, column=col)
            c.font = Font(name=F, size=10, color="1B2333")
            c.border = Border(bottom=fino)
            c.alignment = Alignment(horizontal="center")
        r += 1
        d = date.fromordinal(d.toordinal() + 1)
    ws.cell(row=r, column=1, value="Total")
    ws.cell(row=r, column=3, value=f"=SUM(C4:C{r - 1})")
    ws.cell(row=r, column=4, value=f"=SUM(D4:D{r - 1})")
    for col in range(1, 5):
        c = ws.cell(row=r, column=col)
        c.font = Font(name=F, bold=True, color=AZUL)
        c.fill = PatternFill("solid", fgColor=AZUL_CLARO)
        c.alignment = Alignment(horizontal="center")
    _formato_hoja(ws, [14, 14, 14, 14], "A4", "3:3")


def excel_del_dia(fecha: str, db_path: Optional[Path] = None) -> bytes:
    return excel_del_rango(fecha, fecha, db_path)


def excel_del_rango(desde: str, hasta: Optional[str] = None, db_path: Optional[Path] = None) -> bytes:
    d1, d2 = validar_rango(desde, hasta)
    un_dia = d1 == d2
    datos = tarjetas_del_rango(desde, hasta, db_path)
    filas: List[Dict[str, Any]] = datos["items"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte del día" if un_dia else "Reporte del periodo"
    ws.sheet_view.showGridLines = False
    fino = Side(style="thin", color=BORDE)

    # Encabezado
    ws.merge_cells("A1:L1")
    ws["A1"] = "Escáner TQT · Reporte de tarjetas del " + ("día" if un_dia else "periodo")
    ws["A1"].font = Font(name=F, size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=AZUL)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    ws.merge_cells("A2:L2")
    ws["A2"] = _fecha_larga(d1) if un_dia else f"Del {_fecha_larga(d1).lower()} al {_fecha_larga(d2).lower()} · {datos['dias']} días"
    ws["A2"].font = Font(name=F, size=11, italic=True, color=GRIS)
    ws["A2"].alignment = Alignment(indent=1)

    # Tarjetas KPI (fórmulas: se recalculan si alguien edita el detalle)
    n = len(filas)
    ini, fin = 9, 9 + max(n, 1)   # rango de datos de la tabla (fila 8 = encabezados)
    kpis = [("B", "D", "Completadas", f'=COUNTIF(H{ini}:H{fin - 1},"Sí")', VERDE, VERDE_CLARO),
            ("E", "G", "Entregadas", f'=COUNTIF(I{ini}:I{fin - 1},"Sí")', AZUL, AZUL_CLARO),
            ("H", "L", "Tarjetas en el reporte", f"=COUNTA(A{ini}:A{fin - 1})", GRIS, "F2F4F7")]
    for c1, c2, titulo, formula, color, fondo in kpis:
        ws.merge_cells(f"{c1}4:{c2}4")
        ws.merge_cells(f"{c1}5:{c2}5")
        ws[f"{c1}4"] = titulo
        ws[f"{c1}4"].font = Font(name=F, size=10, bold=True, color=color)
        ws[f"{c1}5"] = formula if n else 0
        ws[f"{c1}5"].font = Font(name=F, size=24, bold=True, color=color)
        for fila in (4, 5):
            for col in range(ws[f"{c1}4"].column, ws[f"{c2}4"].column + 1):
                cel = ws.cell(row=fila, column=col)
                cel.fill = PatternFill("solid", fgColor=fondo)
                cel.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[5].height = 36

    ws["A7"] = "Detalle"
    ws["A7"].font = Font(name=F, size=12, bold=True, color=AZUL)
    cab = ["Tarjeta", "Lote", "R1", "MAC R1", "R2", "MAC R2", "R3", "Completada", "Entregada", "Gabinete", "Fecha finalizado", "Fecha entrega"]
    for i, t in enumerate(cab, 1):
        c = ws.cell(row=8, column=i, value=t)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[8].height = 22
    for r, f in enumerate(filas, ini):
        valores = ["#" + str(f["id_tarjeta_num"]), f["lote"], f["nombre_r1"] or "—", f["mac_r1"] or "—", f["nombre_r2"] or "—", f["mac_r2"] or "—",
                   f["nombre_r3"] or "—", "Sí" if f["completada"] else "No", "Sí" if f["entregada"] else "No", f["gabinete"] or "—",
                   _dia(f["fecha_finalizado"]), _dia(f["fecha_real"])]
        for i, v in enumerate(valores, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(name=F, size=10, color=(VERDE if v == "Sí" else GRIS if v in ("No", "—") else "1B2333"), bold=(v == "Sí"))
            c.border = Border(bottom=fino)
            c.alignment = Alignment(horizontal="center" if i in (1, 8, 9, 10, 11, 12) else "left", vertical="center")
            if isinstance(v, date):
                c.number_format = "dd/mm/yyyy"
    if n:
        tabla = Table(displayName="TarjetasDelDia", ref=f"A8:L{fin - 1}")
        tabla.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(tabla)
    else:
        ws.merge_cells("A9:L9")
        ws["A9"] = "No hubo tarjetas completadas ni entregadas en " + ("esta fecha." if un_dia else "este periodo.")
        ws["A9"].font = Font(name=F, italic=True, color=GRIS)
        ws["A9"].alignment = Alignment(horizontal="center")

    nota = fin + 1
    ws.merge_cells(f"A{nota}:L{nota}")
    ws[f"A{nota}"] = ("Completada = la tarjeta quedó con R1, R2, R3 y MAC de R1 y R2 dentro del " + ("día" if un_dia else "periodo") + " (fecha de finalizado). "
                      "Entregada = fecha real de entrega dentro del " + ("día" if un_dia else "periodo") + ". Fuente: base de datos de Escáner TQT.")
    ws[f"A{nota}"].font = Font(name=F, size=9, italic=True, color=GRIS)
    ws[f"A{nota}"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[nota].height = 28

    _formato_hoja(ws, [11, 17, 22, 20, 22, 20, 22, 13, 12, 14, 15, 15], "A9", "8:8")
    if not un_dia:
        _hoja_por_dia(wb, d1, d2, filas)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
