import _aislamiento  # noqa: F401  (debe ser lo primero)
"""v1.4.34: Costos por módulo. Cada módulo tiene sus materiales y su tipo de cambio; costo por tarjeta, recurrentes aparte,
importar el formato del cotizador (filas «RX y TX» repartidas) y exportar con fórmulas."""
import io
import unittest

import openpyxl

from app.database import db
from app.main import app
from app.modulos import ANTIJAMMER, SIGFOX, modulo_de_ruta
from app.services import usuarios

from _aislamiento import cliente_sin_sesion, verificar_aislamiento


def cotizador(filas, tc=20.0) -> bytes:
    """Excel con el formato de los cotizadores: hoja Componentes (encabezado en fila 5) y Dashboard de costos."""
    wb = openpyxl.Workbook()
    hc = wb.active; hc.title = "Componentes"
    hc["B2"] = "LISTA DE COMPONENTES"
    hc.append([]); hc.append([])
    hc.append(["Tipo de tarjeta", "Categoria", "Descripción del Material", "Cantidad por caja/bobina/lote", "Unidad de medida",
               "Cantidad por ensamble", "Rendimiento por caja/bobina/lote", "Moneda", "Costo unitario", "Subtotal", "Subtotal (MXN)"])
    for f in filas:
        hc.append(list(f[:6]) + [None] + list(f[6:8]))
    hd = wb.create_sheet("Dashboard de costos")
    hd["D5"] = "Tipo de Cambio Referencia:"; hd["E5"] = tc
    hd["C9"] = "Arduino Nano"; hd["D9"] = "Kuma Electrónica"
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


class TestCosteo(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        verificar_aislamiento()
        db.init_db()
        usuarios.sembrar(iteraciones=1000)
        usuarios.limitador.reiniciar()
        cls.c = cliente_sin_sesion(app)
        r = cls.c.post("/api/auth/login", json={"email": "developer@skyguardian.mx", "password": usuarios.CLAVE_INICIAL})
        assert r.status_code == 200, r.text

    def importar(self, slug, datos):
        return self.c.post(f"/api/{slug}/costos/importar", content=datos, headers={"Content-Type": "application/octet-stream"})

    def test_rutas_pertenecen_a_su_modulo(self):
        self.assertEqual(modulo_de_ruta("/api/antijammer/costos"), ANTIJAMMER)
        self.assertEqual(modulo_de_ruta("/api/sigfox/costos/exportar"), SIGFOX)
        self.assertEqual(modulo_de_ruta("/api/tpms/costos"), "tpms")
        self.assertEqual(self.c.get("/api/caiquen/costos").json()["tarjetas"], ["Caiquén"])   # v1.4.37

    def test_antijammer_reparte_rx_y_tx_y_calcula_por_tarjeta(self):
        r = self.importar("antijammer", cotizador([
            ("Tarjeta RX", "Tarjeta PCB", "PCB RX", 1, "pza", 1, "USD", 10),
            ("Tarjeta TX", "Tarjeta PCB", "PCB TX", 1, "pza", 1, "USD", 10),
            ("Tarjeta RX y TX", "Componentes electrónicos", "Arduino Nano", 2, "pza", 2, "MXN", 250),
            ("Tarjeta RX", "Accesorios", "Relevador", 10, "pza", 1, "MXN", 45),
        ]))
        self.assertEqual(r.status_code, 200, r.text)
        d = r.json()
        self.assertEqual(d["importados"], 5)
        self.assertEqual(d["tipo_cambio"], 20.0)
        self.assertEqual(d["por_tarjeta"], {"RX": 200 + 250 + 45, "TX": 200 + 250})   # relevador solo RX, sin dividir
        self.assertEqual(d["total_produccion"], 945)
        nano = [x for x in d["items"] if x["descripcion"] == "Arduino Nano"]
        self.assertEqual(sorted(x["tarjeta"] for x in nano), ["RX", "TX"])
        self.assertTrue(all(x["cantidad"] == 1 and x["proveedor"] == "Kuma Electrónica" for x in nano))
        # el tipo de cambio es del módulo y recalcula
        d = self.c.patch("/api/antijammer/costos/parametros", json={"tipo_cambio": 10}).json()
        self.assertEqual(d["por_tarjeta"]["TX"], 100 + 250)

    def test_sigfox_licencia_aparte_y_modulos_independientes(self):
        r = self.importar("sigfox", cotizador([
            ("Tarjeta Sigfox", "Tarjeta PCB", "PCB SKUDO SIGFOX", 1, "pza", 1, "USD", 8),
            ("Tarjeta Sigfox", "Componentes electrónicos", "Módulo Sigfox", 1, "pza", 1, "USD", 27.5),
            ("Licencia (anual)", "Otros", "Licencia sigfox 1 año", 1, "pza", 1, "USD", 5),
        ], tc=18.5))
        self.assertEqual(r.status_code, 200, r.text)
        d = r.json()
        self.assertEqual(d["por_tarjeta"], {"Sigfox": round((8 + 27.5) * 18.5, 2)})
        self.assertEqual(d["recurrente"], 92.5)
        self.assertNotIn("PCB RX", {x["descripcion"] for x in d["items"]})   # no ve materiales de AntiJammer

    def test_tpms_y_speedlocker_una_tarjeta(self):
        """v1.4.36: TPMS y SpeedLocker tienen su propio cotizador; todo cae en su única tarjeta."""
        for slug, tipo, tarjeta in (("tpms", "TPMS", "TPMS"), ("speedlocker", "Tarjeta Speedlocker", "SpeedLocker")):
            r = self.importar(slug, cotizador([
                (tipo, "Tarjeta PCB", f"PCB {tarjeta}", 1, "pza", 1, "USD", 18),
                (tipo, "Cable", "Cable calibre 22 | Rojo", 100, "m", 0.5, "MXN", 4.26),
            ], tc=18.5))
            self.assertEqual(r.status_code, 200, r.text)
            self.assertEqual(r.json()["por_tarjeta"], {tarjeta: round(18 * 18.5 + 2.13, 2)})
            self.assertEqual(self.c.get(f"/api/{slug}/materiales").json()["tarjetas"], [tarjeta])

    def test_crud_validaciones_y_exportar(self):
        b = "/api/sigfox/costos"
        r = self.c.post(b, json={"descripcion": "Caja de protección", "cantidad": 1, "costo_unitario": 52, "moneda": "MXN"})
        self.assertEqual(r.status_code, 201, r.text)
        cid = r.json()["id"]
        self.assertEqual(self.c.post(b, json={"descripcion": "Caja de protección"}).status_code, 422)   # duplicado
        self.assertEqual(self.c.post(b, json={"descripcion": "X", "tarjeta": "RX"}).status_code, 422)  # tarjeta de otro módulo
        self.assertEqual(self.c.post(b, json={"descripcion": "X", "costo_unitario": -1}).status_code, 422)
        self.assertEqual(self.c.patch(f"{b}/{cid}", json={"costo_unitario": 60}).json()["costo_unitario"], 60)
        self.assertEqual(self.c.patch(f"/api/antijammer/costos/{cid}", json={"costo_unitario": 1}).status_code, 404)
        x = self.c.get(f"{b}/exportar")
        self.assertEqual(x.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(x.content))
        self.assertEqual(wb.sheetnames, ["Componentes", "Resumen"])
        self.assertTrue(str(wb["Componentes"]["K2"].value).startswith("=IF("))
        self.assertEqual(self.c.delete(f"{b}/{cid}").status_code, 204)
        self.assertEqual(self.importar("sigfox", b"no es excel").status_code, 422)


if __name__ == "__main__":
    unittest.main()
