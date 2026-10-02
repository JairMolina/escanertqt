"""v1.1.12: Excel con Botón R3 (Producción!H, Catálogo K:M y O3#), exportación fiel a la referencia e importación con validación de R3."""
import _aislamiento  # noqa: F401  (debe ser lo primero)
import re
import subprocess
import tempfile
import unittest
import zipfile
from pathlib import Path

import openpyxl

from app.database import db
from app.services.excel_sync import ExcelSyncEngine, structure_counts

from _aislamiento import crear_par, verificar_aislamiento

RAIZ = Path(__file__).resolve().parents[1]


def _formulas(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    out = {}
    for ws in wb:
        for fila in ws.iter_rows():
            for c in fila:
                v = c.value
                if hasattr(v, "text") or (isinstance(v, str) and v.startswith("=")):
                    out[f"{ws.title}!{c.coordinate}"] = v.text if hasattr(v, "text") else v
    return out


def _zip_txt(path, nombre):
    with zipfile.ZipFile(path) as z:
        return z.read(nombre).decode("utf-8")


class TestExcelR3(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        verificar_aislamiento()
        cls.tmp = tempfile.TemporaryDirectory()
        cls.dir = Path(cls.tmp.name)
        cls.template = str(ExcelSyncEngine().get_template_path())
        cls.n = 0

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def setUp(self):
        type(self).n += 1
        self.path = self.dir / f"r3_{self.n}.db"
        db.init_db(self.path)
        self.engine = ExcelSyncEngine(db_path=self.path)
        self.excel = self.engine.create_monthly_excel(9, 2026, target_path=str(self.dir / f"m{self.n}.xlsx"))
        self.lote = db.create_lote("R3", 9, 2026, ruta_excel=self.excel, activo=True, db_path=self.path)["id"]

    def _con_r3(self, num="0011"):
        return crear_par(self.lote, num=num, r3=True, path=self.path)

    # ---------------------------------------------------------------- exportación
    def test_01_exporta_r3_en_produccion_h_y_catalogo(self):
        self._con_r3("0011")
        res = self.engine.export_to_excel(self.excel, self.lote, db_path=self.path)
        self.assertFalse(any("R3" in a for a in res["avisos"]), res["avisos"])
        wb = openpyxl.load_workbook(self.excel)
        prod, cat = wb["Producción"], wb["Catálogo PCB"]
        self.assertEqual(prod["H1"].value, "Botón R3")
        self.assertEqual(prod["A2"].value, "0011")
        self.assertEqual(prod["H2"].value, "TQT-R3-V30-0011")
        self.assertEqual(prod["H2"].number_format, "@")
        fila = next(r for r in range(3, 103) if cat.cell(r, 11).value == "TQT-R3-V30-0011")
        self.assertIn(cat.cell(fila, 12).value, ("SIN PROBAR", "FUNCIONAL", "NO FUNCIONAL", "REPARACIÓN"))
        self.assertTrue(str(cat.cell(fila, 13).value).startswith("=IF(K"))   # 'Asignada a' sigue siendo fórmula
        self.assertEqual([cat.cell(2, c).value for c in (11, 12, 13)], ["Nombre PCB", "Estado PCB", "Asignada a"])   # R3 sin MAC

    def test_02_estructura_identica_a_la_referencia(self):
        self._con_r3("0011")
        self.engine.export_to_excel(self.excel, self.lote, db_path=self.path)
        self.assertEqual(structure_counts(self.template), structure_counts(self.excel))
        f0, f1 = _formulas(self.template), _formulas(self.excel)
        self.assertEqual(len(f1), 2211)
        self.assertEqual(f0, f1)                       # ni una fórmula cambió (incluidas Etiquetas!G -> Producción!I)
        self.assertIn("Producción!I2", f1["Etiquetas!G17"])
        self.assertEqual(f1["Catálogo PCB!O3"], '=_xlfn._xlws.FILTER(K3:K102,M3:M102="DISPONIBLE","SIN DISPONIBLES")')

    def test_03_xml_filter_dinamico_metadata_y_nombres(self):
        self._con_r3("0011")
        self.engine.export_to_excel(self.excel, self.lote, db_path=self.path)
        cat = _zip_txt(self.excel, "xl/worksheets/sheet4.xml")
        self.assertRegex(cat, r'<c r="O3"[^>]*\scm="1"[^>]*><f t="array" ref="O3:O102">_xlfn\._xlws\.FILTER\(')
        with zipfile.ZipFile(self.excel) as z:
            self.assertIn("xl/metadata.xml", z.namelist())
        self.assertIn("XLDAPR", _zip_txt(self.excel, "xl/metadata.xml"))
        self.assertIn("sheetMetadata", _zip_txt(self.excel, "[Content_Types].xml"))
        prod = _zip_txt(self.excel, "xl/worksheets/sheet1.xml")
        self.assertIn("_xlfn.ANCHORARRAY('Catálogo PCB'!$O$3)", prod)               # validación de H (Stop, mensaje 'Botón R3')
        self.assertIn('promptTitle="Botón R3"', prod)
        self.assertIn("COUNTIF('Catálogo PCB'!$K$3:$K$102,H2)=0", prod)              # CF de R3 inexistente
        wbxml = _zip_txt(self.excel, "xl/workbook.xml")
        self.assertIn('fullCalcOnLoad="1"', wbxml)
        self.assertIn("FirmwareList", wbxml)
        self.assertIn('name="ProduccionTQT"', _zip_txt(self.excel, "xl/tables/table1.xml"))
        self.assertIn('ref="A1:K101"', _zip_txt(self.excel, "xl/tables/table1.xml"))

    # ---------------------------------------------------------------- importación
    def test_04_ciclo_exportar_importar_exportar(self):
        for n in ("0011", "0012"):
            self._con_r3(n)
        self.engine.export_to_excel(self.excel, self.lote, db_path=self.path)

        db2 = self.dir / f"r3b_{self.n}.db"
        db.init_db(db2)
        motor2 = ExcelSyncEngine(db_path=db2)
        lote2 = db.create_lote("R3b", 9, 2026, activo=True, db_path=db2)["id"]
        res = motor2.import_from_excel(self.excel, lote2, db_path=db2)
        self.assertEqual(res["errores"], [], res)
        self.assertEqual(res["imported_tarjetas"], 2)
        tarjetas, _ = db.list_tarjetas(lote_id=lote2, limit=100, db_path=db2)
        self.assertEqual(sorted(t["nombre_r3"] for t in tarjetas), ["TQT-R3-V30-0011", "TQT-R3-V30-0012"])
        self.assertTrue(all(t["mac_r3"] is None for t in tarjetas))

        excel2 = motor2.create_monthly_excel(9, 2026, target_path=str(self.dir / f"m2_{self.n}.xlsx"))
        motor2.export_to_excel(excel2, lote2, db_path=db2)
        a, b = openpyxl.load_workbook(self.excel), openpyxl.load_workbook(excel2)
        for hoja in ("Producción", "Catálogo PCB"):
            for fila_a, fila_b in zip(a[hoja].iter_rows(), b[hoja].iter_rows()):
                for ca, cb in zip(fila_a, fila_b):
                    va = ca.value.text if hasattr(ca.value, "text") else ca.value
                    vb = cb.value.text if hasattr(cb.value, "text") else cb.value
                    self.assertEqual(va, vb, f"{hoja}!{ca.coordinate}")
        self.assertEqual(structure_counts(self.excel), structure_counts(excel2))

    def test_05_conflictos_de_r3_se_reportan(self):
        for n in ("0011", "0012", "0013"):
            self._con_r3(n)
        self.engine.export_to_excel(self.excel, self.lote, db_path=self.path)
        wb = openpyxl.load_workbook(self.excel)
        prod = wb["Producción"]
        prod["H3"].value = "TQT-R3-V30-0011"      # repetida (0012 usa la de 0011)
        prod["H4"].value = "TQT-R3-V30-9999"      # no existe en el catálogo
        conf = self.dir / f"conf{self.n}.xlsx"
        wb.save(conf)
        db2 = self.dir / f"r3c_{self.n}.db"
        db.init_db(db2)
        motor2 = ExcelSyncEngine(db_path=db2)
        lote2 = db.create_lote("R3c", 9, 2026, activo=True, db_path=db2)["id"]
        res = motor2.import_from_excel(str(conf), lote2, db_path=db2)
        texto = " ".join(res["errores"])
        self.assertIn("ya está en la tarjeta 0011", texto)
        self.assertIn("no existe en el Catálogo PCB", texto)
        self.assertEqual(res["status"], "partial")

    def test_06_excel_anterior_sin_r3_se_importa(self):
        viejo = self.dir / f"v0_{self.n}.xlsx"
        viejo.write_bytes(subprocess.run(["git", "show", "v1.1.11:templates/Control_Produccion_TQT_Template.xlsx"],
                                         cwd=RAIZ, capture_output=True, check=True).stdout)
        crear_par(self.lote, num="0011", path=self.path)
        self.engine.export_to_excel(str(viejo), self.lote, db_path=self.path)   # exporta en su layout anterior
        db2 = self.dir / f"r3d_{self.n}.db"
        db.init_db(db2)
        motor2 = ExcelSyncEngine(db_path=db2)
        lote2 = db.create_lote("R3d", 9, 2026, activo=True, db_path=db2)["id"]
        res = motor2.import_from_excel(str(viejo), lote2, db_path=db2)
        self.assertEqual(res["errores"], [], res)
        self.assertEqual(res["imported_tarjetas"], 1)
        tarjetas, _ = db.list_tarjetas(lote_id=lote2, limit=10, db_path=db2)
        self.assertIsNone(tarjetas[0]["r3"])

    def test_07_encabezados_distintos_dan_error_claro(self):
        wb = openpyxl.load_workbook(self.excel)
        wb["Producción"]["E1"].value = "Respaldo"
        malo = self.dir / f"malo{self.n}.xlsx"
        wb.save(malo)
        with self.assertRaises(ValueError) as cm:
            self.engine.import_from_excel(str(malo), self.lote, db_path=self.path)
        self.assertIn("Producción!E1", str(cm.exception))


if __name__ == "__main__":
    unittest.main()
