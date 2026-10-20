"""v1.4.14: módulo Impresora, registro de activos con etiqueta DYMO «ACT 0000» y QR."""
import _aislamiento  # noqa: F401  (debe ser lo primero)
import io
import re
import unittest
from unittest import mock
from urllib.parse import quote

from app.database import db
from app.main import app
from app.services import activos as svc
from app.services import dymo_remoto, usuarios

from _aislamiento import cliente_sin_sesion, verificar_aislamiento


def entrar(email="developer@skyguardian.mx"):
    c = cliente_sin_sesion(app)
    r = c.post("/api/auth/login", json={"email": email, "password": usuarios.CLAVE_INICIAL})
    assert r.status_code == 200, r.text
    return c


class TestActivos(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        verificar_aislamiento()
        db.init_db()
        usuarios.sembrar(iteraciones=1000)

    def setUp(self):
        usuarios.limitador.reiniciar()
        self.c = entrar()

    def test_requiere_sesion(self):
        self.assertEqual(cliente_sin_sesion(app).get("/api/activos").status_code, 401)

    def test_alta_con_y_sin_numero_y_conflicto(self):
        r = self.c.post("/api/activos", json={"numero": 8001, "nombre": "Osciloscopio Rigol", "categoria": "Instrumentos",
                                              "marca": "Rigol", "modelo": "DS1054Z", "serie": "DS1ZA0001", "ubicacion": "Laboratorio",
                                              "responsable": "Ana", "fecha_adquisicion": "2026-03-15", "valor": 7500.5, "notas": "Con sondas"})
        self.assertEqual(r.status_code, 201, r.text)
        a = r.json()
        self.assertEqual((a["numero"], a["codigo"], a["etiqueta"], a["estado"]), (8001, "ACT-8001", "ACT 8001", "activo"))
        self.assertEqual((a["valor"], a["fecha_adquisicion"], a["impresiones"], a["ultima_impresion"]), (7500.5, "2026-03-15", 0, None))
        self.assertTrue(a["qr"].endswith("/impresora?activo=ACT-8001"), a["qr"])
        self.assertEqual(a["creado_por"], "developer@skyguardian.mx")
        # sin número: max + 1
        r2 = self.c.post("/api/activos", json={"nombre": "Fuente de poder"})
        self.assertEqual(r2.status_code, 201, r2.text)
        self.assertGreater(r2.json()["numero"], 8001)
        # número repetido
        r3 = self.c.post("/api/activos", json={"numero": 8001, "nombre": "Otro"})
        self.assertEqual(r3.status_code, 409)
        self.assertIn("ACT-8001", r3.json()["detail"])
        self.assertEqual([e["evento"] for e in self.c.get("/api/activos/ACT-8001").json()["eventos"]], ["ALTA"])

    def test_validacion_422(self):
        for cuerpo in ({}, {"nombre": ""}, {"nombre": "x" * 121}, {"nombre": "A", "estado": "roto"}, {"nombre": "A", "valor": -1},
                       {"nombre": "A", "fecha_adquisicion": "15/03/2026"}, {"nombre": "A", "fecha_adquisicion": "2026-02-31"},
                       {"nombre": "A", "numero": 0}):
            self.assertEqual(self.c.post("/api/activos", json=cuerpo).status_code, 422, cuerpo)

    def test_siguiente_y_catalogos(self):
        self.c.post("/api/activos", json={"numero": 8010, "nombre": "Multímetro", "categoria": "Medición", "ubicacion": "Banco 2", "responsable": "Luis"})
        s = self.c.get("/api/activos/siguiente").json()
        self.assertGreaterEqual(s["numero"], 8011)
        self.assertEqual(s["codigo"], f"ACT-{s['numero']:04d}")
        self.assertEqual(s["etiqueta"], f"ACT {s['numero']:04d}")
        cat = self.c.get("/api/activos/catalogos").json()
        self.assertIn("Medición", cat["categorias"])
        self.assertIn("Banco 2", cat["ubicaciones"])
        self.assertIn("Luis", cat["responsables"])
        self.assertEqual([e["id"] for e in cat["estados"]], ["activo", "resguardo", "reparacion", "baja"])

    def test_listado_filtros_y_orden(self):
        self.c.post("/api/activos", json={"numero": 8020, "nombre": "Soldador Hakko", "marca": "Hakko", "serie": "ZX-9", "estado": "reparacion"})
        self.c.post("/api/activos", json={"numero": 8021, "nombre": "Cautín", "marca": "Weller", "ubicacion": "Taller 100%"})
        d = self.c.get("/api/activos", params={"q": "hakko"}).json()
        self.assertEqual([x["numero"] for x in d["items"]], [8020])
        self.assertEqual(d["total"], 1)
        self.assertEqual(self.c.get("/api/activos", params={"q": "ACT 8021"}).json()["items"][0]["numero"], 8021)
        self.assertEqual(self.c.get("/api/activos", params={"q": "100%"}).json()["total"], 1)
        self.assertEqual([x["numero"] for x in self.c.get("/api/activos", params={"estado": "reparacion", "q": "ZX"}).json()["items"]], [8020])
        todos = [x["numero"] for x in self.c.get("/api/activos", params={"limite": 500}).json()["items"]]
        self.assertEqual(todos, sorted(todos, reverse=True))
        self.assertEqual(len(self.c.get("/api/activos", params={"limite": 1}).json()["items"]), 1)
        self.assertEqual(self.c.get("/api/activos", params={"estado": "nada"}).status_code, 422)
        self.assertIn("siguiente", d)

    def test_busqueda_flexible_incluida_url(self):
        self.c.post("/api/activos", json={"numero": 8030, "nombre": "Analizador lógico"})
        for texto in ("ACT-8030", "ACT 8030", "act8030", "8030", "  act_8030 ", "https://tqt.example.mx/impresora?activo=ACT-8030",
                      "http://x/impresora?foo=1&activo=ACT%208030#z"):
            r = self.c.get("/api/activos/buscar", params={"codigo": texto})
            self.assertEqual(r.status_code, 200, texto)
            self.assertEqual(r.json()["codigo"], "ACT-8030")
            self.assertIn("eventos", r.json())
        self.assertEqual(self.c.get("/api/activos/8030").json()["numero"], 8030)
        self.assertEqual(self.c.get("/api/activos/" + quote("ACT 8030")).status_code, 200)
        for malo in ("ACT-9999", "hola", "0"):
            self.assertEqual(self.c.get("/api/activos/buscar", params={"codigo": malo}).status_code, 404, malo)
            self.assertEqual(self.c.get("/api/activos/" + malo).status_code, 404, malo)

    def test_patch_parcial_evento_y_conflicto(self):
        self.c.post("/api/activos", json={"numero": 8040, "nombre": "Router", "ubicacion": "Rack"})
        self.c.post("/api/activos", json={"numero": 8041, "nombre": "Switch"})
        r = self.c.patch("/api/activos/ACT-8040", json={"ubicacion": "Oficina", "estado": "resguardo", "valor": 1200})
        self.assertEqual(r.status_code, 200, r.text)
        a = r.json()
        self.assertEqual((a["nombre"], a["ubicacion"], a["estado"], a["valor"]), ("Router", "Oficina", "resguardo", 1200))
        ev = a["eventos"][0]
        self.assertEqual(ev["evento"], "EDICION")
        self.assertIn("ubicacion: Rack → Oficina", ev["detalle"])
        self.assertEqual(self.c.patch("/api/activos/8040", json={"numero": 8041}).status_code, 409)
        r = self.c.patch("/api/activos/8040", json={"numero": 8042, "valor": None})
        self.assertEqual((r.status_code, r.json()["codigo"], r.json()["valor"]), (200, "ACT-8042", None))
        self.assertEqual(self.c.get("/api/activos/ACT-8040").status_code, 404)
        self.assertEqual(self.c.patch("/api/activos/ACT-8042", json={"nombre": ""}).status_code, 422)
        self.assertEqual(self.c.patch("/api/activos/ACT-8042", json={"nombre": None}).status_code, 422)
        self.assertEqual(self.c.patch("/api/activos/ACT-0404", json={"nombre": "x"}).status_code, 404)
        # sin cambios reales: no genera evento nuevo
        n = len(self.c.get("/api/activos/ACT-8042").json()["eventos"])
        self.c.patch("/api/activos/ACT-8042", json={"nombre": "Router"})
        self.assertEqual(len(self.c.get("/api/activos/ACT-8042").json()["eventos"]), n)

    def test_etiqueta_svg_xml_y_archivo(self):
        self.c.post("/api/activos", json={"numero": 8050, "nombre": "Impresora 3D Prusa MK4 con enclosure y extras"})
        r = self.c.get("/api/activos/ACT-8050/etiqueta.svg")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.headers["content-type"], "image/svg+xml")
        self.assertEqual(re.findall(r"<text[^>]*>([^<]*)</text>", r.text), ["ACT", "8050"])   # v1.4.18: dos líneas
        ys = [float(y) for y in re.findall(r'<text x="[\d.]+" y="([\d.]+)"', r.text)]
        sizes = [float(v) for v in re.findall(r'font-size="([\d.]+)"', r.text)]
        self.assertLess(ys[0], ys[1])
        self.assertGreater(sizes[1], sizes[0])   # el número es la línea dominante
        self.assertIn('width="57mm" height="32mm"', r.text)
        self.assertNotIn("Prusa", r.text)   # v1.4.16: la etiqueta lleva solo «ACT 0000» y el QR, sin nombre
        for tipo in ("dymo", "label"):
            x = self.c.get("/api/activos/ACT-8050/etiqueta/xml", params={"tipo": tipo}).text
            self.assertNotIn("Prusa", x)
            self.assertNotIn("NOMBRE", x)
            self.assertNotIn("DETALLE", x)
        self.assertIn('rx="2.5"', self.c.get("/api/activos/8050/etiqueta.svg", params={"borde": "true"}).text)
        for tipo, raiz in (("dymo", "<DesktopLabel"), ("label", "<DieCutLabel")):
            x = self.c.get("/api/activos/ACT-8050/etiqueta/xml", params={"tipo": tipo})
            self.assertEqual(x.status_code, 200, tipo)
            self.assertIn(raiz, x.text)
            etiqueta_texto = "<Text>{}</Text>" if tipo == "dymo" else "<String>{}</String>"
            self.assertIn(etiqueta_texto.format("ACT"), x.text)
            self.assertIn(etiqueta_texto.format("8050"), x.text)
            self.assertNotIn("ACT 8050", x.text)
            self.assertIn("<Name>QR</Name>", x.text)
            self.assertIn("<Name>NUMERO</Name>", x.text)
        dcd = self.c.get("/api/activos/ACT-8050/etiqueta/xml").text
        self.assertRegex(dcd, r"<DataString>https?://[^<]+/impresora\?activo=ACT-8050</DataString>")
        for fmt, ext in (("dymo", "dymo"), ("label", "label")):
            a = self.c.get("/api/activos/ACT-8050/etiqueta/archivo", params={"formato": fmt})
            self.assertEqual(a.status_code, 200)
            self.assertIn(f'filename="ACT-8050.{ext}"', a.headers["content-disposition"])
        self.assertTrue(self.c.get("/api/activos/ACT-8050/etiqueta/archivo").content.startswith(b"\xef\xbb\xbf<?xml"))
        self.assertEqual(self.c.get("/api/activos/ACT-0404/etiqueta.svg").status_code, 404)
        self.assertEqual(self.c.get("/api/activos/ACT-8050/etiqueta/archivo", params={"formato": "pdf"}).status_code, 422)

    def test_etiqueta_usa_tqt_public_url(self):
        import os
        self.c.post("/api/activos", json={"numero": 8060, "nombre": "Cámara"})
        os.environ["TQT_PUBLIC_URL"] = "https://tqt.skyguardian.mx/"
        try:
            self.assertEqual(self.c.get("/api/activos/ACT-8060").json()["qr"], "https://tqt.skyguardian.mx/impresora?activo=ACT-8060")
            self.assertIn("https://tqt.skyguardian.mx/impresora?activo=ACT-8060", self.c.get("/api/activos/ACT-8060/etiqueta/xml").text)
        finally:
            del os.environ["TQT_PUBLIC_URL"]

    def test_impresion_sin_agente_503_y_con_agente_encola(self):
        self.c.post("/api/activos", json={"numero": 8070, "nombre": "Laptop"})
        with db.get_db() as conn:   # sin agentes en línea (otros tests pueden haber dejado uno): se fuerzan como desconectados
            dymo_remoto.init()
            conn.execute("UPDATE dymo_agentes SET last_seen=0")
            conn.execute("UPDATE dymo_trabajos SET estado='cancelada' WHERE estado IN ('pendiente','enviada')")   # cola limpia
        r = self.c.post("/api/activos/ACT-8070/imprimir")
        self.assertEqual(r.status_code, 503, r.text)
        self.assertTrue(r.json()["detail"])
        self.assertEqual(self.c.get("/api/activos/ACT-8070").json()["impresiones"], 0)
        # con agente en línea y la impresora lista
        ag = dymo_remoto.agente_crear("Raspberry activos")
        dymo_remoto.latido(ag["id"], {"impresora_ok": True})
        r = self.c.post("/api/activos/ACT-8070/imprimir")
        self.assertEqual(r.status_code, 200, r.text)
        d = r.json()
        self.assertEqual((d["ok"], d["via"]), (True, "compartida"))
        self.assertTrue(d["trabajo_id"])
        a = self.c.get("/api/activos/ACT-8070").json()
        self.assertEqual(a["impresiones"], 1)
        self.assertTrue(a["ultima_impresion"])
        self.assertEqual(a["eventos"][0]["evento"], "IMPRESION")
        # el agente recibe el SVG de la etiqueta del activo
        cab = {"Authorization": "Bearer " + ag["token"]}
        j = dymo_remoto.tomar(ag["id"])
        self.assertEqual((j["id"], j["tipo"]), (d["trabajo_id"], "activo"))
        s = self.c.get(f"/api/dymo/remoto/agente/trabajos/{j['id']}/svg", headers=cab)
        self.assertEqual(s.status_code, 200, s.text)
        self.assertEqual(re.findall(r"<text[^>]*>([^<]*)</text>", s.text), ["ACT", "8070"])
        self.assertNotIn('rx="2.5"', s.text)
        # impresión local con DYMO Connect
        self.assertEqual(self.c.post("/api/activos/ACT-8070/impreso").json(), {"ok": True})
        self.assertEqual(self.c.get("/api/activos/ACT-8070").json()["impresiones"], 2)
        self.assertEqual(self.c.post("/api/activos/ACT-0404/impreso").status_code, 404)

    def test_etiqueta_cabe_en_la_zona_segura(self):
        from app.services.dymo_service import DymoService, MARGEN_SEGURO_MM
        url = "https://tqt.skyguardian.mx/impresora?activo=ACT-0001"
        for fmt in ("30334", "30252"):
            d, lineas = DymoService.diseno_activo(url, "ACT 0001", fmt)
            f = d.formato
            self.assertEqual([t for t, _, _ in lineas], ["ACT", "0001"])
            self.assertAlmostEqual(lineas[0][1].x, lineas[1][1].x)   # alineadas a la izquierda, junto al QR
            self.assertGreater(lineas[1][2], lineas[0][2])
            self.assertGreaterEqual(lineas[0][1].x, d.qr.derecha)
            for caja in (d.qr, lineas[0][1], lineas[1][1]):
                self.assertGreaterEqual(caja.x, MARGEN_SEGURO_MM - 1e-6)
                self.assertGreaterEqual(caja.y, MARGEN_SEGURO_MM - 1e-6)
                self.assertLessEqual(caja.derecha, f.ancho_mm - MARGEN_SEGURO_MM + 1e-6)
                self.assertLessEqual(caja.abajo, f.alto_mm - MARGEN_SEGURO_MM + 1e-6)
            self.assertGreaterEqual(d.puntos_por_modulo, 4)
            bloque_centro = (lineas[0][1].y + lineas[1][1].abajo) / 2
            self.assertAlmostEqual(bloque_centro, d.qr.y + d.qr.h / 2, places=2)   # bloque centrado en vertical contra el QR
        d, lineas = DymoService.diseno_activo(url, "ACT 0001", "30334")
        self.assertEqual(d.puntos_por_modulo, 7)   # el QR no cambió respecto a v1.4.16 (más grande que la etiqueta de tarjeta, 6)
        self.assertGreater(d.qr.w, 21)
        self.assertGreaterEqual(lineas[1][2], 30)   # el número es grande
        # un número de 5 dígitos sigue cabiendo
        largo = DymoService.diseno_activo(url, "ACT 12345", "30334")[1]
        self.assertLessEqual(max(c.derecha for _, c, _ in largo), 57 - MARGEN_SEGURO_MM + 1e-6)

    def test_numero_de(self):
        self.assertEqual([svc.numero_de(t) for t in ("ACT-0001", "ACT 0001", "act0001", "1", "ACT-12345")], [1, 1, 1, 1, 12345])
        for malo in ("", "ACT", "ACT-0", "ACT-1234567", "abc", None):
            with self.assertRaises(ValueError):
                svc.numero_de(malo)


class TestLotes(unittest.TestCase):
    """v1.4.16: filtros de lote, impresión por filtro y avance de los trabajos en la cola compartida."""

    @classmethod
    def setUpClass(cls):
        verificar_aislamiento()
        db.init_db()
        usuarios.sembrar(iteraciones=1000)
        c = entrar()
        for n, alta in ((9101, "2031-03-10T08:00:00"), (9102, "2031-03-11T23:59:00"), (9103, "2031-03-12T00:01:00")):
            c.post("/api/activos", json={"numero": n, "nombre": f"Lote {n}", "departamento": "Lotes"})
            with db.get_db() as conn:
                conn.execute("UPDATE activos SET creado_en=? WHERE numero=?", (alta, n))
        c.post("/api/activos/ACT-9102/impreso")   # 9102 ya salió una vez

    def setUp(self):
        usuarios.limitador.reiniciar()
        self.c = entrar()
        with db.get_db() as conn:   # estado de impresión conocido, sin depender del orden de los tests
            conn.execute("UPDATE activos SET impresiones=0 WHERE numero IN (9101,9103)")
            conn.execute("UPDATE activos SET impresiones=1 WHERE numero=9102")

    def nums(self, **params):
        r = self.c.get("/api/activos", params=params)
        self.assertEqual(r.status_code, 200, r.text)
        return [x["numero"] for x in r.json()["items"]]

    def agente_en_linea(self):
        with db.get_db() as conn:
            dymo_remoto.init()
            conn.execute("UPDATE dymo_trabajos SET estado='cancelada' WHERE estado IN ('pendiente','enviada')")
        ag = dymo_remoto.agente_crear("Raspberry lotes")
        dymo_remoto.latido(ag["id"], {"impresora_ok": True})
        return ag

    def test_filtros_desde_hasta_impresa_y_codigos(self):
        self.assertEqual(self.nums(desde="2031-03-10"), [9101])   # solo desde = ese día
        self.assertEqual(self.nums(desde="2031-03-10", hasta="2031-03-11"), [9102, 9101])
        self.assertEqual(self.nums(desde="2031-03-11", hasta="2031-03-12"), [9103, 9102])
        self.assertEqual(self.nums(hasta="2031-03-11", departamento="Lotes"), [9102, 9101])
        self.assertEqual(self.nums(desde="2031-03-10", hasta="2031-03-12", impresa=0), [9103, 9101])
        self.assertEqual(self.nums(desde="2031-03-10", hasta="2031-03-12", impresa=1), [9102])
        self.assertEqual(self.nums(codigos="ACT-9101, 9103,zzz"), [9103, 9101])
        self.assertEqual(self.nums(codigos="zzz"), [])
        self.assertIn("creado_en", self.c.get("/api/activos", params={"codigos": "9101"}).json()["items"][0])
        for malo in ({"desde": "10/03/2031"}, {"desde": "2031-03-12", "hasta": "2031-03-10"}, {"impresa": 2}):
            self.assertEqual(self.c.get("/api/activos", params=malo).status_code, 422, malo)

    def test_lote_por_filtro_y_estado_de_trabajos(self):
        ag = self.agente_en_linea()
        r = self.c.post("/api/activos/imprimir-lote", json={"desde": "2031-03-10", "hasta": "2031-03-12", "impresa": 0})
        self.assertEqual(r.status_code, 200, r.text)
        d = r.json()
        self.assertEqual((d["ok"], d["encolados"], d["errores"]), (True, 2, []))
        self.assertEqual([t["codigo"] for t in d["trabajos"]], ["ACT-9101", "ACT-9103"])
        ids = [t["trabajo_id"] for t in d["trabajos"]]
        self.assertEqual(self.nums(desde="2031-03-10", hasta="2031-03-12", impresa=0), [])   # ya cuentan como impresas
        # misma cola que las etiquetas de tarjeta: aparecen en el estado compartido
        self.assertGreaterEqual(self.c.get("/api/activos/impresora/estado").json()["en_cola"], 2)
        e = self.c.get("/api/activos/lote/estado", params={"trabajos": ",".join(ids)}).json()
        self.assertEqual([t["estado"] for t in e["trabajos"]], ["pendiente", "pendiente"])
        self.assertEqual((e["resumen"]["pendiente"], e["resumen"]["total"], e["terminado"]), (2, 2, False))
        self.assertEqual([t["delante"] for t in e["trabajos"]], [0, 1])
        j = dymo_remoto.tomar(ag["id"])   # la Raspberry toma el primero
        self.assertEqual(j["id"], ids[0])
        e = self.c.get("/api/activos/lote/estado", params={"trabajos": ",".join(ids)}).json()
        self.assertEqual([t["estado"] for t in e["trabajos"]], ["imprimiendo", "pendiente"])
        dymo_remoto.resultado(ag["id"], ids[0], "impresa", "", "")
        j = dymo_remoto.tomar(ag["id"])
        dymo_remoto.resultado(ag["id"], j["id"], "error", "Sin papel", "")
        e = self.c.get("/api/activos/lote/estado", params={"trabajos": ",".join(ids) + ",noexiste"}).json()
        self.assertEqual([t["estado"] for t in e["trabajos"]], ["impresa", "error", "error"])
        self.assertEqual(e["trabajos"][1]["detalle"], "Sin papel")
        self.assertEqual((e["resumen"]["impresa"], e["resumen"]["error"], e["terminado"]), (1, 2, True))
        self.assertEqual(self.c.get("/api/activos/lote/estado", params={"trabajos": ""}).status_code, 422)
        self.assertEqual(self.c.get("/api/activos/lote/estado", params={"trabajos": ",".join(["a"] * 501)}).status_code, 422)

    def test_lote_por_codigos_con_inexistentes_y_validaciones(self):
        self.agente_en_linea()
        d = self.c.post("/api/activos/imprimir-lote", json={"codigos": ["ACT-9101", "ACT-9999"]}).json()
        self.assertEqual((d["ok"], d["encolados"], d["errores"][0]["codigo"]), (False, 1, "ACT-9999"))
        self.assertEqual(self.c.post("/api/activos/imprimir-lote", json={}).status_code, 422)   # sin criterio
        self.assertEqual(self.c.post("/api/activos/imprimir-lote", json={"desde": "2031-01-01", "hasta": "2031-01-02"}).status_code, 404)
        self.assertEqual(self.c.post("/api/activos/imprimir-lote", json={"desde": "mañana"}).status_code, 422)

    def test_limite_de_500_por_lote(self):
        self.agente_en_linea()
        r = self.c.post("/api/activos/imprimir-lote", json={"codigos": [f"ACT-{n:04d}" for n in range(1, 502)]})
        self.assertEqual(r.status_code, 422)
        from app.routers import activos as router_activos
        before = self.c.get("/api/activos/ACT-9103").json()["impresiones"]
        with mock.patch.object(router_activos, "LOTE_MAX", 2):   # el filtro que resuelve más del máximo se rechaza completo
            r = self.c.post("/api/activos/imprimir-lote", json={"departamento": "Lotes"})
        self.assertEqual(r.status_code, 422, r.text)
        self.assertIn("máximo", r.json()["detail"])
        self.assertEqual(self.c.get("/api/activos/ACT-9103").json()["impresiones"], before)   # nada se encoló

    def test_sin_dymo_el_lote_da_503(self):
        with db.get_db() as conn:
            dymo_remoto.init()
            conn.execute("UPDATE dymo_agentes SET last_seen=0")
        r = self.c.post("/api/activos/imprimir-lote", json={"codigos": ["ACT-9101"]})
        self.assertEqual(r.status_code, 503, r.text)
        self.assertFalse(self.c.get("/api/activos/impresora/estado").json()["disponible"])


class TestExportarYEliminar(unittest.TestCase):
    """v1.4.18: exportar a Excel con los mismos filtros del listado y eliminar en lote con contraseña."""

    @classmethod
    def setUpClass(cls):
        verificar_aislamiento()
        db.init_db()
        usuarios.sembrar(iteraciones=1000)
        c = entrar()
        c.post("/api/activos", json={"numero": 9201, "nombre": "Laptop export", "departamento": "Export", "categoria": "Cómputo",
                                     "valor": 1234.5, "fecha_adquisicion": "2024-05-06", "serie": "SER-1", "marca": "Dell",
                                     "caracteristicas": [{"etiqueta": "RAM", "valor": "16 GB"}, {"etiqueta": "Procesador", "valor": "i5"}]})
        c.post("/api/activos", json={"numero": 9202, "nombre": "=1+1", "departamento": "Export", "estado": "resguardo",
                                     "caracteristicas": [{"etiqueta": "Zeta", "valor": "z"}]})
        c.post("/api/activos", json={"numero": 9203, "nombre": "Otro depto", "departamento": "Otro"})
        for n, alta in ((9201, "2031-04-01T09:00:00"), (9202, "2031-04-02T09:00:00"), (9203, "2031-04-03T09:00:00")):
            with db.get_db() as conn:
                conn.execute("UPDATE activos SET creado_en=? WHERE numero=?", (alta, n))

    def setUp(self):
        usuarios.limitador.reiniciar()
        self.c = entrar()

    def libro(self, **params):
        r = self.c.get("/api/activos/exportar.xlsx", params=params)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertIn("spreadsheetml", r.headers["content-type"])
        self.assertRegex(r.headers["content-disposition"], r'filename="activos_\d{8}-\d{4}\.xlsx"')
        from openpyxl import load_workbook
        return load_workbook(io.BytesIO(r.content))["Activos"]

    def test_exportar_columnas_filas_y_formato(self):
        ws = self.libro(departamento="Export")
        cab = [c.value for c in ws[1]]
        self.assertEqual(cab, ["Código", "Nombre", "Estado", "Departamento", "Categoría", "Ubicación", "Responsable", "Serie", "Marca",
                               "Modelo", "Valor", "Fecha de adquisición", "Alta", "Impresiones", "Última impresión", "Notas",
                               "Procesador", "RAM", "Zeta"])   # + unión alfabética de las etiquetas de características
        self.assertEqual(ws.max_row, 3)
        a, b = [[c.value for c in ws[i]] for i in (2, 3)]
        self.assertEqual((a[0], a[1], a[2], a[3], a[4], a[7], a[8], a[10]), ("ACT-9201", "Laptop export", "Activo", "Export", "Cómputo", "SER-1", "Dell", 1234.5))
        self.assertEqual((a[16], a[17], a[18]), ("i5", "16 GB", None))   # celda vacía
        self.assertEqual((b[0], b[2], b[18]), ("ACT-9202", "En resguardo", "z"))
        self.assertEqual(a[11].date().isoformat(), "2024-05-06")
        self.assertEqual(a[12].isoformat(timespec="minutes"), "2031-04-01T09:00")
        self.assertEqual((a[13], a[14]), (0, None))
        self.assertIn("$", ws.cell(2, 11).number_format)
        self.assertEqual(ws.cell(2, 12).number_format, "yyyy-mm-dd")
        self.assertTrue(all(c.font.bold for c in ws[1]))
        self.assertEqual((ws.freeze_panes, ws.auto_filter.ref), ("C2", ws.dimensions))
        self.assertGreater(ws.column_dimensions["B"].width, 20)
        self.assertEqual(ws.cell(3, 2).data_type, "s")   # «=1+1» queda como texto, no como fórmula
        self.assertEqual(ws.cell(3, 2).value, "=1+1")

    def test_exportar_aplica_los_mismos_filtros_que_el_listado(self):
        def codigos(ws):
            return [ws.cell(i, 1).value for i in range(2, ws.max_row + 1)]
        casos = ({"q": "laptop"}, {"estado": "resguardo", "departamento": "Export"}, {"desde": "2031-04-02", "hasta": "2031-04-03"},
                 {"desde": "2031-04-01"}, {"codigos": "9203, ACT-9201"}, {"impresa": 0, "departamento": "Export"}, {"categoria": "Cómputo"})
        for params in casos:
            esperado = sorted(f"ACT-{x['numero']:04d}" for x in self.c.get("/api/activos", params=params).json()["items"])
            self.assertEqual(codigos(self.libro(**params)), esperado, params)
        self.assertEqual(codigos(self.libro(desde="2031-04-02", hasta="2031-04-03")), ["ACT-9202", "ACT-9203"])
        self.assertEqual(codigos(self.libro(codigos="9203, ACT-9201")), ["ACT-9201", "ACT-9203"])
        self.assertEqual(self.libro(q="no-existe-xyz").max_row, 1)   # solo encabezados
        self.assertEqual(self.c.get("/api/activos/exportar.xlsx", params={"desde": "ayer"}).status_code, 422)
        self.assertEqual(cliente_sin_sesion(app).get("/api/activos/exportar.xlsx").status_code, 401)

    def crear(self, *numeros):
        for n in numeros:
            self.c.post("/api/activos", json={"numero": n, "nombre": f"Borrar {n}"})

    def eliminar(self, codigos, password):
        return self.c.post("/api/activos/eliminar-lote", json={"codigos": codigos, "password": password})

    def test_eliminar_con_contrasena_incorrecta_no_borra(self):
        self.crear(9211)
        for malo in ("mala", ""):
            r = self.eliminar(["ACT-9211"], malo)
            self.assertEqual((r.status_code, r.json()["detail"]), (401, "Contraseña incorrecta."))
        self.assertEqual(self.c.get("/api/activos/ACT-9211").status_code, 200)

    def test_eliminar_con_contrasena_correcta(self):
        self.crear(9212, 9213)
        self.c.patch("/api/activos/ACT-9212", json={"ubicacion": "X"})   # genera más eventos
        r = self.eliminar(["ACT-9212", "9213", "ACT-9999", "ACT-9212", "basura"], usuarios.CLAVE_INICIAL)
        self.assertEqual(r.status_code, 200, r.text)
        d = r.json()
        self.assertEqual((d["eliminados"], d["codigos"], d["no_encontrados"]), (2, ["ACT-9212", "ACT-9213"], ["ACT-9999", "basura"]))
        self.assertEqual(self.c.get("/api/activos/ACT-9212").status_code, 404)
        self.assertEqual(self.c.get("/api/activos/ACT-9213").status_code, 404)
        with db.get_db() as conn:
            huerfanos = conn.execute("SELECT COUNT(*) n FROM activos_eventos WHERE activo_id NOT IN (SELECT id FROM activos)").fetchone()["n"]
            bit = conn.execute("SELECT detalle, operador FROM escaneos WHERE evento='ACTIVOS_ELIMINADOS' ORDER BY id DESC LIMIT 1").fetchone()
        self.assertEqual(huerfanos, 0)
        self.assertIn("ACT-9212, ACT-9213", bit["detalle"])
        self.assertEqual(bit["operador"], "developer@skyguardian.mx")

    def test_eliminar_lista_vacia_o_demasiado_larga_da_422(self):
        self.assertEqual(self.eliminar([], usuarios.CLAVE_INICIAL).status_code, 422)
        self.assertEqual(self.eliminar([f"ACT-{n:04d}" for n in range(1, 502)], usuarios.CLAVE_INICIAL).status_code, 422)

    def test_eliminar_se_bloquea_tras_muchos_fallos(self):
        self.crear(9214)
        codigos = [self.eliminar(["ACT-9214"], "mala").status_code for _ in range(12)]
        self.assertEqual(codigos[0], 401)
        self.assertIn(429, codigos)
        r = self.eliminar(["ACT-9214"], usuarios.CLAVE_INICIAL)   # bloqueado: ni la correcta pasa hasta que expire
        self.assertEqual(r.status_code, 429)
        self.assertIn("Retry-After", r.headers)
        self.assertEqual(self.c.get("/api/activos/ACT-9214").status_code, 200)
        usuarios.limitador.reiniciar()
        self.assertEqual(self.eliminar(["ACT-9214"], usuarios.CLAVE_INICIAL).status_code, 200)


class TestAccesoImpresora(unittest.TestCase):
    EMAIL = "developer4@skyguardian.mx"

    @classmethod
    def setUpClass(cls):
        verificar_aislamiento()
        db.init_db()
        usuarios.sembrar(iteraciones=1000)

    def setUp(self):
        usuarios.limitador.reiniciar()
        self.addCleanup(usuarios.actualizar_cuenta, self.EMAIL, "administrador", None, None, ["escaner-tqt"])

    def test_consultor_lee_pero_no_escribe(self):
        entrar().post("/api/activos", json={"numero": 8080, "nombre": "Taladro"})
        usuarios.actualizar_cuenta(self.EMAIL, rol="consultor", modulos=["impresora"])
        c = entrar(self.EMAIL)
        self.assertEqual(c.get("/api/activos").status_code, 200)
        self.assertEqual(c.get("/api/activos/ACT-8080").status_code, 200)
        self.assertEqual(c.get("/api/activos/exportar.xlsx").status_code, 200)       # exportar es solo lectura
        self.assertEqual(c.post("/api/activos/eliminar-lote", json={"codigos": ["ACT-8080"], "password": usuarios.CLAVE_INICIAL}).status_code, 403)
        self.assertEqual(c.get("/api/activos/ACT-8080").status_code, 200)            # y no se borró nada
        self.assertEqual(c.get("/api/activos/impresora/estado").status_code, 200)   # el estado de la DYMO sí está en el módulo
        self.assertEqual(c.get("/api/dymo/remoto/estado").status_code, 403)         # el de Escáner TQT no
        self.assertEqual(c.post("/api/activos", json={"nombre": "No"}).status_code, 403)
        self.assertEqual(c.patch("/api/activos/ACT-8080", json={"nombre": "No"}).status_code, 403)
        self.assertEqual(c.post("/api/activos/ACT-8080/imprimir").status_code, 403)

    def test_general_sin_modulo_impresora_403_y_con_modulo_entra(self):
        usuarios.actualizar_cuenta(self.EMAIL, rol="general", modulos=["escaner-tqt"])
        c = entrar(self.EMAIL)
        self.assertEqual(c.get("/api/activos").status_code, 403)
        self.assertEqual(c.post("/api/activos", json={"nombre": "No"}).status_code, 403)
        usuarios.actualizar_cuenta(self.EMAIL, rol="general", modulos=["escaner-tqt", "impresora"])
        c = entrar(self.EMAIL)
        self.assertEqual(c.get("/api/activos").status_code, 200)
        r = c.post("/api/activos", json={"nombre": "Prueba general"})
        self.assertEqual(r.status_code, 201, r.text)
        self.assertEqual(r.json()["creado_por"], self.EMAIL)


if __name__ == "__main__":
    unittest.main()
