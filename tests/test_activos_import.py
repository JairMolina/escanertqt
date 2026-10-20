"""v1.4.15: importación de activos desde Excel (libro generado en memoria; sin datos reales)."""
import _aislamiento  # noqa: F401  (debe ser lo primero)
import io
import re
import sqlite3
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from app.database import db
from app.main import app
from app.services import activos as svc
from app.services import activos_import as ai
from app.services import dymo_remoto, usuarios

from _aislamiento import cliente_sin_sesion, verificar_aislamiento

VERDE = PatternFill("solid", fgColor="FFE2EFD9")
ESTANDAR = ["Nombre Activo", "Status", "Departamento", "Activo Number", "Fecha de revisión", "Valor del Activo", "Asignado a",
            "Última modificación hecha por", "Ubicación del Activo", "Número de Serie", "Fotografía del item", "Descripción del item",
            "Colaborador", "Comentarios Estado Físico", "Notas", "Fecha de Baja", "Razón de Baja", "Motivo de Baja",
            "Comentarios de Baja", "Ultima Modificación", "Origen", "Fecha de creación", "Procesador", "Memoria RAM", "Disco Duro",
            "Tamaño Pantalla (Pulgadas)"]


def fila_std(nombre, num, status=None, depto="Administración", revision=44861.0, valor=1000.0, asignado=None, ubicacion=None,
             serie=None, desc=None, colab=None, fisico=None, notas=None, baja=None, ultima=None, creacion=None, cpu=None):
    f = [None] * 26
    f[0], f[1], f[2], f[3], f[4], f[5], f[6] = nombre, status, depto, num, revision, valor, asignado
    f[8], f[9], f[11], f[12], f[13], f[14] = ubicacion, serie, desc, colab, fisico, notas
    f[15], f[16], f[17] = baja or (None, None, None)
    f[19], f[21], f[22] = ultima, creacion, cpu
    return f


def libro() -> bytes:
    wb = Workbook()
    std = wb.active
    std.title = "STD"
    std.append(ESTANDAR)
    std.append(fila_std("LAPTOP DELL", "ACT501", "Activo", valor=15000.0, asignado="Employees::::Ana", ubicacion=108.0, serie="S/N",
                        desc="Laptop 14 pulgadas", colab="vtcmcolaboradores::::LUIS", ultima="27-10-2022 04:20 p.m.",
                        creacion="16-11-2018 06:27 p.m.", cpu="Intel Core i5"))
    std.append(fila_std("SILLA NO REVISADA", "ACT502"))
    std.append(fila_std("FILA ROTA", "XX"))
    std.append(fila_std("MESA SIN NUMERO", None, asignado="x"))
    for r in (2, 4, 5):
        std.cell(r, 1).fill = VERDE
    sky = wb.create_sheet("SKY")   # encabezados 1-10 malos, datos en formato estándar
    sky.append(["NOMBRE DEL ACTIVO", "ESTATUS", "DEPARTAMENTO", "NUMERO DE ACTIVO", "FECHA DE REVISION", "VALOR DEL ACTIVO",
                "ASIGNADO A ", "UBICACION", "NUMERO DE SERIE", "COMENTARIOS", None] + ESTANDAR[11:])
    sky.append(fila_std("SILLA EJECUTIVA", "ACT503", "Baja - Perdida", depto="Ventas", serie="SN-77", colab="Pedro",
                        baja=(46015.0, "Se perdió", "Obsoleto"), ultima="01-01-2026 10:00 a.m."))
    sky.append(fila_std("LAPTOP DELL 2", "ACT501", "Activo", ultima="02-03-2026 09:00 a.m.", colab="Rosa"))   # duplicado más reciente
    for r in (2, 3):
        sky.cell(r, 1).fill = VERDE
    vipe = wb.create_sheet("VIPE")   # otro orden y mojibake
    vipe.append(["Nombre Activo", "Status", "Departamento", "Activo Number", "Valor del Activo", "Asignado a", "Ubicacion del Activo",
                 "Numero de Serie", "Descripcion del articulo ", "Colaborador", "Comentarios Estado F�\xadsico",
                 "Fecha de creaci�\xb3n", "Tama�\xb1o Pantalla (Pulgadas)"])
    vipe.append(["MONITOR LG", "Resguardo", "Vicepresidencia", "ACT504", "$2,500.50", "Ana", "Sala", "N/A", "Monitor 24\"", "MarÃ_x00AD_a",
                 "Revisi�\xb3n pendiente", 44861.0, "24"])
    vipe.cell(2, 1).fill = VERDE
    return _bytes(wb)


def _bytes(wb) -> bytes:
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


class TestLimpieza(unittest.TestCase):
    def test_departamentos_con_una_sola_escritura(self):
        items = [{"departamento": d} for d in ("Administracion", "Administración", "Administración", "OPERACIONES", "Operaciones", "")]
        ai._unificar_departamentos(items)
        self.assertEqual([i["departamento"] for i in items],
                         ["Administración", "Administración", "Administración", "Operaciones", "Operaciones", ""])

    def test_reparar_mojibake(self):
        self.assertEqual(ai.reparar("F�\xadsico"), "Físico")
        self.assertEqual(ai.reparar("Dise�_x00B1_o"), "Diseño")
        self.assertEqual(ai.reparar("GARCÃ_x008d_A"), "GARCÍA")
        self.assertEqual(ai.reparar("TamaÃ±o Ãšltima"), "Tamaño Última")
        self.assertEqual(ai.reparar("AﾑO"), "AÑO")

    def test_texto_vacios_y_prefijos(self):
        for v in ("S/N", "sn", "N/A", "undefined", ".", None, "  ", "Solo electronicos campos en adelante"):
            self.assertEqual(ai.texto(v), "", v)
        self.assertEqual(ai.texto("Employees::::Ana Pérez"), "Ana Pérez")
        self.assertEqual(ai.texto("vtcmcolaboradores::::LUIS"), "LUIS")
        self.assertEqual(ai.texto(108.0), "108")

    def test_fechas(self):
        self.assertEqual(ai.fecha_hora(44861.0).date().isoformat(), "2022-10-27")
        d = ai.fecha_hora("16-11-2018 06:27 p.m.")
        self.assertEqual((d.year, d.month, d.day, d.hour, d.minute), (2018, 11, 16, 18, 27))
        self.assertEqual(ai.fecha_hora("12-01-2020 12:05 a.m.").hour, 0)
        self.assertIsNone(ai.fecha_hora("mañana"))

    def test_estado_categoria_valor(self):
        self.assertEqual([ai.estado_de(s, False) for s in ("", "Activo", "Baja - Obsoleto", "Resguardo", "Reparación")],
                         ["activo", "activo", "baja", "resguardo", "reparacion"])
        self.assertEqual(ai.estado_de("", True), "baja")
        self.assertEqual([ai.categoria_de(n) for n in ("LAPTOP ACER", "MONITOR PARA COMPUTADORA", "Silla ejecutiva", "CAMARA IP", "MOUSE", "AIRE ACONDICIONADO",
                                                       "DIADEMA", "Alexa Echo")],
                         ["Cómputo", "Monitores", "Mobiliario", "Seguridad", "Periféricos", "Equipo de oficina", "Telefonía", "General"])
        self.assertEqual((ai.valor_num("$2,500.50"), ai.valor_num("x"), ai.valor_num(-3)), (2500.5, None, None))


class TestAnalisis(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        verificar_aislamiento()
        db.init_db()

    def test_analizar_solo_verdes(self):
        r = ai.analizar(libro(), existentes=set())
        h = {x["hoja"]: x for x in r["hojas"]}
        self.assertEqual((h["STD"]["formato"], h["SKY"]["formato"], h["VIPE"]["formato"]), ("estandar", "estandar", "por_nombre"))
        self.assertEqual([h[k]["filas_verdes"] for k in ("STD", "SKY", "VIPE")], [3, 2, 1])
        self.assertEqual([i["numero"] for i in r["items"]], [501, 503, 504])   # ACT502 no es verde; XX se omite; ACT501 repetido
        self.assertEqual(r["totales"], {"nuevos": 3, "actualizar": 0, "omitidos": 3})
        self.assertTrue(any("inválido" in o["motivo"] for o in h["STD"]["omitidos"]))
        # duplicado: gana la fila con «Ultima Modificación» más reciente (la de SKY) y se avisa
        self.assertEqual(r["avisos"], [{"codigo": "ACT-0501", "gana": "SKY!3", "descartada": "STD!2"}])
        it = {i["numero"]: i for i in r["items"]}
        self.assertEqual((it[501]["hoja"], it[501]["nombre"], it[501]["responsable"]), ("SKY", "LAPTOP DELL 2", "Rosa"))
        b = it[503]
        self.assertEqual((b["estado"], b["departamento"], b["serie"], b["categoria"]), ("baja", "Ventas", "SN-77", "Mobiliario"))
        self.assertIn(("Fecha de baja", "2025-12-24"), [(c["etiqueta"], c["valor"]) for c in b["caracteristicas"]])
        self.assertIn(("Razón de baja", "Se perdió"), [(c["etiqueta"], c["valor"]) for c in b["caracteristicas"]])
        v = it[504]
        self.assertEqual((v["estado"], v["valor"], v["serie"], v["responsable"], v["categoria"], v["fecha_adquisicion"]),
                         ("resguardo", 2500.5, "", "María", "Monitores", "2022-10-27"))
        car = {c["etiqueta"]: c["valor"] for c in v["caracteristicas"]}
        self.assertEqual((car["Comentarios de estado físico"], car["Tamaño de pantalla"], car["Hoja de origen"]), ("Revisión pendiente", "24", "VIPE"))

    def test_analizar_sin_filtro_de_verde_y_caracteristicas_estandar(self):
        r = ai.analizar(libro(), solo_verdes=False, existentes={502})
        it = {i["numero"]: i for i in r["items"]}
        self.assertIn(502, it)
        self.assertEqual(it[502]["accion"], "actualizar")

    def test_item_estandar_completo(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "STD"
        ws.append(ESTANDAR)
        ws.append(fila_std("LAPTOP DELL", "ACT601", "Activo", valor=15000.0, asignado="Employees::::Ana", ubicacion=108.0, serie="S/N",
                           desc="Laptop 14", colab="vtcmcolaboradores::::LUIS", creacion="16-11-2018 06:27 p.m.", cpu="Intel Core i5"))
        ws.cell(2, 1).fill = VERDE
        i = ai.analizar(_bytes(wb), existentes=set())["items"][0]
        self.assertEqual((i["ubicacion"], i["serie"], i["responsable"], i["valor"], i["fecha_adquisicion"], i["notas"], i["categoria"]),
                         ("108", "", "LUIS", 15000.0, "2018-11-16", "Laptop 14", "Cómputo"))
        self.assertEqual([c["etiqueta"] for c in i["caracteristicas"]],
                         ["Asignado a", "Colaborador", "Descripción", "Procesador", "Fecha de revisión", "Hoja de origen"])
        self.assertEqual(i["caracteristicas"][0]["valor"], "Ana")

    def test_archivo_invalido(self):
        with self.assertRaises(ValueError):
            ai.analizar(b"no soy un xlsx")

    def test_hoja_sin_columnas_conocidas(self):
        wb = Workbook()
        wb.active.title = "RARA"
        wb.active.append(["a", "b"])
        h = ai.analizar(_bytes(wb), existentes=set())["hojas"][0]
        self.assertEqual(h["formato"], "desconocido")


class TestAplicarYApi(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        verificar_aislamiento()
        db.init_db()
        usuarios.sembrar(iteraciones=1000)

    def setUp(self):
        usuarios.limitador.reiniciar()
        self.c = cliente_sin_sesion(app)
        r = self.c.post("/api/auth/login", json={"email": "developer@skyguardian.mx", "password": usuarios.CLAVE_INICIAL})
        self.assertEqual(r.status_code, 200, r.text)

    def importar(self, datos=None, aplicar=0, nombre="inventario.xlsx", cliente=None):
        return (cliente or self.c).post("/api/activos/importar", files={"archivo": (nombre, datos or libro(), "application/octet-stream")},
                                        data={"solo_verdes": "1", "aplicar": str(aplicar)})

    def test_aplicar_crea_actualiza_y_no_borra_con_vacios(self):
        db.init_db()
        # el activo 504 ya existe con datos propios: la importación no debe borrar lo que el Excel trae vacío
        self.c.post("/api/activos", json={"numero": 504, "nombre": "Monitor viejo", "serie": "SERIE-PROPIA", "ubicacion": "Taller",
                                          "caracteristicas": [{"etiqueta": "Garantía", "valor": "2 años"}]})
        vista = self.importar()
        self.assertEqual(vista.status_code, 200, vista.text)
        d = vista.json()
        self.assertFalse(d["aplicado"])
        self.assertEqual(d["totales"], {"nuevos": 2, "actualizar": 1, "omitidos": 3})
        self.assertEqual(len(d["muestra"]), 3)
        self.assertEqual(self.c.get("/api/activos/ACT-0501").status_code, 404)   # la vista previa no escribe
        r = self.importar(aplicar=1)
        d = r.json()
        self.assertTrue(d["aplicado"])
        self.assertEqual(d["codigos"], ["ACT-0501", "ACT-0503", "ACT-0504"])
        self.assertEqual((d["resultado"]["creados"], d["resultado"]["actualizados"]), (2, 1))
        a = self.c.get("/api/activos/ACT-0501").json()
        self.assertEqual((a["nombre"], a["departamento"], a["categoria"], a["estado"]), ("LAPTOP DELL 2", "Administración", "Cómputo", "activo"))
        self.assertEqual(a["eventos"][0]["evento"], "IMPORTACION")
        self.assertIn("hoja SKY, fila 3", a["eventos"][0]["detalle"])
        self.assertEqual(self.c.get("/api/activos/ACT-0503").json()["estado"], "baja")
        m = self.c.get("/api/activos/ACT-0504").json()
        self.assertEqual((m["nombre"], m["serie"], m["ubicacion"], m["estado"], m["valor"]), ("MONITOR LG", "SERIE-PROPIA", "Sala", "resguardo", 2500.5))
        etiquetas = {c["etiqueta"]: c["valor"] for c in m["caracteristicas"]}
        self.assertEqual((etiquetas["Garantía"], etiquetas["Hoja de origen"]), ("2 años", "VIPE"))
        self.assertEqual(m["eventos"][0]["evento"], "IMPORTACION")
        # segunda pasada: todo existe y no cambia nada
        d2 = self.importar(aplicar=1).json()
        self.assertEqual((d2["totales"]["nuevos"], d2["totales"]["actualizar"]), (0, 3))
        self.assertEqual((d2["resultado"]["creados"], d2["resultado"]["actualizados"], d2["resultado"]["sin_cambios"]), (0, 0, 3))
        self.assertEqual(len([e for e in self.c.get("/api/activos/ACT-0501").json()["eventos"] if e["evento"] == "IMPORTACION"]), 1)

    def test_validaciones_de_la_carga(self):
        self.assertEqual(self.importar(b"texto", nombre="x.xlsx").status_code, 422)
        self.assertEqual(self.importar(libro(), nombre="x.csv").status_code, 422)
        self.assertEqual(self.importar(b"PK\x03\x04roto").status_code, 422)
        self.assertEqual(self.c.post("/api/activos/importar", json={}).status_code, 422)
        self.assertEqual(self.c.post("/api/activos/importar", files={"otro": ("a.xlsx", libro())}).status_code, 422)

    def test_consultor_no_importa(self):
        email = "developer4@skyguardian.mx"
        self.addCleanup(usuarios.actualizar_cuenta, email, "administrador", None, None, ["escaner-tqt"])
        usuarios.actualizar_cuenta(email, rol="consultor", modulos=["impresora"])
        c = cliente_sin_sesion(app)
        self.assertEqual(c.post("/api/auth/login", json={"email": email, "password": usuarios.CLAVE_INICIAL}).status_code, 200)
        self.assertEqual(self.importar(cliente=c).status_code, 403)
        self.assertEqual(c.post("/api/activos/imprimir-lote", json={"codigos": ["ACT-0001"]}).status_code, 403)

    def test_zip_de_etiquetas_y_ultima_importacion(self):
        self.importar(aplicar=1)
        z = self.c.get("/api/activos/etiquetas.zip", params={"importacion": "ultima"})
        self.assertEqual(z.status_code, 200)
        self.assertEqual(z.headers["content-type"], "application/zip")
        nombres = zipfile.ZipFile(io.BytesIO(z.content)).namelist()
        self.assertEqual(nombres, ["ACT-0501.dymo", "ACT-0503.dymo", "ACT-0504.dymo"])
        p = self.c.post("/api/activos/etiquetas.zip", json={"codigos": ["ACT-0501", "ACT-0504", "ACT-9999"]})
        self.assertEqual((p.status_code, p.headers["content-type"]), (200, "application/zip"))
        self.assertEqual(sorted(zipfile.ZipFile(io.BytesIO(p.content)).namelist()), ["ACT-0501.dymo", "ACT-0504.dymo"])
        self.assertEqual(self.c.post("/api/activos/etiquetas.zip", json={"codigos": ["ACT-9999"]}).status_code, 404)
        self.assertEqual(self.c.post("/api/activos/etiquetas.zip", json={"codigos": []}).status_code, 422)
        z = self.c.get("/api/activos/etiquetas.zip", params={"codigos": "ACT-0501, act 0503,ACT-9999"})
        self.assertEqual(sorted(zipfile.ZipFile(io.BytesIO(z.content)).namelist()), ["ACT-0501.dymo", "ACT-0503.dymo"])
        self.assertEqual(self.c.get("/api/activos/etiquetas.zip").status_code, 404)
        self.assertEqual(self.c.get("/api/activos/etiquetas.zip", params={"codigos": "ACT-9999"}).status_code, 404)

    def test_imprimir_lote(self):
        self.c.post("/api/activos", json={"numero": 520, "nombre": "Uno"})
        self.c.post("/api/activos", json={"numero": 521, "nombre": "Dos"})
        with db.get_db() as conn:
            dymo_remoto.init()
            conn.execute("UPDATE dymo_agentes SET last_seen=0")
        r = self.c.post("/api/activos/imprimir-lote", json={"codigos": ["ACT-0520"]})
        self.assertEqual(r.status_code, 503, r.text)
        ag = dymo_remoto.agente_crear("Raspberry lote")
        dymo_remoto.latido(ag["id"], {"impresora_ok": True})
        r = self.c.post("/api/activos/imprimir-lote", json={"codigos": ["ACT-0520", "521", "ACT-9998"]})
        self.assertEqual(r.status_code, 200, r.text)
        d = r.json()
        self.assertEqual((d["encolados"], len(d["errores"]), d["errores"][0]["codigo"]), (2, 1, "ACT-9998"))
        self.assertEqual(self.c.get("/api/activos/ACT-0520").json()["impresiones"], 1)
        self.assertEqual(self.c.post("/api/activos/imprimir-lote", json={"codigos": []}).status_code, 422)

    def test_departamento_caracteristicas_busqueda_y_etiqueta(self):
        r = self.c.post("/api/activos", json={"numero": 530, "nombre": "Servidor", "departamento": "Sistemas", "serie": "XYZ-1",
                                              "caracteristicas": [{"etiqueta": "Procesador", "valor": "Xeon unico530"},
                                                                  {"etiqueta": "Vacía", "valor": ""}]})
        self.assertEqual(r.status_code, 201, r.text)
        a = r.json()
        self.assertEqual((a["departamento"], a["caracteristicas"]), ("Sistemas", [{"etiqueta": "Procesador", "valor": "Xeon unico530"}]))
        self.assertEqual([x["numero"] for x in self.c.get("/api/activos", params={"q": "unico530"}).json()["items"]], [530])
        self.assertEqual([x["numero"] for x in self.c.get("/api/activos", params={"q": "sistemas"}).json()["items"]][:1], [530])
        self.assertEqual([x["numero"] for x in self.c.get("/api/activos", params={"departamento": "Sistemas"}).json()["items"]], [530])
        self.assertEqual(self.c.get("/api/activos", params={"departamento": "Nada"}).json()["total"], 0)
        g = self.c.get("/api/activos/530").json()   # POST, GET y PATCH devuelven departamento y caracteristicas
        self.assertEqual((g["departamento"], g["caracteristicas"][0]["etiqueta"]), ("Sistemas", "Procesador"))
        self.assertIn("departamento", self.c.get("/api/activos").json()["items"][0])
        self.assertIn("Sistemas", self.c.get("/api/activos/catalogos").json()["departamentos"])
        p = self.c.patch("/api/activos/ACT-0530", json={"caracteristicas": [{"etiqueta": "RAM", "valor": "64 GB"}], "departamento": "I+D"})
        self.assertEqual(p.json()["caracteristicas"], [{"etiqueta": "RAM", "valor": "64 GB"}])
        self.assertIn("características actualizadas", p.json()["eventos"][0]["detalle"])
        self.assertEqual(self.c.patch("/api/activos/ACT-0530", json={"caracteristicas": []}).json()["caracteristicas"], [])
        self.assertEqual(self.c.post("/api/activos", json={"nombre": "x", "caracteristicas": [{"etiqueta": "", "valor": "a"}]}).status_code, 422)
        # v1.4.16: la etiqueta ya no lleva departamento ni serie, solo «ACT 0530» y el QR
        svg = self.c.get("/api/activos/ACT-0530/etiqueta.svg").text
        self.assertEqual(re.findall(r"<text[^>]*>([^<]*)</text>", svg), ["ACT", "0530"])
        self.assertNotIn("S/N", svg)
        self.assertNotIn("DETALLE", self.c.get("/api/activos/ACT-0530/etiqueta/xml").text)


class TestMigracion(unittest.TestCase):
    def test_agrega_columnas_a_una_tabla_de_v1_4_14(self):
        verificar_aislamiento()
        with tempfile.TemporaryDirectory() as d:
            ruta = Path(d) / "vieja.db"
            con = sqlite3.connect(ruta)
            con.execute("CREATE TABLE activos(id INTEGER PRIMARY KEY AUTOINCREMENT, numero INTEGER NOT NULL UNIQUE, nombre TEXT NOT NULL, "
                        "categoria TEXT NOT NULL DEFAULT '', marca TEXT NOT NULL DEFAULT '', modelo TEXT NOT NULL DEFAULT '', "
                        "serie TEXT NOT NULL DEFAULT '', ubicacion TEXT NOT NULL DEFAULT '', responsable TEXT NOT NULL DEFAULT '', "
                        "estado TEXT NOT NULL DEFAULT 'activo', fecha_adquisicion TEXT, valor REAL, notas TEXT NOT NULL DEFAULT '', "
                        "creado_por TEXT NOT NULL DEFAULT '', creado_en TEXT NOT NULL, actualizado_en TEXT NOT NULL, "
                        "impresiones INTEGER NOT NULL DEFAULT 0, ultima_impresion TEXT)")
            con.execute("INSERT INTO activos(numero,nombre,creado_en,actualizado_en) VALUES(1,'Viejo','x','x')")
            con.commit()
            con.close()
            original = db.get_db
            with mock.patch.object(db, "get_db", lambda p=None: original(ruta)):
                svc._asegurar()
                svc._asegurar()   # idempotente
            con = sqlite3.connect(ruta)
            cols = {r[1] for r in con.execute("PRAGMA table_info(activos)")}
            self.assertTrue({"departamento", "caracteristicas"} <= cols)
            self.assertEqual(con.execute("SELECT nombre,departamento FROM activos").fetchone(), ("Viejo", ""))
            con.close()


if __name__ == "__main__":
    unittest.main()
