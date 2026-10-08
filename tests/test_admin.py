"""Área de administración: contraseña (PBKDF2), token HMAC, bloqueo por intentos, borrados con respaldo y exportación fiel."""
import _aislamiento  # noqa: F401  (debe ser lo primero)
import hashlib
import os
import sqlite3
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

import openpyxl
from fastapi.testclient import TestClient

from app.config import settings
from app.database import admin_ops, db, inventario as inv
from app.main import app
from app.services import admin_auth
from app.services.excel_sync import MESES_ES, _totales, structure_counts

from _aislamiento import crear_par, recibir_ws, verificar_aislamiento

CLAVE = "clave-de-prueba-1"


def _formulas(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    out = {}
    for ws in wb:
        for fila in ws.iter_rows():
            for c in fila:
                v = c.value
                if hasattr(v, "text") or (isinstance(v, str) and v.startswith("=")):
                    out[f"{ws.title}!{c.coordinate}"] = v.text if hasattr(v, "text") else v
    return out


class AdminBase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        verificar_aislamiento()
        cls.tmp = tempfile.TemporaryDirectory()
        cls._db_previa = settings.DB_PATH
        settings.DB_PATH = Path(cls.tmp.name) / "admin.db"          # BD propia: el reset no toca a los demás tests
        cls._parches = [
            mock.patch.dict(os.environ, {"TQT_ADMIN_PASSWORD": CLAVE}),
            mock.patch.object(settings, "ADMIN_RETARDO_MS", 0),
            mock.patch.object(settings, "ADMIN_PBKDF2_ITER", 1000),
        ]
        for p in cls._parches:
            p.start()
        db.init_db()
        cls.client = TestClient(app)
        cls.client.__enter__()

    @classmethod
    def tearDownClass(cls):
        cls.client.__exit__(None, None, None)
        for p in reversed(cls._parches):
            p.stop()
        settings.DB_PATH = cls._db_previa
        cls.tmp.cleanup()

    def setUp(self):
        admin_auth.limitador.reiniciar()
        with db.transaction() as c:  # cada test parte de la clave inicial
            c.execute("DELETE FROM ajustes WHERE clave LIKE 'admin_%'")
        self.lote = db.get_active_lote()["id"]

    def token(self, clave=CLAVE):
        r = self.client.post("/api/admin/login", json={"password": clave})
        self.assertEqual(r.status_code, 200, r.text)
        return {"X-Admin-Token": r.json()["token"]}

    def respaldos(self):
        return sorted(p.name for p in settings.BACKUP_DIR.glob("tqt_*.db")) if settings.BACKUP_DIR.exists() else []

    def eventos(self, prefijo="ADMIN_"):
        with db.get_db() as c:
            return [r["evento"] for r in c.execute("SELECT evento FROM escaneos WHERE evento LIKE ? ORDER BY id", (prefijo + "%",))]


class TestSesion(AdminBase):
    def test_01_deshabilitado_sin_contrasena(self):
        with mock.patch.dict(os.environ, {"TQT_ADMIN_PASSWORD": ""}):
            e = self.client.get("/api/admin/estado").json()
            self.assertFalse(e["habilitado"])
            self.assertIn("TQT_ADMIN_PASSWORD", e["mensaje"])
            r = self.client.post("/api/admin/login", json={"password": "cualquiera"})
            self.assertEqual(r.status_code, 503)
            self.assertIn("TQT_ADMIN_PASSWORD", r.json()["detail"])
            self.assertEqual(self.client.get("/api/admin/resumen").status_code, 503)  # nunca hay contraseña por defecto
        with mock.patch.dict(os.environ, {"TQT_ADMIN_PASSWORD": "corta"}):              # menos de 8 caracteres = no configurada
            self.assertEqual(self.client.post("/api/admin/login", json={"password": "corta"}).status_code, 503)

    def test_02_login_correcto_e_incorrecto(self):
        self.assertTrue(self.client.get("/api/admin/estado").json()["habilitado"])
        bad = self.client.post("/api/admin/login", json={"password": "mala"})
        self.assertEqual(bad.status_code, 401)
        self.assertNotIn("token", bad.text)
        ok = self.client.post("/api/admin/login", json={"password": CLAVE})
        self.assertEqual(ok.status_code, 200)
        self.assertEqual((ok.json()["expira_en"], ok.json()["token"].count(".")), (1800, 1))
        self.assertEqual(self.client.get("/api/admin/resumen", headers={"X-Admin-Token": ok.json()["token"]}).status_code, 200)
        self.assertEqual(self.client.post("/api/admin/login", json={}).status_code, 422)

    def test_03_la_clave_se_guarda_como_hash_pbkdf2_con_sal(self):
        self.token()
        with db.get_db() as c:
            filas = {r["clave"]: r["valor"] for r in c.execute("SELECT clave, valor FROM ajustes WHERE clave LIKE 'admin_%'")}
        self.assertTrue(filas["admin_hash"].startswith("pbkdf2_sha256$"))
        alg, it, sal, h = filas["admin_hash"].split("$")
        self.assertEqual(it, "1000")
        self.assertTrue(len(sal) >= 20 and len(h) >= 40)
        self.assertTrue(all(CLAVE not in v for v in filas.values()))            # jamás en claro
        self.assertNotEqual(filas["admin_hash"], filas["admin_env_hash"])          # sal distinta en cada hash
        self.assertTrue(admin_auth.verificar_clave(CLAVE, filas["admin_hash"]))
        self.assertFalse(admin_auth.verificar_clave(CLAVE + "x", filas["admin_hash"]))
        self.assertFalse(admin_auth.verificar_clave(CLAVE, "basura"))
        self.assertNotEqual(admin_auth.hash_clave(CLAVE), admin_auth.hash_clave(CLAVE))  # sal aleatoria

    def test_04_bloqueo_tras_5_fallos_y_desbloqueo_por_tiempo(self):
        reloj = [1000.0]
        admin_auth.limitador._reloj = lambda: reloj[0]
        self.addCleanup(lambda: setattr(admin_auth.limitador, "_reloj", __import__("time").monotonic))
        codigos = [self.client.post("/api/admin/login", json={"password": "x"}).status_code for _ in range(5)]
        self.assertEqual(codigos, [401, 401, 401, 401, 429])                         # el 5.º fallo ya bloquea
        r = self.client.post("/api/admin/login", json={"password": CLAVE})           # aun con la clave correcta
        self.assertEqual(r.status_code, 429)
        self.assertEqual(int(r.headers["retry-after"]), 300)
        reloj[0] += 299
        self.assertEqual(self.client.post("/api/admin/login", json={"password": CLAVE}).status_code, 429)
        reloj[0] += 2                                                                 # pasaron 5 minutos
        self.assertEqual(self.client.post("/api/admin/login", json={"password": CLAVE}).status_code, 200)

    def test_05_un_login_correcto_reinicia_el_contador(self):
        for _ in range(4):
            self.client.post("/api/admin/login", json={"password": "x"})
        self.token()
        for _ in range(4):
            self.assertEqual(self.client.post("/api/admin/login", json={"password": "x"}).status_code, 401)

    def test_06_retardo_uniforme(self):
        import time
        with mock.patch.object(settings, "ADMIN_RETARDO_MS", 250):
            for clave in (CLAVE, "mala"):
                t0 = time.monotonic()
                self.client.post("/api/admin/login", json={"password": clave})
                self.assertGreaterEqual(time.monotonic() - t0, 0.24, clave)

    def test_07_token_ausente_manipulado_o_caducado(self):
        h = self.token()
        self.assertEqual(self.client.get("/api/admin/resumen").status_code, 401)
        self.assertEqual(self.client.get("/api/admin/resumen", headers={"X-Admin-Token": ""}).status_code, 401)
        self.assertEqual(self.client.get("/api/admin/resumen", headers={"X-Admin-Token": "abc"}).status_code, 401)
        cuerpo, firma = h["X-Admin-Token"].split(".")
        for falso in (f"{cuerpo}.{firma[:-2]}AA", f"{cuerpo[:-2]}AA.{firma}", f"{firma}.{cuerpo}", f"{cuerpo}."):
            r = self.client.get("/api/admin/resumen", headers={"X-Admin-Token": falso})
            self.assertEqual(r.status_code, 401, falso)
        caducado = admin_auth.emitir_token(ahora=__import__("time").time() - 3600)["token"]  # emitido hace 1 h (TTL 30 min)
        r = self.client.get("/api/admin/resumen", headers={"X-Admin-Token": caducado})
        self.assertEqual(r.status_code, 401)
        self.assertIn("caduc", r.json()["detail"])
        fresco = admin_auth.emitir_token(ahora=__import__("time").time() - 1700)["token"]      # 28 min: aún vale
        self.assertEqual(self.client.get("/api/admin/resumen", headers={"X-Admin-Token": fresco}).status_code, 200)

    def test_08_token_firmado_con_otro_secreto_no_vale(self):
        h = self.token()
        with db.transaction() as c:
            c.execute("UPDATE ajustes SET valor = ? WHERE clave = 'admin_secret'", (admin_auth._b64(b"otro-secreto-de-32-bytes-xxxxxxx"),))
        self.assertEqual(self.client.get("/api/admin/resumen", headers=h).status_code, 401)

    def test_09_cambiar_clave(self):
        h = self.token()
        url = "/api/admin/cambiar-clave"
        self.assertEqual(self.client.post(url, json={"actual": CLAVE, "nueva": "x"}).status_code, 401)     # sin token
        self.assertEqual(self.client.post(url, headers=h, json={"actual": CLAVE, "nueva": "corta"}).status_code, 400)
        self.assertEqual(self.client.post(url, headers=h, json={"actual": CLAVE, "nueva": CLAVE}).status_code, 400)
        self.assertEqual(self.client.post(url, headers=h, json={"actual": "incorrecta", "nueva": "nueva-clave-99"}).status_code, 401)
        r = self.client.post(url, headers=h, json={"actual": CLAVE, "nueva": "nueva-clave-99"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(self.client.get("/api/admin/resumen", headers=h).status_code, 401)            # token viejo invalidado
        nuevo = {"X-Admin-Token": r.json()["token"]}
        self.assertEqual(self.client.get("/api/admin/resumen", headers=nuevo).status_code, 200)
        self.assertEqual(self.client.post("/api/admin/login", json={"password": CLAVE}).status_code, 401)
        self.assertEqual(self.client.post("/api/admin/login", json={"password": "nueva-clave-99"}).status_code, 200)
        self.assertIn("ADMIN_CLAVE_CAMBIADA", self.eventos())

    def test_10_la_clave_cambiada_sobrevive_al_reinicio_y_la_variable_nueva_la_restablece(self):
        h = self.token()
        self.client.post("/api/admin/cambiar-clave", headers=h, json={"actual": CLAVE, "nueva": "nueva-clave-99"})
        admin_auth.sincronizar()                                    # "reinicio" con la misma variable de entorno
        self.assertEqual(self.client.post("/api/admin/login", json={"password": "nueva-clave-99"}).status_code, 200)
        with mock.patch.dict(os.environ, {"TQT_ADMIN_PASSWORD": "clave-nueva-del-entorno"}):  # el admin la olvidó: cambia la variable
            self.assertEqual(self.client.post("/api/admin/login", json={"password": "clave-nueva-del-entorno"}).status_code, 200)
            self.assertEqual(self.client.post("/api/admin/login", json={"password": "nueva-clave-99"}).status_code, 401)

    def test_11_cambiar_clave_con_actual_incorrecta_cuenta_para_el_bloqueo(self):
        h = self.token()
        codigos = [self.client.post("/api/admin/cambiar-clave", headers=h, json={"actual": "x", "nueva": "otra-clave-123"}).status_code
                   for _ in range(5)]
        self.assertEqual(codigos, [401, 401, 401, 401, 429])


class TestBorrados(AdminBase):
    def test_12_resumen(self):
        t = crear_par(self.lote, r3=True)
        r = self.client.get("/api/admin/resumen", headers=self.token()).json()
        lote = next(x for x in r["lotes"] if x["id"] == self.lote)
        self.assertGreaterEqual(lote["tarjetas"], 1)
        self.assertEqual(lote["por_estado"]["PENDIENTE"], lote["tarjetas"])
        self.assertGreaterEqual(r["inventario"]["R3"]["ASIGNADA"], 1)
        self.assertGreater(r["bd_bytes"], 0)
        self.assertEqual(self.client.get("/api/admin/resumen").status_code, 401)
        self.assertTrue(t["id"])

    def _delete(self, url, json, h):
        return self.client.request("DELETE", url, json=json, headers=h)

    def test_13_borrar_tarjetas_liberando_las_pcb(self):
        a, b = crear_par(self.lote, r3=True), crear_par(self.lote)
        db.set_prueba(a["id"], "soldadura", "OK")
        h = self.token()
        antes = self.respaldos()
        with self.client.websocket_connect("/ws?client_type=monitor") as ws:
            recibir_ws(ws, evento="CONEXION_ESTABLECIDA")
            r = self._delete("/api/admin/tarjetas", {"ids": [a["id"], 999999], "liberar_pcb": True}, h)
            msg = recibir_ws(ws, evento="ADMIN_CAMBIO")
        self.assertEqual(r.status_code, 200, r.text)
        d = r.json()
        self.assertEqual((d["tarjetas_borradas"], d["pcb_liberadas"], d["pcb_eliminadas"], d["no_encontradas"]), (1, 3, 0, [999999]))
        self.assertEqual(msg["data"]["accion"], "BORRAR_TARJETAS")
        self.assertEqual(msg["data"]["respaldo"], d["respaldo"])
        # respaldo creado ANTES de borrar: contiene la tarjeta borrada y sus pruebas
        self.assertEqual(len(self.respaldos()), len(antes) + 1)
        self.assertIn(d["respaldo"], self.respaldos())
        c = sqlite3.connect(settings.BACKUP_DIR / d["respaldo"])
        try:
            self.assertEqual(c.execute("SELECT COUNT(*) FROM tarjetas_produccion WHERE id = ?", (a["id"],)).fetchone()[0], 1)
            self.assertEqual(c.execute("SELECT soldadura FROM pruebas_historial WHERE tarjeta_id = ?", (a["id"],)).fetchone()[0], "OK")
        finally:
            c.close()
        # efectos
        self.assertIsNone(db.get_tarjeta_by_id(a["id"]))
        self.assertIsNotNone(db.get_tarjeta_by_id(b["id"]))
        for pid in (a["pcb_r1_id"], a["pcb_r2_id"], a["pcb_r3_id"]):
            p = inv.get_pcb(pid)
            self.assertEqual((p["estado_ciclo"], p["tarjeta_id"]), ("DISPONIBLE", None))
        with db.get_db() as c2:
            self.assertEqual(c2.execute("SELECT COUNT(*) FROM pruebas_historial WHERE tarjeta_id = ?", (a["id"],)).fetchone()[0], 0)
        self.assertIn("ADMIN_TARJETAS_BORRADAS", self.eventos())

    def test_14_borrar_tarjetas_eliminando_tambien_las_pcb(self):
        a = crear_par(self.lote)
        r = self._delete("/api/admin/tarjetas", {"ids": [a["id"]], "liberar_pcb": False}, self.token())
        self.assertEqual((r.json()["pcb_eliminadas"], r.json()["pcb_liberadas"]), (2, 0))
        self.assertIsNone(inv.get_pcb(a["pcb_r1_id"]))
        self.assertIsNone(inv.get_pcb(a["pcb_r2_id"]))

    def test_15_borrar_tarjetas_inexistentes_es_404_y_no_hace_respaldo(self):
        h = self.token()
        antes = self.respaldos()
        r = self._delete("/api/admin/tarjetas", {"ids": [987654]}, h)
        self.assertEqual(r.status_code, 404)
        self.assertEqual(self.respaldos(), antes)
        self.assertEqual(self._delete("/api/admin/tarjetas", {"ids": []}, h).status_code, 422)
        self.assertEqual(self._delete("/api/admin/tarjetas", {"ids": [1]}, None).status_code, 401)
        self.assertEqual(self.respaldos(), antes)  # sin token no se hace nada, ni respaldo

    def test_16_borrar_pcb_deja_la_tarjeta_incompleta_y_no_liberada(self):
        t = crear_par(self.lote)
        for etapa in ("soldadura", "programacion", "prueba_pcb", "integracion", "prueba_final"):
            db.set_prueba(t["id"], etapa, "OK")
        self.assertEqual(db.get_tarjeta_by_id(t["id"])["estado_general"], "LIBERADO")
        suelta = inv.registrar_manual("R1", None, None, 1)["pcbs"][0]["id"]
        antes = self.respaldos()
        r = self._delete("/api/admin/pcb", {"ids": [t["pcb_r2_id"], suelta, 999999]}, self.token())
        self.assertEqual(r.status_code, 200, r.text)
        d = r.json()
        self.assertEqual((d["pcb_eliminadas"], d["tarjetas_afectadas"], d["no_encontradas"]), (2, 1, [999999]))
        self.assertEqual(len(self.respaldos()), len(antes) + 1)
        t2 = db.get_tarjeta_by_id(t["id"])
        self.assertFalse(t2["completa"])
        self.assertEqual(t2["estado_general"], "EN PROCESO")   # nunca LIBERADA sin R2
        self.assertEqual(self._delete("/api/admin/pcb", {"ids": [999999]}, self.token()).status_code, 404)
        self.assertIn("ADMIN_PCB_BORRADAS", self.eventos())

    def test_17_vaciar_lote_pide_confirmacion(self):
        t = crear_par(self.lote)
        h = self.token()
        antes = self.respaldos()
        for mala in ("", "vaciar", "VACIAR ", "SI"):
            r = self.client.post(f"/api/admin/lote/{self.lote}/vaciar", headers=h, json={"confirmar": mala})
            self.assertEqual(r.status_code, 400, mala)
        self.assertEqual(self.respaldos(), antes)                       # confirmación mala: sin respaldo y sin borrar
        self.assertIsNotNone(db.get_tarjeta_by_id(t["id"]))
        self.assertEqual(self.client.post("/api/admin/lote/999999/vaciar", headers=h, json={"confirmar": "VACIAR"}).status_code, 404)

        r = self.client.post(f"/api/admin/lote/{self.lote}/vaciar", headers=h, json={"confirmar": "VACIAR"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertGreaterEqual(r.json()["tarjetas_borradas"], 1)
        self.assertIn(r.json()["respaldo"], self.respaldos())
        self.assertEqual(db.list_tarjetas(lote_id=self.lote)[1], 0)
        self.assertEqual(inv.get_pcb(t["pcb_r1_id"])["estado_ciclo"], "DISPONIBLE")   # las PCB vuelven al inventario
        self.assertIn("ADMIN_LOTE_VACIADO", self.eventos())

    def test_18_vaciar_lote_eliminando_pcb(self):
        t = crear_par(self.lote)
        r = self.client.post(f"/api/admin/lote/{self.lote}/vaciar", headers=self.token(),
                             json={"confirmar": "VACIAR", "eliminar_pcb": True})
        self.assertEqual(r.status_code, 200)
        self.assertIsNone(inv.get_pcb(t["pcb_r1_id"]))

    def test_19_reset_pide_confirmacion_y_conserva_lotes_y_ajustes(self):
        crear_par(self.lote, r3=True)
        inv.set_version_defecto("31")
        lotes = [l["id"] for l in db.list_lotes()]
        h = self.token()
        antes = self.respaldos()
        for mala in ("", "borrar todo", "BORRAR", "VACIAR"):
            self.assertEqual(self.client.post("/api/admin/reset", headers=h, json={"confirmar": mala}).status_code, 400, mala)
        self.assertEqual(self.respaldos(), antes)

        r = self.client.post("/api/admin/reset", headers=h, json={"confirmar": "BORRAR TODO"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertGreaterEqual(r.json()["borrado"]["pcb_inventario"], 3)
        self.assertIn(r.json()["respaldo"], self.respaldos())
        res = admin_ops.resumen()
        self.assertEqual((res["totales"]["tarjetas"], res["totales"]["pcb"]), (0, 0))
        self.assertGreaterEqual(res["totales"]["bitacora"], 2)               # el historial de movimientos se conserva
        self.assertIn("ADMIN_RESET", self.eventos(""))                        # y queda el registro del propio borrado
        self.assertEqual([l["id"] for l in db.list_lotes()], lotes)          # lotes conservados
        self.assertEqual(inv.get_version_defecto(), "31")                    # ajustes conservados
        self.assertEqual(self.client.post("/api/admin/login", json={"password": CLAVE}).status_code, 200)  # y la clave del admin
        c = sqlite3.connect(settings.BACKUP_DIR / r.json()["respaldo"])
        try:
            self.assertGreaterEqual(c.execute("SELECT COUNT(*) FROM pcb_inventario").fetchone()[0], 3)   # el respaldo sí tiene los datos
        finally:
            c.close()

    def test_20_respaldos_con_nombre_unico_y_listados(self):
        n1 = admin_ops.respaldar("prueba uno!")
        n2 = admin_ops.respaldar("prueba uno!")
        self.assertNotEqual(n1, n2)
        self.assertRegex(n1, r"^tqt_\d{8}_\d{6}_\d{6}_prueba_uno\.db$")
        self.assertEqual({n1, n2} <= {x["nombre"] for x in admin_ops.listar_respaldos(50)}, True)


class TestExportacion(AdminBase):
    def setUp(self):
        super().setUp()
        admin_ops.reset_total("prueba", None)  # cada test parte sin tarjetas ni PCB (la plantilla empieza en 0011)
    def _con_datos(self):
        t = crear_par(self.lote, num="0011", r3=True)
        for etapa in ("soldadura", "programacion", "prueba_pcb"):
            db.set_prueba(t["id"], etapa, "OK")
        inv.actualizar_datos_tarjeta(t["id"], {"firmware_r1": "3.2", "semana_produccion": 38, "fecha_real": "2026-09-20"})
        return db.get_tarjeta_by_id(t["id"])

    def _mensual_nombre(self):
        l = db.get_lote_by_id(self.lote)
        return f"Control_Produccion_TQT_{MESES_ES[l['mes'] - 1]}_{l['anio']}.xlsx"

    def test_21_descarga_con_la_estructura_de_la_plantilla(self):
        t = self._con_datos()
        mensual = settings.EXCEL_DIR / self._mensual_nombre()
        antes_dir = sorted(p.name for p in settings.EXCEL_DIR.glob("*")) if settings.EXCEL_DIR.exists() else []
        r = self.client.get(f"/api/admin/export/excel?lote_id={self.lote}", headers=self.token())
        self.assertEqual(r.status_code, 200, r.text[:200])
        self.assertIn("spreadsheetml", r.headers["content-type"])
        self.assertIn(self._mensual_nombre(), r.headers["content-disposition"])
        self.assertEqual(r.headers["x-integridad-valida"], "true")
        self.assertTrue(r.content.startswith(b"PK"))

        destino = Path(self.tmp.name) / "descarga.xlsx"
        destino.write_bytes(r.content)
        wb = openpyxl.load_workbook(destino)                                    # abre con openpyxl
        self.assertEqual(sorted(wb.sheetnames), sorted(["Producción", "Catálogo PCB", "Pruebas", "Etiquetas"]))
        from app.services.excel_sync import ExcelSyncEngine
        plantilla = str(ExcelSyncEngine().get_template_path())
        antes, despues = _formulas(plantilla), _formulas(destino)
        self.assertEqual((len(antes), len(despues)), (2211, 2211))              # las 2,211 fórmulas
        self.assertEqual([k for k in antes if antes[k] != despues.get(k)], [])  # ninguna cambió
        self.assertEqual(_totales(structure_counts(plantilla)), _totales(structure_counts(str(destino))))  # validaciones y formatos
        self.assertEqual(structure_counts(str(destino))["Producción"]["validaciones_x14"], 2)
        # contiene los datos
        prod, cat, pru = wb["Producción"], wb["Catálogo PCB"], wb["Pruebas"]
        self.assertEqual((prod["A2"].value, prod["D2"].value, prod["I2"].value), ("0011", "3.2", "2026-W38"))
        self.assertEqual((cat["B3"].value, cat["G3"].value), (t["mac_r1"], t["mac_r2"]))
        self.assertEqual(cat["C3"].value, "FUNCIONAL")
        bitacora = [(pru.cell(r_, 11).value, pru.cell(r_, 12).value) for r_ in range(3, 12) if pru.cell(r_, 11).value]
        self.assertIn(("Soldadura", "OK"), bitacora)
        # el archivo mensual en uso no se tocó y no quedaron temporales
        self.assertEqual(sorted(p.name for p in settings.EXCEL_DIR.glob("*")) if settings.EXCEL_DIR.exists() else [], antes_dir)
        self.assertFalse(mensual.exists() and mensual.stat().st_size == 0)

    def test_22_no_toca_el_mensual_en_uso(self):
        self._con_datos()
        excel = settings.EXCEL_DIR / self._mensual_nombre()
        excel.parent.mkdir(parents=True, exist_ok=True)
        excel.write_bytes(b"PK-contenido-en-uso")
        db.set_ruta_excel(self.lote, str(excel))
        sha = hashlib.sha256(excel.read_bytes()).hexdigest()
        r = self.client.get("/api/admin/export/excel", headers=self.token())     # lote activo
        self.assertEqual(r.status_code, 200)
        self.assertEqual(hashlib.sha256(excel.read_bytes()).hexdigest(), sha)
        self.assertFalse(list(settings.EXCEL_DIR.glob("*.bak")))
        excel.unlink()

    def test_23_guardar_en_carpeta_permitida(self):
        self._con_datos()
        h = self.token()
        carpeta = settings.EXCEL_DIR / "exportaciones"
        r = self.client.get("/api/admin/export/excel", params={"lote_id": self.lote, "ruta": str(carpeta)}, headers=h)
        self.assertEqual(r.status_code, 200, r.text)
        d = r.json()
        self.assertTrue(d["success"] and d["integridad_valida"])
        self.assertEqual(Path(d["ruta"]), (carpeta / self._mensual_nombre()).resolve())
        self.assertTrue(Path(d["ruta"]).exists())
        self.assertEqual(len(_formulas(d["ruta"])), 2211)
        self.assertEqual([p.name for p in carpeta.glob("*")], [self._mensual_nombre()])  # sin .bak ni temporales
        # ruta a un archivo concreto
        r2 = self.client.get("/api/admin/export/excel", params={"ruta": str(carpeta / "otro.xlsx")}, headers=h)
        self.assertEqual(r2.status_code, 200)
        self.assertTrue((carpeta / "otro.xlsx").exists())

    def test_24_rutas_fuera_de_las_carpetas_permitidas_y_sin_token(self):
        h = self.token()
        fuera = Path(tempfile.gettempdir()) / "otro_lugar_admin"
        self.assertEqual(self.client.get("/api/admin/export/excel", params={"ruta": str(fuera)}, headers=h).status_code, 400)
        self.assertEqual(self.client.get("/api/admin/export/excel", params={"ruta": r"C:\Windows\x.xlsx"}, headers=h).status_code, 400)
        self.assertFalse(fuera.exists())
        self.assertIn(self.client.get("/api/admin/export/excel").status_code, (200, 404))   # v1.3.44: la descarga no pide clave
        self.assertEqual(self.client.get("/api/admin/export/excel", params={"ruta": str(fuera)}).status_code, 401)   # guardar en disco sí
        self.assertEqual(self.client.get("/api/admin/export/excel?lote_id=987654", headers=h).status_code, 404)


if __name__ == "__main__":
    unittest.main()
