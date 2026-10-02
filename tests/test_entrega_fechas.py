"""v1.2.16: fechas de llegada / finalizado / entrega y gabinete (Quintalock/Translock): disparadores, API de datos, Excel ida y vuelta."""
import _aislamiento  # noqa: F401  (debe ser lo primero)
import tempfile
import unittest
from datetime import date
from pathlib import Path

import openpyxl

from app.database import db, inventario as inv
from app.services.excel_sync import ExcelSyncEngine

from _aislamiento import crear_par, verificar_aislamiento


class TestEntregaFechas(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        verificar_aislamiento()
        cls.tmp = tempfile.TemporaryDirectory(); cls.dir = Path(cls.tmp.name); cls.n = 0

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def setUp(self):
        type(self).n += 1
        self.path = self.dir / f"e{self.n}.db"; db.init_db(self.path)
        self.engine = ExcelSyncEngine(db_path=self.path)
        self.excel = self.engine.create_monthly_excel(9, 2026, target_path=str(self.dir / f"e{self.n}.xlsx"))
        self.lote = db.create_lote("E", 9, 2026, ruta_excel=self.excel, activo=True, db_path=self.path)["id"]
        self.hoy = date.today().isoformat()

    def test_01_llegada_y_finalizado_automaticos(self):
        t = crear_par(self.lote, num="0011", r3=True, macs=False, path=self.path)
        self.assertEqual(t["fecha_llegada"], self.hoy)           # recepción confirmada hoy
        self.assertIsNone(t["fecha_finalizado"])                 # sin MAC todavía
        self.assertIsNone(t["gabinete"])
        for pid in (t["pcb_r1_id"], t["pcb_r2_id"]):
            inv.set_mac(pid, "70:4B:CA:5B:00:%02X" % pid, db_path=self.path)
        self.assertEqual(db.get_tarjeta_by_id(t["id"], db_path=self.path)["fecha_finalizado"], self.hoy)

    def test_02_sin_r3_no_se_finaliza(self):
        t = crear_par(self.lote, num="0012", r3=False, macs=True, path=self.path)
        self.assertIsNone(t["fecha_finalizado"])
        r3 = inv.registrar_manual("R3", "30", "0012", db_path=self.path)["pcbs"][0]["id"]
        inv.confirmar_recepcion([r3], db_path=self.path)
        inv.asignar_pcb(t["id"], "R3", r3, db_path=self.path)
        self.assertEqual(db.get_tarjeta_by_id(t["id"], db_path=self.path)["fecha_finalizado"], self.hoy)

    def test_03_editar_fechas_y_gabinete(self):
        t = crear_par(self.lote, num="0013", r3=True, path=self.path)
        r = inv.actualizar_datos_tarjeta(t["id"], {"fecha_real": "2026-09-30", "gabinete": "translock", "fecha_llegada": "2026-09-01"}, db_path=self.path)
        self.assertEqual((r["fecha_real"], r["gabinete"], r["fecha_llegada"]), ("2026-09-30", "Translock", "2026-09-01"))
        with self.assertRaises(ValueError):
            inv.actualizar_datos_tarjeta(t["id"], {"gabinete": "Otro"}, db_path=self.path)
        with self.assertRaises(ValueError):
            inv.actualizar_datos_tarjeta(t["id"], {"fecha_real": "30/09/2026"}, db_path=self.path)
        r = inv.actualizar_datos_tarjeta(t["id"], {"gabinete": ""}, db_path=self.path)
        self.assertIsNone(r["gabinete"])

    def test_04_excel_exporta_e_importa_las_fechas_y_el_gabinete(self):
        t = crear_par(self.lote, num="0014", r3=True, path=self.path)
        inv.actualizar_datos_tarjeta(t["id"], {"fecha_real": "2026-09-30", "gabinete": "Quintalock", "fecha_llegada": "2026-09-02", "fecha_finalizado": "2026-09-20"}, db_path=self.path)
        self.engine.export_to_excel(self.excel, self.lote, db_path=self.path)
        ws = openpyxl.load_workbook(self.excel)["Producción"]
        enc = {str(c.value): c.column for c in ws[1] if c.value}
        self.assertIn("Fecha Llegada", enc); self.assertIn("Fecha Finalizado", enc)
        fila = next(r for r in range(2, 102) if str(ws.cell(r, 1).value) == "0014")
        self.assertEqual(ws.cell(fila, 12).value, "Quintalock")                                  # L: Gabinete
        self.assertEqual(ws.cell(fila, 11).value.date().isoformat(), "2026-09-30")               # K: Fecha Real Entrega
        self.assertEqual(ws.cell(fila, enc["Fecha Llegada"]).value.date().isoformat(), "2026-09-02")
        self.assertEqual(ws.cell(fila, enc["Fecha Finalizado"]).value.date().isoformat(), "2026-09-20")
        # ida y vuelta: borra los datos de la BD y vuelve a importar
        with db.transaction(self.path) as c:
            c.execute("UPDATE tarjetas_produccion SET fecha_real=NULL, gabinete=NULL, fecha_llegada=NULL, fecha_finalizado=NULL")
        self.engine.import_from_excel(self.excel, self.lote, db_path=self.path)
        r = db.get_tarjeta_by_id(t["id"], db_path=self.path)
        self.assertEqual((r["fecha_real"], r["gabinete"], r["fecha_llegada"], r["fecha_finalizado"]), ("2026-09-30", "Quintalock", "2026-09-02", "2026-09-20"))

    def test_05_migracion_de_una_base_anterior(self):
        p = self.dir / "vieja.db"; db.init_db(p)
        lote = db.create_lote("V", 9, 2026, activo=True, db_path=p)["id"]
        t = crear_par(lote, num="0020", r3=True, path=p)
        with db.transaction(p) as c:   # simula una base anterior: sin columnas ni disparadores
            for tr in ("trg_tarjeta_fechas_alta", "trg_tarjeta_fechas_placas", "trg_tarjeta_fechas_mac"):
                c.execute(f"DROP TRIGGER IF EXISTS {tr}")
            for col in ("fecha_llegada", "fecha_finalizado", "gabinete"):
                c.execute(f"ALTER TABLE tarjetas_produccion DROP COLUMN {col}")
        db.init_db(p)   # migra
        r = db.get_tarjeta_by_id(t["id"], db_path=p)
        self.assertEqual(r["fecha_llegada"], self.hoy); self.assertEqual(r["fecha_finalizado"], self.hoy)

class TestFmtFecha(unittest.TestCase):
    """v1.3.19: la importación del Excel solo acepta fechas reales; lo demás se ignora en vez de guardar basura."""
    def test_formatos(self):
        from datetime import datetime
        from app.services.excel_sync import _fmt_fecha
        self.assertEqual(_fmt_fecha(datetime(2026, 10, 5, 13, 0)), "2026-10-05")
        self.assertEqual(_fmt_fecha("2026-10-05"), "2026-10-05")
        self.assertEqual(_fmt_fecha("05/10/2026"), "2026-10-05")
        self.assertIsNone(_fmt_fecha("pronto")); self.assertIsNone(_fmt_fecha(46300)); self.assertIsNone(_fmt_fecha(None))

if __name__ == "__main__":
    unittest.main()
