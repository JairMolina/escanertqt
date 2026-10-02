"""Excel (openpyxl): motor, endpoints REST y rutas permitidas. (La etiqueta DYMO se prueba en test_dymo.py.)"""
import _aislamiento  # noqa: F401  (debe ser lo primero)
import os
import shutil
import tempfile
import unittest
from unittest import mock

from app.services import admin_auth
from pathlib import Path

import openpyxl
from starlette.testclient import TestClient

from app.config import settings
from app.database import db, inventario as inv
from app.main import app
from app.services.excel_sync import ExcelSyncEngine, get_sheet_by_name

from _aislamiento import crear_par, nuevo_lote, verificar_aislamiento


class TestExcelIntegration(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        verificar_aislamiento()
        db.init_db()
        cls.client = TestClient(app)
        # export/import/create-monthly exigen sesión de administrador
        cls._env = mock.patch.dict(os.environ, {"TQT_ADMIN_PASSWORD": "clave-excel-123"})
        cls._env.start()
        admin_auth.limitador.reiniciar()        # otros tests (test_admin) dejan la IP de pruebas bloqueada
        tok = cls.client.post("/api/admin/login", json={"password": "clave-excel-123"}).json()["token"]
        cls.client_sin_token = TestClient(app)
        cls.client.headers["X-Admin-Token"] = tok
        cls.test_dir = Path(tempfile.mkdtemp())
        cls.db_path = cls.test_dir / "excel.db"
        db.init_db(cls.db_path)
        cls.engine = ExcelSyncEngine(db_path=cls.db_path)
        cls.template = cls.engine.get_template_path()
        settings.EXCEL_DIR.mkdir(parents=True, exist_ok=True)

    @classmethod
    def tearDownClass(cls):
        admin_auth.limitador.reiniciar()
        cls._env.stop()
        shutil.rmtree(cls.test_dir, ignore_errors=True)

    def test_00_rutas_de_base_de_datos_exigen_contrasena(self):
        """Descarga completa, importación y creación de mensuales sin token de admin => 401."""
        c = self.client_sin_token
        self.assertEqual(c.get("/api/export/excel").status_code, 401)
        self.assertEqual(c.post("/api/excel/import", json={"excel_path": "x.xlsx"}).status_code, 401)
        self.assertEqual(c.post("/api/excel/create-monthly", json={"mes": 10, "anio": 2026}).status_code, 401)
        self.assertEqual(c.get("/docs").status_code, 404)          # explorador de API apagado por defecto
        self.assertEqual(c.get("/openapi.json").status_code, 404)

    # ------------------------------------------------------------------ Excel (motor)
    def test_01_crear_mensual_desde_plantilla(self):
        destino = self.test_dir / "Control_Produccion_TQT_Octubre_2026.xlsx"
        ruta = self.engine.create_monthly_excel(mes=10, anio=2026, target_path=str(destino))
        self.assertTrue(os.path.exists(ruta))
        wb = openpyxl.load_workbook(ruta, data_only=False)
        for hoja in ("Producción", "Catálogo PCB", "Pruebas", "Etiquetas"):
            self.assertIsNotNone(get_sheet_by_name(wb, hoja))
        with self.assertRaises(ValueError):
            self.engine.create_monthly_excel(13, 2026)

    def test_02_export_conserva_formulas_clave(self):
        excel = self.test_dir / "preservacion.xlsx"
        shutil.copy2(self.template, excel)
        lote = db.create_lote("2026-10-X", 10, 2026, db_path=self.db_path)["id"]
        t = crear_par(lote, num="0011", path=self.db_path)
        inv.actualizar_datos_tarjeta(t["id"], {"firmware_r1": "2.5", "firmware_r2": "2.1", "semana_produccion": 40,
                                               "fecha_proyectada": "2026-10-15", "fecha_real": "2026-10-10"}, db_path=self.db_path)
        for etapa in ("soldadura", "programacion", "prueba_pcb", "integracion", "prueba_final"):
            db.set_prueba(t["id"], etapa, "OK", db_path=self.db_path)
        t = db.get_tarjeta_by_id(t["id"], db_path=self.db_path)

        antes = openpyxl.load_workbook(excel, data_only=False)
        formulas = {(h, ref): get_sheet_by_name(antes, h)[ref].value for h, ref in
                    (("Producción", "B2"), ("Producción", "C2"), ("Producción", "F2"), ("Catálogo PCB", "D3"), ("Catálogo PCB", "I3"),
                     ("Pruebas", "A2"), ("Pruebas", "G2"))}
        res = self.engine.export_to_excel(str(excel), lote)
        self.assertEqual(res["status"], "success")
        self.assertGreaterEqual(res["exported_tarjetas"], 1)
        self.assertEqual(res["avisos"], [])          # nombre R1 coincide con la fórmula: nada que avisar

        despues = openpyxl.load_workbook(excel, data_only=False)
        prod, cat, pru = (get_sheet_by_name(despues, n) for n in ("Producción", "Catálogo PCB", "Pruebas"))
        self.assertEqual((prod["A2"].value, prod["D2"].value, prod["E2"].value, prod["G2"].value, prod["I2"].value),
                         ("0011", "2.5", "TQT-R2-V30-0011", "2.1", "2026-W40"))
        self.assertEqual((cat["B3"].value, cat["G3"].value), (t["mac_r1"], t["mac_r2"]))
        self.assertEqual(str(prod["J2"].value)[:10], "2026-10-15")
        self.assertEqual(str(prod["K2"].value)[:10], "2026-10-10")
        self.assertTrue(str(pru["B2"].value.text if hasattr(pru["B2"].value, "text") else pru["B2"].value).startswith("="))
        for (h, ref), formula in formulas.items():
            self.assertEqual(get_sheet_by_name(despues, h)[ref].value, formula, f"{h}!{ref}")
        rep = self.engine.verify_excel_integrity(str(excel))
        self.assertTrue(rep["is_valid"], rep["summary"])

    def test_03_importar_plantilla_vacia(self):
        lote = db.create_lote("2026-11-IMP", 11, 2026, db_path=self.db_path)["id"]
        res = self.engine.import_from_excel(str(self.template), lote)
        self.assertEqual(res["status"], "success")
        self.assertEqual((res["imported_pcbs"], res["imported_tarjetas"]), (0, 0))  # los nombres de la plantilla sin MAC no son PCB reales

    # ------------------------------------------------------------------ API
    def test_04_api_verify(self):
        r = self.client.get("/api/excel/verify", params={"excel_path": str(self.template)})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(r.json()["is_valid"])
        self.assertGreater(r.json()["total_formulas"], 1000)

    def test_05_api_create_monthly_sin_sobrescribir(self):
        destino = settings.EXCEL_DIR / "Control_Octubre_API.xlsx"
        r = self.client.post("/api/excel/create-monthly", json={"mes": 10, "anio": 2026, "target_path": str(destino), "activar": True})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(r.json()["success"])
        self.assertFalse(r.json()["reutilizado"])
        self.assertTrue(os.path.exists(r.json()["excel_path"]))
        r2 = self.client.post("/api/excel/create-monthly", json={"mes": 10, "anio": 2026, "target_path": str(destino)})
        self.assertTrue(r2.json()["reutilizado"])   # el mismo mes no reemplaza el archivo con datos

    def test_06_api_rutas_fuera_de_las_carpetas_permitidas_se_rechazan(self):
        fuera = Path(tempfile.gettempdir()) / "otro_lugar" / "x.xlsx"
        self.assertEqual(self.client.post("/api/excel/create-monthly", json={"mes": 10, "anio": 2026, "target_path": str(fuera)}).status_code, 400)
        self.assertEqual(self.client.post("/api/excel/sync", json={"excel_path": str(fuera)}).status_code, 400)
        self.assertEqual(self.client.get("/api/excel/verify", params={"excel_path": r"C:\Windows\win.ini"}).status_code, 400)
        self.assertEqual(self.client.post("/api/excel/import", json={"excel_path": str(settings.EXCEL_DIR / "..\\..\\x.xlsx")}).status_code, 400)

    def test_07_api_sync_excel_end_to_end(self):
        """Lote nuevo -> tarjeta -> pruebas -> POST /api/sync/excel (botón del monitor) -> descargar el .xlsx."""
        lote = self.client.post("/api/lotes", json={"codigo_lote": "2026-12-E2E", "mes": 12, "anio": 2026}).json()
        lote_id = lote["id"]
        t = crear_par(lote_id, num="0128")
        self.client.put(f"/api/tarjetas/{t['id']}/pruebas", json={"etapa": "soldadura", "estado": "OK"})

        r = self.client.post(f"/api/sync/excel?lote_id={lote_id}")
        self.assertEqual(r.status_code, 200, r.text)
        d = r.json()
        self.assertTrue(d["success"])
        self.assertEqual(d["filename"], "Control_Produccion_TQT_Diciembre_2026.xlsx")   # plantilla mensual, no un libro nuevo
        self.assertEqual(d["total_tarjetas"], 1)
        self.assertTrue(d["integridad"]["is_valid"], d["integridad"]["summary"])
        self.assertEqual(d["integridad"]["total_formulas"], 2211)
        # (la BD de pruebas es compartida y acumula placas: el aviso de "no caben en el catálogo" no es lo que se prueba aquí)
        self.assertEqual([a for a in d["avisos"] if "PCB sueltas" not in a], [])
        self.assertEqual(d["download_url"], f"/api/export/excel?lote_id={lote_id}")

        dl = self.client.get(d["download_url"])
        self.assertEqual(dl.status_code, 200)
        self.assertTrue(dl.content.startswith(b"PK"))
        tmp = self.test_dir / "descarga.xlsx"
        tmp.write_bytes(dl.content)
        prod = get_sheet_by_name(openpyxl.load_workbook(tmp), "Producción")
        self.assertIn("0128", [prod.cell(r_, 1).value for r_ in range(2, 102)])

        r2 = self.client.post("/api/excel/sync", json={"lote_id": lote_id})   # ruta nativa del motor
        self.assertEqual((r2.status_code, r2.json()["status"]), (200, "success"))

    def test_08_api_export_sin_excel_404(self):
        lote = self.client.post("/api/lotes", json={"codigo_lote": "2027-05-SINXL", "mes": 5, "anio": 2027, "crear_excel": False}).json()
        self.assertEqual(self.client.get(f"/api/export/excel?lote_id={lote['id']}").status_code, 404)
        self.assertEqual(self.client.get("/api/export/excel?lote_id=9999999").status_code, 404)


if __name__ == "__main__":
    unittest.main()
