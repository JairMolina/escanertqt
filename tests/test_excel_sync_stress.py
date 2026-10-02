"""Estrés de integridad Excel: ciclo mensual completo con 50 tarjetas y verificación celda por celda.

- create_monthly_excel -> import_from_excel -> 50 emparejamientos con pruebas -> export_to_excel
- Reapertura con openpyxl: las 2,211 fórmulas de la plantilla siguen idénticas y sin errores de Excel.
"""
import _aislamiento  # noqa: F401  (debe ser lo primero)
import tempfile
import time
import unittest
from pathlib import Path

import openpyxl

from app.database import db, inventario as inv
from app.database.models import parse_tarjeta_code
from app.services.excel_sync import EXCEL_ERROR_CODES, ExcelSyncEngine, get_sheet_by_name

from _aislamiento import mac_unica, verificar_aislamiento

HOJAS = ["Producción", "Catálogo PCB", "Pruebas", "Etiquetas"]


def _mapa_formulas(wb):
    mapa, por_hoja = {}, {}
    for h in HOJAS:
        ws = get_sheet_by_name(wb, h)
        n = 0
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if v is None:
                    continue
                txt = v.text if hasattr(v, "text") else str(v)
                if txt.startswith("=") or hasattr(v, "text"):
                    mapa[f"{h}!{openpyxl.utils.get_column_letter(c)}{r}"] = txt
                    n += 1
        por_hoja[h] = n
    return mapa, por_hoja


class TestExcelSyncStress(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        verificar_aislamiento()
        cls.tmp = tempfile.TemporaryDirectory()
        cls.path = Path(cls.tmp.name) / "stress.db"
        db.init_db(cls.path)
        cls.engine = ExcelSyncEngine(db_path=cls.path)
        wb = openpyxl.load_workbook(str(cls.engine.get_template_path()), data_only=False)
        cls.orig, cls.por_hoja = _mapa_formulas(wb)
        wb.close()

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_01_plantilla_tiene_2211_formulas(self):
        self.assertEqual(len(self.orig), 2211)
        self.assertEqual(self.por_hoja, {"Producción": 300, "Catálogo PCB": 301, "Pruebas": 700, "Etiquetas": 910})

    def test_02_ciclo_mensual_50_tarjetas(self):
        destino = Path(self.tmp.name) / "Control_Produccion_TQT_Noviembre_2026.xlsx"
        creado = self.engine.create_monthly_excel(11, 2026, target_path=str(destino))
        lote_id = db.create_lote("2026-11-STRESS", 11, 2026, ruta_excel=creado, activo=True, db_path=self.path)["id"]

        imp = self.engine.import_from_excel(creado, lote_id, db_path=self.path)
        self.assertEqual(imp["status"], "success")
        self.assertEqual((imp["imported_pcbs"], imp["imported_tarjetas"]), (0, 0))   # plantilla vacía: nombres sin MAC no son PCB

        # En la plantilla el R2 de cada tarjeta viene prellenado (Producción!E) y NO siempre coincide con su número de R1
        wb_plantilla = openpyxl.load_workbook(str(self.engine.get_template_path()))
        prod_p = get_sheet_by_name(wb_plantilla, "Producción")
        r2_de = {prod_p.cell(r, 1).value: prod_p.cell(r, 5).value for r in range(2, 102)}
        wb_plantilla.close()
        self.assertTrue(any(r2_de[f"{i:04d}"] != f"TQT-R2-V30-{i:04d}" for i in range(11, 61)))  # hay casos con numeración distinta

        t0 = time.perf_counter()
        for i in range(11, 61):
            num = f"{i:04d}"
            nombre_r2 = parse_tarjeta_code(r2_de[num], "R2")
            r1 = inv.registrar_manual("R1", "30", num, db_path=self.path)["pcbs"][0]["id"]
            r2 = inv.registrar_manual("R2", nombre_r2["version"], nombre_r2["numero"], db_path=self.path)["pcbs"][0]["id"]
            inv.confirmar_recepcion([r1, r2], db_path=self.path)
            t = inv.crear_tarjeta(lote_id, num, r1, r2, db_path=self.path)
            inv.set_mac(r1, mac_unica(), db_path=self.path)
            inv.set_mac(r2, mac_unica(), db_path=self.path)
            inv.actualizar_datos_tarjeta(t["id"], {"firmware_r1": "3.2", "firmware_r2": "2.1", "semana_produccion": 45,
                                                   "fecha_proyectada": "2026-11-10", "fecha_real": "2026-11-08"}, db_path=self.path)
            estados = ["OK"] * 5 if i % 3 == 0 else ["OK", "OK", "PENDIENTE", "PENDIENTE", "PENDIENTE"]
            for etapa, est in zip(("soldadura", "programacion", "prueba_pcb", "integracion", "prueba_final"), estados):
                db.set_prueba(t["id"], etapa, est, db_path=self.path)
        t_pair = (time.perf_counter() - t0) * 1000 / 50

        self.assertGreaterEqual(db.list_tarjetas(lote_id=lote_id, limit=100, db_path=self.path)[1], 50)
        t0 = time.perf_counter()
        res = self.engine.export_to_excel(creado, lote_id, db_path=self.path)
        t_export = (time.perf_counter() - t0) * 1000
        self.assertEqual(res["status"], "success")
        self.assertEqual(res["exported_tarjetas"], 50)
        self.assertEqual(res["omitidas"], [])
        self.assertEqual(res["avisos"], [])

        wb = openpyxl.load_workbook(creado, data_only=False)
        despues, _ = _mapa_formulas(wb)
        self.assertEqual(len(despues), 2211)
        self.assertEqual([k for k, v in self.orig.items() if despues.get(k) != v], [])

        errores = []
        for h in HOJAS:
            for fila in get_sheet_by_name(wb, h).iter_rows():
                for c in fila:
                    txt = c.value.text if hasattr(c.value, "text") else str(c.value or "")
                    errores += [(h, c.coordinate, e) for e in EXCEL_ERROR_CODES if e in txt]
        self.assertEqual(errores, [])

        prod, cat, pru = (get_sheet_by_name(wb, n) for n in ("Producción", "Catálogo PCB", "Pruebas"))
        self.assertEqual((prod["A2"].value, prod["D2"].value, prod["E2"].value, prod["G2"].value, prod["I2"].value),
                         ("0011", "3.2", "TQT-R2-V30-0011", "2.1", "2026-W45"))
        # los IDs de la plantilla NO son consecutivos: cada tarjeta debe caer en SU renglón (donde A == su ID)
        for i in range(11, 61):
            fila = next(r for r in range(2, 102) if prod.cell(r, 1).value == f"{i:04d}")
            self.assertEqual(prod.cell(fila, 5).value, r2_de[f"{i:04d}"])
            self.assertEqual(prod.cell(fila, 9).value, "2026-W45")
        self.assertRegex(cat["B3"].value, r"^02:")            # MAC de R1 de la 0011 en el catálogo
        self.assertRegex(cat["G3"].value, r"^02:")

        # Bitácora masiva: cada etapa/estado agrupa hasta 15 IDs por fila
        filas = [(pru.cell(r, 11).value, pru.cell(r, 12).value, [pru.cell(r, c).value for c in range(13, 28) if pru.cell(r, c).value])
                 for r in range(3, 200) if pru.cell(r, 11).value]
        soldadura_ok = [x for x in filas if x[0] == "Soldadura" and x[1] == "OK"]
        self.assertEqual(sum(len(x[2]) for x in soldadura_ok), 50)
        self.assertTrue(all(len(x[2]) <= 15 for x in filas))
        self.assertEqual(sum(len(x[2]) for x in filas if x[0] == "Prueba Final" and x[1] == "OK"), 17)  # i % 3 == 0 -> 12,15,...,60

        rep = self.engine.verify_excel_integrity(creado)
        self.assertTrue(rep["is_valid"], rep["summary"])
        self.assertEqual(rep["total_formulas"], 2211)
        wb.close()
        print(f"\n  [métricas] emparejar+pruebas {t_pair:.1f} ms/tarjeta · exportar 50 tarjetas {t_export:.0f} ms")


if __name__ == "__main__":
    unittest.main()
