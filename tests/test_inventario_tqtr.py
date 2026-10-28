"""Inventario TQTR: semilla desde el Excel, cálculos del dashboard (comparados con los valores que Excel tenía calculados), CRUD y export."""
import _aislamiento  # noqa: F401  (debe ser lo primero)
import io
import tempfile
import unittest
from pathlib import Path

from app.database.db import ConflictoError, NoEncontradoError
from app.services import inventario_tqtr as svc

from _aislamiento import verificar_aislamiento


class TestInventarioTQTR(unittest.TestCase):
    def setUp(self):
        verificar_aislamiento()
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "inv_tqtr.db"

    def tearDown(self):
        svc._reiniciar_cache()
        self.tmp.cleanup()

    def comp(self, dash, desc):
        return next(f for f in dash["componentes"] if f["descripcion"] == desc)

    def efecto(self, valor, efecto):
        cat = svc.configuracion(db_path=self.path)["catalogos"]["movimiento"]
        cid = next(x["id"] for x in cat if x["valor"] == valor)
        svc.editar_catalogo(cid, {"efecto": efecto}, db_path=self.path)

    # ------------------------------------------------------------------ semilla
    def test_01_semilla_idempotente(self):
        d1 = svc.dashboard(db_path=self.path)
        self.assertEqual(d1["kpis"]["movimientos"], 52)  # 51 del Excel + entrada de sensores v1.3.39
        self.assertEqual(d1["kpis"]["componentes"], 24)  # v1.4.64: 3 relevadores → 1
        svc._reiniciar_cache()
        d2 = svc.dashboard(db_path=self.path)  # segunda preparación: no duplica nada
        self.assertEqual(d2["kpis"]["movimientos"], 52)
        conf = svc.configuracion(db_path=self.path)
        self.assertEqual(len(conf["productos"]), 5)
        self.assertIn("Devolución por defecto", [x["valor"] for x in conf["catalogos"]["movimiento"]])
        self.assertTrue(conf["catalogos"]["regla"])

    def test_02_dashboard_igual_al_excel(self):
        # El Excel solo suma "Entrada" en el dashboard; con la garantía en 0 los números deben coincidir exactamente.
        self.efecto("Desinstalado o Garantía", 0)
        d = svc.dashboard(db_path=self.path)
        esperado = {  # descripción: (disponible, consumida, ensambles) — valores en caché del Excel
            "PCB TQT-R1 V2.0 | PRINCIPAL": (89, 21, 89),
            "Cable uso rudo 4 hilos calibre 18 | Negro": (298.5, 5.5, 597),
            "Relevador automotriz 12 VCD": (232, 88, 29),  # v1.4.64: R1 58 + R2 116 + R3 58 en un solo stock (8 por armado)
            "Cinchos de nylon": (534, 66, 89),
            "Cable calibre 22 | R3 | Rojo": (84.5, 15.5, 169),
            "Caja Quintalock Metal Reforzada": (37, 10, 37),
            "Caja Translock Metal Reforzada": (16, 30, 16),
            "Actuador lineal de 12 VCD": (19, 11, 19),
        }
        for desc, (disp, cons, ens) in esperado.items():
            f = self.comp(d, desc)
            self.assertAlmostEqual(f["disponible"], disp, places=6, msg=desc)
            self.assertAlmostEqual(f["consumida"], cons, places=6, msg=desc)
            self.assertEqual(f["ensambles"], ens, desc)
        self.assertEqual(d["pcba_posibles"], {"R1": 29, "R2": 87, "R3": 49})  # v1.4.64: el relevador unificado cuenta en R1
        self.assertEqual(d["producto_final_posible"], 29)
        self.assertEqual((d["quintalock"], d["translock"]), (19, 16))
        self.assertEqual(d["sin_coincidencia"], [])

    def test_03_devueltas_en_revision(self):
        d = svc.dashboard(db_path=self.path)  # v1.3.40: la garantía no suma; queda en "Devueltas / en revisión"
        tl = self.comp(d, "Caja Translock Metal Reforzada")
        self.assertEqual((tl["disponible"], tl["en_revision"], d["translock"]), (16, 2, 16))
        r = svc.reclasificar_devueltas(tl["id"], "disponible", 1, db_path=self.path)
        self.assertEqual((r["disponible"], r["en_revision"]), (17, 1))
        regs = svc.listar_registros(movimiento=svc.MOV_A_DISPONIBLE, db_path=self.path)["items"]
        self.assertEqual([(x["modelo"], x["cantidad"]) for x in regs], [("Caja Translock Metal Reforzada", 1)])
        r = svc.reclasificar_devueltas(tl["id"], "baja", 1, db_path=self.path)
        self.assertEqual((r["disponible"], r["en_revision"]), (17, 0))
        with self.assertRaises(ValueError):
            svc.reclasificar_devueltas(tl["id"], "baja", 1, db_path=self.path)  # ya no hay en revisión
        svc._reiniciar_cache()
        self.assertEqual(self.comp(svc.dashboard(db_path=self.path), "Caja Translock Metal Reforzada")["disponible"], 17)  # migración idempotente

    def test_04_alertas_y_productos_bom(self):
        d = svc.dashboard(db_path=self.path)
        sensor = self.comp(d, "Sensor magnético CS1-U")   # v1.4.64: viene con el actuador; ya no limita ni se descuenta aparte
        self.assertEqual((sensor["en_dashboard"] if "en_dashboard" in sensor else 0, sensor["cantidad_ensamble"]), (0, 2))
        quinta = next(p for p in d["productos"] if p["producto"] == "Quinta completa | R1+R2+R3")
        self.assertEqual((quinta["posibles"], quinta["limitante"]), (19, "Actuador lineal de 12 VCD"))
        timer = next(p for p in d["productos"] if p["producto"] == "Solo Timer | R3")
        self.assertGreater(timer["posibles"], 0)
        # stock mínimo configurable
        pcb = self.comp(d, "PCB TQT-R1 V2.0 | PRINCIPAL")
        svc.editar_componente(pcb["id"], {"stock_minimo": 100}, db_path=self.path)
        self.assertEqual(self.comp(svc.dashboard(db_path=self.path), "PCB TQT-R1 V2.0 | PRINCIPAL")["estado"], "bajo")

    # ------------------------------------------------------------------ movimientos
    def test_05_crud_movimientos(self):
        p = self.path
        m = svc.crear_registro({"modelo": "Cinchos de nylon", "movimiento": "Salida", "cantidad": 34, "fecha_entrada": "2026-10-06", "destino": "Producción"}, db_path=p)
        self.assertEqual((m["categoria"], m["unidad"]), ("Accesorios para ensamble PCB", "piezas"))  # se completan del componente
        self.assertEqual(self.comp(svc.dashboard(db_path=p), "Cinchos de nylon")["disponible"], 500)
        svc.editar_registro(m["id"], {"cantidad": 4}, db_path=p)
        self.assertEqual(self.comp(svc.dashboard(db_path=p), "Cinchos de nylon")["disponible"], 530)
        r = svc.listar_registros(q="Producción", movimiento="Salida", db_path=p)
        self.assertIn(m["id"], [x["id"] for x in r["items"]])
        svc.eliminar_registro(m["id"], db_path=p)
        with self.assertRaises(NoEncontradoError):
            svc.obtener_registro(m["id"], db_path=p)
        for malo in ({"modelo": "X", "movimiento": "Robo", "cantidad": 1}, {"modelo": "X", "movimiento": "Entrada", "cantidad": 0},
                     {"modelo": "", "movimiento": "Entrada", "cantidad": 1}, {"modelo": "X", "movimiento": "Entrada", "cantidad": 1, "fecha_entrada": "31/12/2026"}):
            with self.assertRaises(ValueError):
                svc.crear_registro(malo, db_path=p)

    def test_15_migracion_1464_junta_consumos_de_relevadores(self):
        """v1.4.64: una BD con consumos separados R1/R2 de relevadores queda con una sola fila por tarjeta (sumada) y un solo componente."""
        from app.database import db
        p = self.path
        svc.dashboard(db_path=p)
        with db.transaction(p) as c:   # se recrea el estado previo: 2 relevadores separados y sus consumos en la misma tarjeta
            keep = c.execute("SELECT id FROM inv_tqtr_componentes WHERE descripcion=?", (svc.RELEVADOR,)).fetchone()["id"]
            c.execute("UPDATE inv_tqtr_componentes SET descripcion=?, cantidad_ensamble=2 WHERE id=?", (svc.RELEVADOR + " | R1", keep))
            r2 = c.execute("INSERT INTO inv_tqtr_componentes (grupo, categoria, descripcion, cantidad_ensamble) VALUES ('R2','x',?,4)", (svc.RELEVADOR + " | R2",)).lastrowid
            for cid, q in ((keep, 2), (r2, 4)):
                c.execute("INSERT INTO inv_tqtr_consumos (tarjeta_id, componente_id, origen, cantidad_estandar, cantidad_real) VALUES (777,?,'placas',?,?)", (cid, q, q))
            c.execute("DELETE FROM inv_tqtr_meta WHERE clave='migr_1464_relevadores_sensores'")
            svc._migrar_1464(c)
            filas = [tuple(r) for r in c.execute("SELECT componente_id, cantidad_estandar, cantidad_real FROM inv_tqtr_consumos WHERE tarjeta_id=777")]
            n = c.execute("SELECT COUNT(*) FROM inv_tqtr_componentes WHERE descripcion LIKE 'Relevador%'").fetchone()[0]
        self.assertEqual(filas, [(keep, 6, 6)])
        self.assertEqual(n, 1)

    def test_06_salida_de_producto_final_descuenta_bom(self):
        p = self.path
        antes = self.comp(svc.dashboard(db_path=p), "Relevador automotriz 12 VCD")["disponible"]
        svc.crear_registro({"modelo": "Quinta sin timer | R1+R2", "movimiento": "Salida", "cantidad": 2}, db_path=p)
        d = svc.dashboard(db_path=p)
        self.assertEqual(self.comp(d, "Relevador automotriz 12 VCD")["disponible"], antes - 12)  # (2 R1 + 4 R2) por ensamble × 2
        self.assertEqual(self.comp(d, "Caja Translock Metal Reforzada")["disponible"], 16)  # no lleva caja Translock

    def test_07_filtros_y_paginacion(self):
        r = svc.listar_registros(limit=10, offset=0, db_path=self.path)
        self.assertEqual((r["total"], len(r["items"])), (52, 10))
        r2 = svc.listar_registros(limit=10, offset=50, db_path=self.path)
        self.assertEqual(len(r2["items"]), 2)
        g = svc.listar_registros(categoria="Gabinete", limit=5000, db_path=self.path)
        self.assertTrue(g["items"] and all(x["categoria"] == "Gabinete" for x in g["items"]))
        f = svc.listar_registros(desde="2026-10-01", hasta="2026-10-31", limit=5000, db_path=self.path)
        self.assertTrue(all("2026-10-01" <= x["fecha_entrada"] <= "2026-10-31" for x in f["items"]))

    # ------------------------------------------------------------------ componentes y catálogos
    def test_08_componentes(self):
        p = self.path
        c = svc.crear_componente({"grupo": "A", "categoria": "Otros", "descripcion": "Fusible 5 A", "cantidad_lote": 10, "cantidad_ensamble": 1,
                                  "bom": {"Quinta completa | R1+R2+R3": 1}}, db_path=p)
        lista = svc.listar_componentes(db_path=p)
        f = next(x for x in lista["items"] if x["id"] == c["id"])
        self.assertEqual((f["rendimiento_lote"], f["bom"]), (10, {"Quinta completa | R1+R2+R3": 1}))
        with self.assertRaises(ConflictoError):
            svc.crear_componente({"descripcion": "Fusible 5 A"}, db_path=p)
        svc.crear_registro({"modelo": "Fusible 5 A", "movimiento": "Entrada", "cantidad": 3}, db_path=p)
        svc.editar_componente(c["id"], {"descripcion": "Fusible 5 A rápido"}, db_path=p)  # renombra también el registro
        self.assertEqual(svc.listar_registros(modelo="Fusible 5 A rápido", db_path=p)["total"], 1)
        with self.assertRaises(ConflictoError):
            svc.eliminar_componente(c["id"], db_path=p)
        svc.eliminar_registro(svc.listar_registros(modelo="Fusible 5 A rápido", db_path=p)["items"][0]["id"], db_path=p)
        svc.eliminar_componente(c["id"], db_path=p)
        self.assertNotIn(c["id"], [x["id"] for x in svc.listar_componentes(db_path=p)["items"]])

    def test_09_catalogos_y_productos(self):
        p = self.path
        cat = svc.crear_catalogo("condicion", "Dañado", db_path=p)
        with self.assertRaises(ConflictoError):
            svc.crear_catalogo("condicion", "Dañado", db_path=p)
        with self.assertRaises(ValueError):
            svc.crear_catalogo("movimiento", "Préstamo", efecto=5, db_path=p)
        svc.eliminar_catalogo(cat["id"], db_path=p)
        entrada = next(x for x in svc.configuracion(db_path=p)["catalogos"]["movimiento"] if x["valor"] == "Entrada")
        with self.assertRaises(ConflictoError):
            svc.eliminar_catalogo(entrada["id"], db_path=p)
        prod = svc.crear_producto("Kit de prueba", db_path=p)
        self.assertIn("Kit de prueba", svc.configuracion(db_path=p)["modelos"])
        svc.eliminar_producto(prod["id"], db_path=p)

    def test_10_export_xlsx(self):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(svc.exportar_xlsx(db_path=self.path)))
        self.assertEqual(wb.sheetnames, ["Dashboard de inventario", "Registro_inventario", "Componentes", "Configuración", "Costos", "Stock mínimo"])
        self.assertEqual(wb["Registro_inventario"].max_row, 3 + 52)

    # ------------------------------------------------------------------ costos
    def test_11_costo_por_armado_igual_al_excel(self):
        cos = svc.costos(db_path=self.path)
        self.assertEqual(cos["parametros"]["producto"], "Translock")
        self.assertEqual(cos["parametros"]["tipo_cambio"], 18.5)
        # G44 del Excel = 4426.284; v1.3.39 cambia el silicón a 400 g / 100 g a 0.226 MXN/g (226 por envase de 1 kg):
        # R1 593.023 − 74.58 + 90.40 = 608.843 y R3 417.201 − 13.56 + 22.60 = 426.241
        self.assertAlmostEqual(cos["armado"]["total"], 4451.144, places=3)
        g = cos["armado"]["por_grupo"]
        for k, v in (("R1", 608.843 + 162.8), ("R2", 694.06 - 162.8), ("R3", 426.241), ("Otros", 2722)):  # v1.4.64: el costo de los 4 relevadores de R2 pasa a la fila unificada (R1)
            self.assertAlmostEqual(g[k], v, places=3, msg=k)
        self.assertAlmostEqual(cos["por_variante"]["Quintalock"]["total"], 4451.144 - 1276 + 1450, places=3)
        self.assertEqual(cos["sin_componente"], ["Relevador automotriz 12 VCD | R3"])  # v1.4.64: la fila R3 «no incluida» queda como referencia
        self.assertEqual(cos["componentes_sin_costo"], [])
        self.assertGreater(cos["valor_existencia"], 0)
        self.assertEqual(cos["armados_posibles"]["cantidad"], 16)  # min(producto final 29, Translock 16)
        self.assertAlmostEqual(cos["armados_posibles"]["costo"], round(16 * 4451.144, 2), places=2)
        self.assertIn("costos", svc.dashboard(db_path=self.path))

    def test_12_crud_costos_y_parametros(self):
        p = self.path
        svc.editar_parametros_costos({"tipo_cambio": 20}, db_path=p)
        self.assertGreater(svc.costos(db_path=p)["armado"]["total"], 4451.144)
        svc.editar_parametros_costos({"producto": "Quintalock", "tipo_cambio": 18.5}, db_path=p)
        self.assertAlmostEqual(svc.costos(db_path=p)["armado"]["total"], 4625.144, places=3)
        with self.assertRaises(ValueError):
            svc.editar_parametros_costos({"producto": "Otro"}, db_path=p)
        n = svc.crear_costo({"descripcion": "Etiqueta", "grupo": "Otros", "costo_unitario": 2, "moneda": "USD", "cantidad_armado": 1}, db_path=p)
        self.assertAlmostEqual(svc.costos(db_path=p)["armado"]["total"], 4625.144 + 37, places=3)
        self.assertIn("Etiqueta", svc.costos(db_path=p)["sin_componente"])
        svc.editar_costo(n["id"], {"incluido": False}, db_path=p)
        self.assertAlmostEqual(svc.costos(db_path=p)["armado"]["total"], 4625.144, places=3)
        with self.assertRaises(ValueError):
            svc.editar_costo(n["id"], {"moneda": "EUR"}, db_path=p)
        svc.eliminar_costo(n["id"], db_path=p)
        with self.assertRaises(NoEncontradoError):
            svc.eliminar_costo(n["id"], db_path=p)


class TestConsumoTarjetas(unittest.TestCase):
    """v1.3.38: al completar una tarjeta se descuentan sus materiales; gabinete y actuador según el gabinete asignado."""
    def setUp(self):
        verificar_aislamiento()
        from app.database import db
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "consumo.db"
        db.init_db(self.path)

    def tearDown(self):
        svc._reiniciar_cache()
        self.tmp.cleanup()

    def disp(self, desc):
        return next(f for f in svc.dashboard(db_path=self.path)["componentes"] if f["descripcion"] == desc)["disponible"]

    def tarjeta(self, r3=True, firmware=True):
        """Tarjeta NUEVA con MAC; `firmware=True` programa R1 y R2 (evento «placas»). Antes se asegura la migración v1.4.6
        (primer sync) para que la tarjeta no caiga en la línea base."""
        from _aislamiento import crear_par
        from app.database import inventario as inv
        svc.sincronizar_consumos(self.path)
        t = crear_par(r3=r3, path=self.path)
        if firmware:
            self.programar(t)
        return t

    def programar(self, t, fw="1.0"):
        from app.database import inventario as inv
        for k in ("pcb_r1_id", "pcb_r2_id"):
            inv.set_firmware(t[k], fw, db_path=self.path)

    def estados(self, tid):
        return {i["descripcion"]: i["estado"] for i in svc.materiales_tarjeta(tid, db_path=self.path)["items"]}

    def test_01_programar_descuenta_placas_y_r3(self):
        """v1.4.6: R1+R2 con MAC y firmware descuentan placas; la R3 asignada, sus materiales."""
        from app.database import inventario as inv
        D = ("PCB TQT-R1 V2.0 | PRINCIPAL", "Cable calibre 22 | R3 | Rojo", "Relevador automotriz 12 VCD", ACT)
        base = {d: self.disp(d) for d in D}
        t = self.tarjeta(firmware=False)  # MAC sin firmware: solo la R3 asignada consume
        self.assertEqual(self.disp("PCB TQT-R1 V2.0 | PRINCIPAL"), base["PCB TQT-R1 V2.0 | PRINCIPAL"])
        self.assertEqual(self.disp("Relevador automotriz 12 VCD"), base["Relevador automotriz 12 VCD"])
        self.assertAlmostEqual(self.disp("Cable calibre 22 | R3 | Rojo"), base["Cable calibre 22 | R3 | Rojo"] - 0.55)
        inv.set_firmware(t["pcb_r1_id"], "1.0", db_path=self.path)  # solo R1 programada: aún no
        self.assertEqual(self.disp("PCB TQT-R1 V2.0 | PRINCIPAL"), base["PCB TQT-R1 V2.0 | PRINCIPAL"])
        inv.set_firmware(t["pcb_r2_id"], "1.0", db_path=self.path)
        self.assertEqual(self.disp("PCB TQT-R1 V2.0 | PRINCIPAL"), base["PCB TQT-R1 V2.0 | PRINCIPAL"] - 1)
        self.assertEqual(self.disp("Relevador automotriz 12 VCD"), base["Relevador automotriz 12 VCD"] - 6)  # v1.4.64: 2 R1 + 4 R2 juntos
        self.assertEqual(self.disp(ACT), base[ACT])  # sin gabinete no se toca el actuador
        f = svc.materiales_tarjeta(t["id"], db_path=self.path)
        self.assertTrue(f["editable"])
        self.assertEqual({i["origen"] for i in f["items"]}, {"placas", "r3"})
        svc.sincronizar_consumos(self.path); svc.sincronizar_consumos(self.path)  # idempotente
        self.assertEqual(self.disp("PCB TQT-R1 V2.0 | PRINCIPAL"), base["PCB TQT-R1 V2.0 | PRINCIPAL"] - 1)
        self.assertAlmostEqual(self.disp("Cable calibre 22 | R3 | Rojo"), base["Cable calibre 22 | R3 | Rojo"] - 0.55)

    def test_02_gabinete_y_cambio(self):
        from app.database import inventario as inv
        t = self.tarjeta()
        a0, q0, t0 = self.disp(ACT), self.disp(CQ), self.disp(CT)
        inv.actualizar_datos_tarjeta(t["id"], {"gabinete": "Translock"}, db_path=self.path)
        self.assertEqual((self.disp(ACT), self.disp(CQ), self.disp(CT)), (a0 - 1, q0, t0 - 1))
        inv.actualizar_datos_tarjeta(t["id"], {"gabinete": "Quintalock"}, db_path=self.path)
        self.assertEqual((self.disp(ACT), self.disp(CQ), self.disp(CT)), (a0 - 1, q0 - 1, t0))
        inv.actualizar_datos_tarjeta(t["id"], {"gabinete": ""}, db_path=self.path)
        self.assertEqual((self.disp(ACT), self.disp(CQ), self.disp(CT)), (a0, q0, t0))

    def test_03_merma_cambia_existencia_y_costo(self):
        t = self.tarjeta()
        f = svc.materiales_tarjeta(t["id"], db_path=self.path)
        cab = next(i for i in f["items"] if i["descripcion"] == "Cable uso rudo 4 hilos calibre 18 | Negro")
        d0 = self.disp(cab["descripcion"])
        f2 = svc.editar_materiales_tarjeta(t["id"], [{"id": cab["id"], "cantidad_real": 1.55, "nota": "corte mal medido"}], db_path=self.path)
        self.assertAlmostEqual(self.disp(cab["descripcion"]), d0 - 1.0)
        self.assertAlmostEqual(f2["totales"]["merma_costo"], 28.0, places=2)  # 1 m extra × 28 MXN
        self.assertGreater(f2["totales"]["costo_real"], f2["totales"]["costo_estandar"])
        with self.assertRaises(ValueError):
            svc.editar_materiales_tarjeta(t["id"], [{"id": cab["id"], "cantidad_real": -1}], db_path=self.path)

    def test_04_desemparejar_y_borrar_devuelven_solo_actuador_y_gabinete(self):
        from app.database import inventario as inv
        base = {d: self.disp(d) for d in (ACT, CT, "PCB TQT-R3 V2.0 | TIMER", "PCB TQT-R1 V2.0 | PRINCIPAL")}
        t = self.tarjeta()
        inv.actualizar_datos_tarjeta(t["id"], {"gabinete": "Translock"}, db_path=self.path)
        self.assertEqual(self.disp(ACT), base[ACT] - 1)  # sincronizado con gabinete
        inv.asignar_pcb(t["id"], "R3", None, db_path=self.path)  # sin R3: sus materiales quedan consumidos; el gabinete sigue puesto
        self.assertEqual((self.disp(ACT), self.disp(CT)), (base[ACT] - 1, base[CT] - 1))
        self.assertEqual(self.disp("PCB TQT-R3 V2.0 | TIMER"), base["PCB TQT-R3 V2.0 | TIMER"] - 1)  # queda consumida (disuelta)
        self.assertTrue(any(i["estado"] == "disuelta" and i["nota"] == svc.NOTA_DISUELTA and i["origen"] == "r3"
                            for i in svc.materiales_tarjeta(t["id"], db_path=self.path)["items"]))
        inv.actualizar_datos_tarjeta(t["id"], {"gabinete": ""}, db_path=self.path)  # sin gabinete vuelve solo lo devolvible
        self.assertEqual((self.disp(ACT), self.disp(CT)), (base[ACT], base[CT]))
        f = svc.materiales_tarjeta(t["id"], db_path=self.path)
        self.assertEqual(f["tarjeta"]["estado"], "activa")  # las placas siguen descontadas
        # desemparejar R1: todo queda disuelto
        inv.asignar_pcb(t["id"], "R1", None, db_path=self.path)
        f = svc.materiales_tarjeta(t["id"], db_path=self.path)
        self.assertEqual(f["tarjeta"]["estado"], "disuelta"); self.assertFalse(f["editable"])
        self.assertEqual(self.disp("PCB TQT-R1 V2.0 | PRINCIPAL"), base["PCB TQT-R1 V2.0 | PRINCIPAL"] - 1)
        # tarjeta borrada (disuelta desde Tarjetas)
        t2 = self.tarjeta()
        inv.actualizar_datos_tarjeta(t2["id"], {"gabinete": "Translock"}, db_path=self.path)
        a1 = self.disp(ACT)
        inv.disolver_tarjeta(t2["id"], forzar=True, db_path=self.path)
        self.assertEqual(self.disp(ACT), a1 + 1)
        estados = {i["tarjeta_id"]: i["estado"] for i in svc.consumos(db_path=self.path)["items"]}
        self.assertEqual((estados[t["id"]], estados[t2["id"]]), ("disuelta", "borrada"))

    def test_05_anteriores_a_consumo_desde_no_cuentan(self):
        from app.database import db
        base = self.disp("PCB TQT-R1 V2.0 | PRINCIPAL")
        t = self.tarjeta()
        with db.transaction(self.path) as c:
            c.execute("UPDATE tarjetas_produccion SET fecha_finalizado='2020-01-01' WHERE id=?", (t["id"],))
        self.assertEqual(self.disp("PCB TQT-R1 V2.0 | PRINCIPAL"), base)
        self.assertEqual(svc.materiales_tarjeta(t["id"], db_path=self.path)["tarjeta"]["estado"], "anterior")
        svc.editar_consumo_desde("2019-12-31", db_path=self.path)
        self.assertEqual(self.disp("PCB TQT-R1 V2.0 | PRINCIPAL"), base - 1)
        with self.assertRaises(ValueError):
            svc.editar_consumo_desde("31/12/2019", db_path=self.path)

    def test_06_programada_y_desarmada_entre_sincronizaciones(self):
        """Por evento: lo que se programa y se deshace sin que nadie sincronice entre medio nunca llegó a descontarse
        (en la app cada cambio sincroniza); si se sincroniza mientras está programada, queda 'disuelta' al desarmar."""
        from app.database import inventario as inv
        base = self.disp("PCB TQT-R2 V2.0 | RESPALDO")
        t = self.tarjeta(r3=False, firmware=False)
        self.programar(t, "2.0"); inv.asignar_pcb(t["id"], "R1", None, db_path=self.path)
        self.assertEqual(self.disp("PCB TQT-R2 V2.0 | RESPALDO"), base)
        t2 = self.tarjeta(r3=False)
        self.assertEqual(self.disp("PCB TQT-R2 V2.0 | RESPALDO"), base - 1)
        inv.asignar_pcb(t2["id"], "R1", None, db_path=self.path)
        self.assertEqual(self.disp("PCB TQT-R2 V2.0 | RESPALDO"), base - 1)
        self.assertEqual(svc.materiales_tarjeta(t2["id"], db_path=self.path)["tarjeta"]["estado"], "disuelta")

    def test_08_borrada_sin_abrir_la_consola(self):
        from app.database import db, inventario as inv
        base = {d: self.disp(d) for d in ("PCB TQT-R1 V2.0 | PRINCIPAL", ACT, CT)}  # prepara tablas y trigger
        t = self.tarjeta()
        inv.actualizar_datos_tarjeta(t["id"], {"gabinete": "Translock"}, db_path=self.path)
        self.assertEqual(self.disp(ACT), base[ACT] - 1)  # gabinete ya consumido
        t2 = self.tarjeta()  # completada y borrada directo en BD, sin ninguna vista de por medio
        with db.transaction(self.path) as c:
            c.execute("DELETE FROM tarjetas_produccion WHERE id IN (?, ?)", (t["id"], t2["id"]))
            self.assertEqual(c.execute("SELECT COUNT(*) FROM inv_tqtr_tarjetas_borradas").fetchone()[0], 2)
        svc.sincronizar_consumos(self.path)
        self.assertEqual(self.disp("PCB TQT-R1 V2.0 | PRINCIPAL"), base["PCB TQT-R1 V2.0 | PRINCIPAL"] - 2)
        self.assertEqual((self.disp(ACT), self.disp(CT)), (base[ACT], base[CT]))  # caja y actuador devueltos
        estados = {i["tarjeta_id"]: i["estado"] for i in svc.consumos(db_path=self.path)["items"]}
        self.assertEqual((estados[t["id"]], estados[t2["id"]]), ("borrada", "borrada"))

    # ------------------------------------------------------------------ v1.4.6: consumo por evento
    def test_10_r3_sin_placas_programadas(self):
        base = {d: self.disp(d) for d in ("PCB TQT-R3 V2.0 | TIMER", "Silicón P-53 | R3", "PCB TQT-R1 V2.0 | PRINCIPAL")}
        t = self.tarjeta(firmware=False)
        self.assertEqual(self.disp("PCB TQT-R3 V2.0 | TIMER"), base["PCB TQT-R3 V2.0 | TIMER"] - 1)
        self.assertEqual(self.disp("Silicón P-53 | R3"), base["Silicón P-53 | R3"] - 100)
        self.assertEqual(self.disp("PCB TQT-R1 V2.0 | PRINCIPAL"), base["PCB TQT-R1 V2.0 | PRINCIPAL"])
        self.assertEqual({i["origen"] for i in svc.materiales_tarjeta(t["id"], db_path=self.path)["items"]}, {"r3"})
        t2 = self.tarjeta(r3=False)  # sin R3 no se descuenta nada de R3
        self.assertEqual(self.disp("PCB TQT-R3 V2.0 | TIMER"), base["PCB TQT-R3 V2.0 | TIMER"] - 1)
        self.assertEqual(svc.materiales_tarjeta(t2["id"], db_path=self.path)["tarjeta"]["estado"], "activa")  # sus placas sí

    def test_11_gabinete_sin_tarjeta_completa_y_cambio(self):
        from app.database import inventario as inv
        t = self.tarjeta(r3=False, firmware=False)  # ni programada ni con R3
        a0, q0, t0 = self.disp(ACT), self.disp(CQ), self.disp(CT)
        inv.actualizar_datos_tarjeta(t["id"], {"gabinete": "Quintalock"}, db_path=self.path)
        self.assertEqual((self.disp(ACT), self.disp(CQ), self.disp(CT)), (a0 - 1, q0 - 1, t0))
        self.assertEqual(self.estados(t["id"])[CQ], "activo")
        inv.actualizar_datos_tarjeta(t["id"], {"gabinete": "Translock"}, db_path=self.path)  # vuelve la Quintalock, sale la Translock
        self.assertEqual((self.disp(ACT), self.disp(CQ), self.disp(CT)), (a0 - 1, q0, t0 - 1))
        est = self.estados(t["id"])
        self.assertEqual((est[CQ], est[CT], est[ACT]), ("devuelto", "activo", "activo"))

    def test_12_quitar_firmware_deja_disuelta_y_reactiva_sin_duplicar(self):
        from app.database import inventario as inv
        pcb = "PCB TQT-R1 V2.0 | PRINCIPAL"
        t = self.tarjeta(r3=False, firmware=False)
        d0 = self.disp(pcb)
        self.programar(t)
        self.assertEqual(self.disp(pcb), d0 - 1)
        n = len(svc.materiales_tarjeta(t["id"], db_path=self.path)["items"])
        inv.set_firmware(t["pcb_r1_id"], "", db_path=self.path)  # se borra el firmware de R1
        self.assertEqual(self.disp(pcb), d0 - 1)  # nada devolvible: sigue descontado
        f = svc.materiales_tarjeta(t["id"], db_path=self.path)
        self.assertEqual({i["estado"] for i in f["items"]}, {"disuelta"})
        self.assertEqual(f["tarjeta"]["estado"], "disuelta")
        self.programar(t, "3.0")  # vuelve a programar: se reactivan las mismas filas
        f = svc.materiales_tarjeta(t["id"], db_path=self.path)
        self.assertEqual((len(f["items"]), {i["estado"] for i in f["items"]}), (n, {"activo"}))
        self.assertEqual(self.disp(pcb), d0 - 1)  # no se descuenta dos veces

    def test_13_linea_base_no_descuenta(self):
        """Lo que ya cumplía un evento antes de la migración v1.4.6 no se descuenta; solo cambia la caja si cambia el gabinete."""
        from app.database import db, inventario as inv
        pcb = "PCB TQT-R1 V2.0 | PRINCIPAL"
        t = self.tarjeta(firmware=False)
        self.programar(t)
        inv.actualizar_datos_tarjeta(t["id"], {"gabinete": "Translock"}, db_path=self.path)
        svc.sincronizar_consumos(self.path)
        with db.transaction(self.path) as c:  # simula una BD anterior a v1.4.6 con la tarjeta ya cumpliendo todo, sin consumos
            c.execute("DELETE FROM inv_tqtr_consumos WHERE tarjeta_id=?", (t["id"],))
            c.execute("DELETE FROM inv_tqtr_consumo_base")
            c.execute("DELETE FROM inv_tqtr_meta WHERE clave='migr_consumo_146'")
        d0 = {d: self.disp(d) for d in (pcb, "PCB TQT-R3 V2.0 | TIMER", ACT, CT, CQ)}
        with db.transaction(self.path) as c:
            svc._sincronizar(c)
            bases = {(r["origen"], r["valor"]) for r in c.execute("SELECT * FROM inv_tqtr_consumo_base WHERE tarjeta_id=?", (t["id"],))}
        self.assertEqual(bases, {("placas", ""), ("r3", ""), ("gabinete", "Translock")})
        self.assertEqual({d: self.disp(d) for d in d0}, d0)  # nada se descuenta
        self.assertEqual(svc.materiales_tarjeta(t["id"], db_path=self.path)["items"], [])
        inv.actualizar_datos_tarjeta(t["id"], {"gabinete": "Quintalock"}, db_path=self.path)  # cambia: solo la caja nueva
        self.assertEqual((self.disp(CT), self.disp(CQ), self.disp(ACT)), (d0[CT], d0[CQ] - 1, d0[ACT]))
        self.assertEqual(self.disp(pcb), d0[pcb])

    def test_09_silicon_en_gabinete(self):
        from app.database import inventario as inv
        s0, c0 = self.disp("Silicón P-53 | R1+R2"), self.disp("Silicón P-53 | R3")
        t = self.tarjeta()
        self.assertEqual((self.disp("Silicón P-53 | R1+R2"), self.disp("Silicón P-53 | R3")), (s0, c0 - 100))  # R3 con las placas
        inv.actualizar_datos_tarjeta(t["id"], {"gabinete": "Quintalock"}, db_path=self.path)
        self.assertEqual(self.disp("Silicón P-53 | R1+R2"), s0 - 400)  # 400 g al montar en gabinete
        self.assertAlmostEqual(self.disp("Catalizador para silicón P-53 | R1+R2"), 294 - 6)

    def test_07_diferencias_entre_excel(self):
        dif = {d["componente"]: d for d in svc.costos(db_path=self.path)["diferencias"]}
        self.assertEqual(dif["Cable calibre 18 | Rojo"]["consumo"], 0.55)
        self.assertNotIn("Catalizador para silicón P-53 | R1+R2", dif)  # v1.3.39: ml en ambos lados (6 ml por armado)


class TestV1339(unittest.TestCase):
    """Unidades del silicón/catalizador, ajuste de sensores, stock mínimo y aviso por correo una sola vez por cruce."""
    def setUp(self):
        verificar_aislamiento()
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "v1339.db"

    def tearDown(self):
        svc._reiniciar_cache()
        self.tmp.cleanup()

    def comp(self, desc):
        return next(f for f in svc.dashboard(db_path=self.path)["componentes"] if f["descripcion"] == desc)

    def test_01_sensores_40_e_idempotente(self):
        self.assertEqual(self.comp("Sensor magnético CS1-U")["disponible"], 73)   # v1.4.64: sin BOM; los 73 recibidos ya no se descuentan aparte
        regs = svc.listar_registros(q="Ajuste v1.3.39", db_path=self.path)["items"]
        self.assertEqual([(r["movimiento"], r["cantidad"]) for r in regs], [("Entrada", 73)])
        svc._reiniciar_cache()
        self.assertEqual(self.comp("Sensor magnético CS1-U")["disponible"], 73)  # la migración no se repite

    def test_02_unidades_y_conversion(self):
        s = self.comp("Silicón P-53 | R1+R2")
        self.assertEqual((s["unidad"], s["disponible"], s["cantidad_ensamble"], s["ensambles"]), ("gramos", 19600, 400, 49))
        c = self.comp("Catalizador para silicón P-53 | R3")
        self.assertEqual((c["unidad"], c["disponible"], c["cantidad_ensamble"]), ("mililitros", 73.5, 1.5))
        regs = svc.listar_registros(modelo="Silicón P-53 | R1+R2", limit=5000, db_path=self.path)["items"]
        self.assertTrue(all(r["unidad"] == "gramos" for r in regs))
        svc.editar_conversiones({"conv_silicon_g_por_envase": 1100}, db_path=self.path)  # factor editable: recalcula desde lo original
        self.assertEqual(self.comp("Silicón P-53 | R1+R2")["disponible"], 24 * 1100 - 4400)
        self.assertAlmostEqual(svc.costos(db_path=self.path)["armado"]["por_grupo"]["R1"], 608.843 + 162.8 - 90.40 + 400 * 226 / 1100, places=3)
        with self.assertRaises(ValueError):
            svc.editar_conversiones({"conv_catalizador_ml_por_envase": 0}, db_path=self.path)

    def test_03_stock_minimo_y_correo_una_vez(self):
        from unittest import mock
        from app.database import db
        svc.dashboard(db_path=self.path)
        with db.transaction(self.path) as c:  # cuentas: un administrador activo y un consultor
            c.execute("CREATE TABLE IF NOT EXISTS usuarios (email TEXT PRIMARY KEY, rol TEXT, activo INTEGER)")
            c.execute("INSERT INTO usuarios VALUES ('jefe@x.mx','administrador',1), ('ver@x.mx','consultor',1)")
        cable = self.comp("Cable calibre 22 | R3 | Rojo")  # 84.5 m
        with mock.patch("app.services.correo.configurado", return_value=True), mock.patch("app.services.correo.enviar") as env:
            r = svc.editar_stock_minimo(cable["id"], {"stock_minimo": 50, "punto_reorden": 60, "entrega_dias": 7, "entrega_nota": "Kuma, pedido semanal",
                                                     "armados_objetivo": 20}, db_path=self.path)
            self.assertFalse(r["bajo_minimo"])
            self.assertEqual(env.call_count, 0)
            svc.crear_registro({"modelo": cable["descripcion"], "movimiento": "Salida", "cantidad": 30}, db_path=self.path)  # 54.5 <= reorden 60
            self.assertEqual(env.call_count, 1)
            para, asunto, texto = env.call_args[0][:3]
            self.assertEqual(para, "jefe@x.mx")
            self.assertIn("Cable calibre 22 | R3 | Rojo", texto)
            self.assertIn("7 días", texto)
            svc.dashboard(db_path=self.path); svc.stock_minimo(db_path=self.path)  # leer no reenvía
            svc.crear_registro({"modelo": cable["descripcion"], "movimiento": "Salida", "cantidad": 1}, db_path=self.path)
            self.assertEqual(env.call_count, 1)
            st = svc.stock_minimo(db_path=self.path)
            x = next(i for i in st["a_comprar"] if i["id"] == cable["id"])
            self.assertAlmostEqual(x["sugerida"], 20 * 0.5 + 60 - 53.5)  # objetivo × por armado + reorden − disponible
            svc.crear_registro({"modelo": cable["descripcion"], "movimiento": "Entrada", "cantidad": 100}, db_path=self.path)  # se recupera
            svc.crear_registro({"modelo": cable["descripcion"], "movimiento": "Salida", "cantidad": 120}, db_path=self.path)  # nuevo cruce
            self.assertEqual(env.call_count, 2)
        with mock.patch("app.services.correo.configurado", return_value=True), mock.patch("app.services.correo.enviar", side_effect=OSError("smtp caído")):
            svc.editar_stock_minimo(self.comp("Cinchos de nylon")["id"], {"stock_minimo": 10000}, db_path=self.path)  # falla en silencio
        self.assertGreaterEqual(svc.dashboard(db_path=self.path)["stock"]["a_comprar"], 2)


ACT, CQ, CT = "Actuador lineal de 12 VCD", "Caja Quintalock Metal Reforzada", "Caja Translock Metal Reforzada"


class TestInventarioTQTRApi(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from starlette.testclient import TestClient
        from app.main import app
        cls.client = TestClient(app)

    def test_endpoints(self):
        c = self.client
        r = c.get("/api/inventario-tqtr/dashboard"); self.assertEqual(r.status_code, 200, r.text)
        self.assertIn("pcba_posibles", r.json())
        self.assertEqual(c.get("/api/inventario-tqtr/registros?limit=5").json()["limit"], 5)
        self.assertEqual(c.get("/api/inventario-tqtr/componentes").status_code, 200)
        self.assertEqual(c.get("/api/inventario-tqtr/configuracion").status_code, 200)
        r = c.post("/api/inventario-tqtr/registros", json={"modelo": "Cinchos de nylon", "movimiento": "Entrada", "cantidad": 5})
        self.assertEqual(r.status_code, 201, r.text)
        mid = r.json()["id"]
        self.assertEqual(c.patch(f"/api/inventario-tqtr/registros/{mid}", json={"notas": "ok"}).json()["notas"], "ok")
        self.assertEqual(c.post("/api/inventario-tqtr/registros", json={"modelo": "X", "movimiento": "Nada", "cantidad": 1}).status_code, 400)
        self.assertEqual(c.delete(f"/api/inventario-tqtr/registros/{mid}").status_code, 204)
        self.assertEqual(c.delete(f"/api/inventario-tqtr/registros/{mid}").status_code, 404)
        r = c.get("/api/inventario-tqtr/costos"); self.assertEqual(r.status_code, 200, r.text)
        self.assertAlmostEqual(r.json()["armado"]["total"], 4451.144, places=3)
        self.assertIsInstance(r.json()["producido"]["tarjetas"], int)
        self.assertEqual(c.patch("/api/inventario-tqtr/costos/parametros", json={"producto": "X"}).status_code, 400)
        self.assertEqual(c.get("/api/inventario-tqtr/consumos").status_code, 200)
        self.assertEqual(c.get("/api/inventario-tqtr/tarjetas/999999/materiales").status_code, 404)
        r = c.get("/api/inventario-tqtr/export")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b"PK"))


if __name__ == "__main__":
    unittest.main()
