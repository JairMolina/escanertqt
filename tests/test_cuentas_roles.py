"""v1.3.35: alta de cuentas por invitación de correo, roles (administrador / general / consultor) y reporte del día."""
import _aislamiento  # noqa: F401  (debe ser lo primero)
import unittest
from unittest import mock

from openpyxl import load_workbook
import io

from app.config import settings
from app.database import db
from app.main import app
from app.routers import admin as admin_router
from app.services import admin_auth, correo, reporte_dia, usuarios

from _aislamiento import cliente_sin_sesion, verificar_aislamiento

ADMIN = "developer@skyguardian.mx"
CLAVE = "ClaveDelInvitado-2026"


class TestCuentas(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        verificar_aislamiento()
        db.init_db()
        usuarios.sembrar(iteraciones=1000)
        app.dependency_overrides[admin_router.admin_requerido] = lambda: {"ok": True}

    @classmethod
    def tearDownClass(cls):
        app.dependency_overrides.pop(admin_router.admin_requerido, None)

    def setUp(self):
        usuarios.limitador.reiniciar()
        admin_auth.limitador.reiniciar()
        p = mock.patch.object(settings, "ADMIN_RETARDO_MS", 0); p.start(); self.addCleanup(p.stop)
        self.adm = cliente_sin_sesion(app)
        self.adm.cookies.set(usuarios.COOKIE, usuarios.emitir_sesion(ADMIN))

    def invitar(self, email, rol):
        with mock.patch.object(correo, "enviar", return_value="<id>") as env:
            r = self.adm.post("/api/admin/usuarios", json={"email": email, "rol": rol}, headers={"host": "tqt.test", "x-forwarded-proto": "https"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(r.json()["enviado"])
        texto = env.call_args[0][2]
        enlace = next(x for x in texto.split() if x.startswith("https://tqt.test/invitacion#"))
        return enlace.split("#", 1)[1]

    def test_01_invitacion_y_rol_consultor(self):
        token = self.invitar("consulta1@ejemplo.com", "consultor")
        u = next(x for x in usuarios.listar() if x["email"] == "consulta1@ejemplo.com")
        self.assertEqual((u["rol"], u["activo"], u["pendiente"]), ("consultor", 0, 1))
        anon = cliente_sin_sesion(app)
        self.assertEqual(anon.post("/api/auth/login", json={"email": "consulta1@ejemplo.com", "password": CLAVE}).status_code, 401)
        self.assertEqual(anon.post("/api/auth/invitacion", json={"token": token}).json()["email"], "consulta1@ejemplo.com")
        self.assertEqual(anon.post("/api/auth/invitacion/aceptar", json={"token": token, "nueva": "corta"}).status_code, 422)
        r = anon.post("/api/auth/invitacion/aceptar", json={"token": token, "nueva": CLAVE})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(anon.post("/api/auth/invitacion/aceptar", json={"token": token, "nueva": CLAVE}).status_code, 404)   # un solo uso
        # consultor: lee, pero no escribe ni entra a otras páginas
        self.assertEqual(anon.get("/api/lotes").status_code, 200)
        self.assertEqual(anon.post("/api/lotes", json={"mes": 1, "anio": 2030}).status_code, 403)
        self.assertEqual(anon.get("/consultar", follow_redirects=False).status_code, 200)
        self.assertEqual(anon.get("/monitor", follow_redirects=False).status_code, 200)   # v1.3.44: consola de escritorio (solo lectura)
        self.assertEqual(anon.get("/api/admin/usuarios").status_code, 403)
        # v1.3.42: cada rechazo explica el motivo (el frontend muestra el aviso "Acceso restringido")
        r = anon.post("/api/lotes", json={"mes": 1, "anio": 2030})
        self.assertEqual((r.headers.get("x-acceso"), r.json()["zona"]), ("rol", "accion"))
        self.assertIn("consulta", r.json()["detail"])
        r = anon.get("/", follow_redirects=False)   # v1.3.44: su inicio es Consultar; en PC/tablet consultar.html sigue a /monitor
        self.assertEqual((r.status_code, r.headers["location"]), (302, "/consultar?inicio=1"))
        for ruta, zona in (("/static/emparejar.html", "emparejar"), ("/programar", "programar"), ("/dymo", "dymo"), ("/admin", "admin"), ("/static/admin.html", "admin")):
            r = anon.get(ruta, follow_redirects=False)
            self.assertEqual((r.status_code, r.headers["location"]), (302, f"/consultar?denegado={zona}"), ruta)
        self.assertEqual(anon.get("/static/js/common.js").status_code, 200)
        self.assertEqual(anon.get("/api/auth/yo").json()["rol"], "consultor")
        self.assertEqual(anon.post("/api/auth/cambiar-clave", json={"actual": "x", "nueva": "y"}).status_code != 403, True)

    def test_02_general_sin_administracion(self):
        token = self.invitar("general1@ejemplo.com", "general")
        anon = cliente_sin_sesion(app)
        self.assertEqual(anon.post("/api/auth/invitacion/aceptar", json={"token": token, "nueva": CLAVE}).status_code, 200)
        self.assertEqual(anon.get("/monitor", follow_redirects=False).status_code, 200)
        self.assertEqual(anon.get("/api/admin/usuarios").status_code, 403)
        r = anon.get("/admin", follow_redirects=False)
        self.assertEqual((r.status_code, r.headers["location"]), (302, "/?denegado=admin"))
        r = anon.post("/api/admin/login", json={"password": "x"})
        self.assertEqual((r.status_code, r.headers.get("x-acceso"), r.json()["zona"]), (403, "rol", "admin"))
        self.assertIn("General", r.json()["detail"])
        self.assertEqual(anon.post("/api/lotes", json={"mes": 1, "anio": 2030}).status_code != 403, True)   # general sí escribe

    def test_03_correo_repetido_rol_invalido_y_smtp_caido(self):
        self.assertEqual(self.adm.post("/api/admin/usuarios", json={"email": ADMIN, "rol": "general"}).status_code, 422)
        self.assertEqual(self.adm.post("/api/admin/usuarios", json={"email": "x@ejemplo.com", "rol": "jefe"}).status_code, 422)
        with mock.patch.object(correo, "enviar", side_effect=OSError("sin red")):
            r = self.adm.post("/api/admin/usuarios", json={"email": "sinred@ejemplo.com", "rol": "general"})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()["enviado"])
        self.assertIn("/invitacion#", r.json()["enlace"])

    def test_04_reenviar_invalida_el_anterior_y_caduca(self):
        viejo = self.invitar("reenvio@ejemplo.com", "general")
        with mock.patch.object(correo, "enviar", return_value="<id>"):
            self.assertEqual(self.adm.post("/api/admin/usuarios/reenvio@ejemplo.com/reenviar").status_code, 200)
        self.assertRaises(PermissionError, usuarios.ver_invitacion, viejo)
        nuevo = usuarios.reenviar_invitacion("reenvio@ejemplo.com", ahora=0)
        self.assertRaises(PermissionError, usuarios.ver_invitacion, nuevo)   # emitido en 1970: caducado

    def test_05_siempre_queda_un_administrador(self):
        with db.transaction() as c:
            c.execute("UPDATE usuarios SET activo = 0 WHERE rol = 'administrador' AND email != ?", (ADMIN,))
        try:
            self.assertEqual(self.adm.patch(f"/api/admin/usuarios/{ADMIN}", json={"rol": "general"}).status_code, 422)
            self.assertEqual(self.adm.delete(f"/api/admin/usuarios/{ADMIN}").status_code, 422)
        finally:
            with db.transaction() as c:
                c.execute("UPDATE usuarios SET activo = 1 WHERE inv_hash IS NULL")

    def test_06_reporte_del_dia_y_excel(self):
        with db.transaction() as c:
            c.execute("INSERT OR IGNORE INTO lotes_mensuales (codigo_lote, mes, anio) VALUES ('2031-03', 3, 2031)")
            lid = c.execute("SELECT id FROM lotes_mensuales WHERE codigo_lote = '2031-03'").fetchone()[0]
            for num, fin, real in (("9001", "2031-03-10", None), ("9002", "2031-03-10", "2031-03-10"), ("9003", "2031-03-09", "2031-03-10")):
                c.execute("INSERT INTO tarjetas_produccion (lote_id, id_tarjeta_num) VALUES (?, ?)", (lid, num))
                c.execute("UPDATE tarjetas_produccion SET fecha_finalizado = ?, fecha_real = ? WHERE lote_id = ? AND id_tarjeta_num = ?", (fin, real, lid, num))
        r = self.adm.get("/api/reporte-dia?fecha=2031-03-10").json()
        self.assertEqual((r["completadas"], r["entregadas"], len(r["items"])), (2, 2, 3))
        self.assertEqual(self.adm.get("/api/reporte-dia?fecha=10/03/2031").status_code, 400)
        x = self.adm.get("/api/reporte-dia/excel?fecha=2031-03-10")
        self.assertEqual(x.status_code, 200)
        ws = load_workbook(io.BytesIO(x.content)).active
        self.assertEqual(ws["A8"].value, "Tarjeta")
        self.assertEqual(ws["B5"].value, '=COUNTIF(H9:H11,"Sí")')
        self.assertEqual(ws["A9"].value, "#9001")
        vacio = load_workbook(io.BytesIO(reporte_dia.excel_del_dia("2031-01-01"))).active
        self.assertIn("No hubo", vacio["A9"].value)

    def test_07_reporte_por_rango(self):
        r = self.adm.get("/api/reporte-dia?desde=2031-03-08&hasta=2031-03-10").json()
        self.assertEqual((r["completadas"], r["entregadas"], r["dias"]), (3, 2, 3))
        self.assertEqual(self.adm.get("/api/reporte-dia?desde=2031-03-10&hasta=2031-03-01").status_code, 400)
        self.assertEqual(self.adm.get("/api/reporte-dia?desde=2030-01-01&hasta=2031-03-01").status_code, 400)
        self.assertEqual(self.adm.get("/api/reporte-dia").status_code, 400)
        x = self.adm.get("/api/reporte-dia/excel?desde=2031-03-08&hasta=2031-03-10")
        self.assertIn("Tarjetas_2031-03-08_a_2031-03-10.xlsx", x.headers["content-disposition"])
        wb = load_workbook(io.BytesIO(x.content))
        self.assertEqual(wb.sheetnames, ["Reporte del periodo", "Por día"])
        self.assertEqual(wb["Por día"]["C7"].value, "=SUM(C4:C6)")
        self.assertEqual(reporte_dia.texto_rango("2031-10-01", "2031-10-06"), "del 1 al 6 de octubre de 2031")

    def test_08_plantilla_de_correo_escapa_y_tiene_marca(self):
        html = correo.plantilla("Título <x>", ["Hola <b>"], kpis=[("Completadas", 3, "ok")], adjunto="a.xlsx", boton=("Abrir", "https://x/?a=1&b=2"))
        self.assertIn("Escáner TQT", html)
        self.assertIn("Título &lt;x&gt;", html)
        self.assertNotIn("<b>", html.split("<body")[1].replace("<b style", ""))
        self.assertIn("a=1&amp;b=2", html)


if __name__ == "__main__":
    unittest.main()
