"""Firewall de fórmulas: tras sincronizar, TODAS las fórmulas críticas siguen intactas (celda por celda).

  * INDEX/MATCH   : Producción C y F (MAC por catálogo)
  * COUNTIF       : Catálogo PCB D e I ('SUELTA', 'DUPLICADA', '00xx')
  * TEXT/VALUE    : Producción B
  * LOOKUP(2,1/…) : Pruebas B:F (leen la bitácora masiva) y G (estado general)
  * Etiquetas     : E3 y la cola de impresión
"""
import _aislamiento  # noqa: F401  (debe ser lo primero)
import tempfile
import unittest
from pathlib import Path

import openpyxl

from app.database import db
from app.services.excel_sync import EXCEL_ERROR_CODES, ExcelSyncEngine, get_sheet_by_name

from app.database import inventario as inv

from _aislamiento import crear_par, verificar_aislamiento


def _texto(v):
    return v.text if hasattr(v, "text") else str(v or "")


class TestExcelSyncFormulasFirewall(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        verificar_aislamiento()
        cls.tmp = tempfile.TemporaryDirectory()
        cls.path = Path(cls.tmp.name) / "fw.db"
        db.init_db(cls.path)
        cls.engine = ExcelSyncEngine(db_path=cls.path)
        cls.excel = cls.engine.create_monthly_excel(9, 2026, target_path=str(Path(cls.tmp.name) / "Firewall.xlsx"))
        cls.lote = db.create_lote("2026-09-FW", 9, 2026, ruta_excel=cls.excel, activo=True, db_path=cls.path)

        t = crear_par(cls.lote["id"], num="0011", path=cls.path)
        inv.actualizar_datos_tarjeta(t["id"], {"firmware_r1": "2", "semana_produccion": 38}, db_path=cls.path)
        for etapa in ("soldadura", "programacion", "prueba_pcb"):
            db.set_prueba(t["id"], etapa, "OK", db_path=cls.path)
        db.set_prueba(t["id"], "integracion", "FALLA", db_path=cls.path)

        cls.original = openpyxl.load_workbook(cls.engine.get_template_path(), data_only=False)
        cls.engine.export_to_excel(cls.excel, cls.lote["id"], db_path=cls.path)
        cls.wb = openpyxl.load_workbook(cls.excel, data_only=False)

    @classmethod
    def tearDownClass(cls):
        cls.wb.close()
        cls.original.close()
        cls.tmp.cleanup()

    def _comparar(self, hoja, celdas):
        antes, despues = get_sheet_by_name(self.original, hoja), get_sheet_by_name(self.wb, hoja)
        for ref in celdas:
            self.assertEqual(_texto(despues[ref].value), _texto(antes[ref].value), f"{hoja}!{ref}")

    def test_01_index_match_produccion(self):
        ws = get_sheet_by_name(self.wb, "Producción")
        for fila in range(2, 102):
            for col in (3, 6):
                v = _texto(ws.cell(fila, col).value)
                self.assertTrue(v.startswith("="), f"Producción!{fila},{col}")
                self.assertIn("INDEX", v)
                self.assertIn("Catálogo PCB", v)

    def test_02_countif_catalogo(self):
        ws = get_sheet_by_name(self.wb, "Catálogo PCB")
        for fila in range(3, 103):
            for col in (4, 9):
                v = _texto(ws.cell(fila, col).value)
                self.assertTrue(v.startswith("=") and "COUNTIF" in v and "SUELTA" in v and "DUPLICADA" in v, f"Catálogo!{fila},{col}")

    def test_03_nombre_text_value(self):
        ws = get_sheet_by_name(self.wb, "Producción")
        for fila in range(2, 102):
            v = _texto(ws.cell(fila, 2).value)
            self.assertTrue(v.startswith("=") and "TQT-R1-V30-" in v and "TEXT" in v and "VALUE" in v)

    def test_04_pruebas_formulas_no_se_sobrescriben(self):
        """Antes el sync pisaba Pruebas B:F con texto fijo, anulando la bitácora masiva. Ahora B:G son fórmulas idénticas."""
        ws = get_sheet_by_name(self.wb, "Pruebas")
        for fila in range(2, 102):
            for col in range(1, 8):
                v = _texto(ws.cell(fila, col).value)
                self.assertTrue(v.startswith("="), f"Pruebas fila {fila} col {col} perdió su fórmula: {v!r}")
        self._comparar("Pruebas", [f"{c}{f}" for c in "ABCDEFG" for f in (2, 3, 50, 101)])

    def test_05_bitacora_masiva_recibe_los_estados(self):
        ws = get_sheet_by_name(self.wb, "Pruebas")
        filas = [[ws.cell(r, c).value for c in (11, 12, 13)] for r in range(3, 10) if ws.cell(r, 11).value]
        self.assertEqual(sorted(filas), sorted([["Soldadura", "OK", "0011"], ["Programación", "OK", "0011"],
                                                ["Prueba PCB", "OK", "0011"], ["Integración", "FALLA", "0011"]]))

    def test_06_etiquetas_intactas(self):
        self._comparar("Etiquetas", ["E3", "B6", "D6", "G3", "G4", "G5", "A17", "H17", "K17", "L17", "B116", "H116"])

    def test_07_datos_en_columnas_de_captura(self):
        prod = get_sheet_by_name(self.wb, "Producción")
        cat = get_sheet_by_name(self.wb, "Catálogo PCB")
        self.assertEqual(prod["A2"].value, "0011")
        self.assertEqual(prod["E2"].value, "TQT-R2-V30-0011")
        self.assertEqual(prod["D2"].value, "2")
        self.assertEqual(prod["I2"].value, "2026-W38")           # formato de la lista desplegable de la plantilla
        self.assertRegex(cat["B3"].value, r"^([0-9A-F]{2}:){5}[0-9A-F]{2}$")
        self.assertEqual(cat["C3"].value, "FUNCIONAL")             # Prueba PCB OK -> FUNCIONAL (lista de la plantilla)
        self.assertRegex(cat["G3"].value, r"^([0-9A-F]{2}:){5}[0-9A-F]{2}$")

    def test_08_cero_errores_de_excel(self):
        for hoja in self.wb.sheetnames:
            for fila in self.wb[hoja].iter_rows():
                for celda in fila:
                    txt = _texto(celda.value)
                    for err in EXCEL_ERROR_CODES:
                        self.assertNotIn(err, txt, f"{hoja}!{celda.coordinate}")


if __name__ == "__main__":
    unittest.main()
