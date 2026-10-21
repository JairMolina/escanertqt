import _aislamiento  # noqa: F401  (debe ser lo primero)
"""v1.4.27: placas TPMS. Las V1 traen la plantilla de la R3 Timer con guiones ('PCB_TQT-R3_V2_0_TIMER-0004'):
TPMS las registra como TPMS-V1-<serie> y Escáner TQT las rechaza; la R3 Timer real (guiones bajos) sigue siendo TQT."""
import unittest

from app.database import db
from app.database.models import parse_tarjeta_code, parse_tpms
from app.main import app
from app.services import usuarios

from _aislamiento import cliente_sin_sesion, verificar_aislamiento


class TestTpms(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        verificar_aislamiento()
        db.init_db()
        usuarios.sembrar(iteraciones=1000)
        usuarios.limitador.reiniciar()
        cls.c = cliente_sin_sesion(app)
        r = cls.c.post("/api/auth/login", json={"email": "developer@skyguardian.mx", "password": usuarios.CLAVE_INICIAL})
        assert r.status_code == 200, r.text

    def test_formatos(self):
        self.assertEqual(parse_tpms("PCB_TQT-R3_V2_0_TIMER-0004")["nombre"], "TPMS-V1-0004")
        self.assertEqual(parse_tpms("TPMS_V1_12")["nombre"], "TPMS-V1-0012")
        self.assertEqual(parse_tpms("pcb tpms 7")["nombre"], "TPMS-V1-0007")
        self.assertEqual(parse_tpms("TPMS-V1_1-0003")["nombre"], "TPMS-V1.1-0003")
        self.assertIsNone(parse_tpms("TPMS_V0_0001"))
        self.assertIsNone(parse_tpms("PCB_TQT_R3_V2_0_TIMER_0073"))             # R3 Timer real
        self.assertEqual(parse_tarjeta_code("PCB_TQT_R3_V2_0_TIMER_0073")["nombre"], "TQT-R3-V20-0073")
        self.assertIsNone(parse_tarjeta_code("PCB_TQT-R3_V2_0_TIMER-0004"))     # TPMS: no es TQT

    def test_alta_y_rechazo_en_tqt(self):
        r = self.c.post("/api/tpms/escanear", json={"codigo": "PCB_TQT-R3_V2_0_TIMER-0004", "lote": "Marzo 2026"}).json()
        self.assertEqual((r["resultado"], r["pcb"]["nombre"], r["pcb"]["lote"]), ("AGREGADA", "TPMS-V1-0004", "Marzo 2026"))
        self.assertEqual(r["pcb"]["operador"], "developer@skyguardian.mx")
        self.assertEqual(self.c.post("/api/tpms/escanear", json={"codigo": "TPMS_V1_0004"}).json()["resultado"], "REPETIDA")
        self.assertEqual(self.c.post("/api/tpms/escanear", json={"codigo": "TQT-R1-V30-0001"}).json()["resultado"], "INVALIDA")
        self.assertIn("TPMS-V1-0004", [p["nombre"] for p in self.c.get("/api/tpms/pcb?lote=Marzo 2026").json()["items"]])
        e = self.c.post("/api/pcb/escanear", json={"codigo": "PCB_TQT-R3_V2_0_TIMER-0004"}).json()
        self.assertEqual(e["resultado"], "INVALIDA")
        self.assertIn("TPMS", e["mensaje"])

    # ---------------- v1.4.28 ----------------
    def _manual(self, serie, version="1", lote="L-prueba"):
        r = self.c.post("/api/tpms/pcb/manual", json={"version": version, "serie": serie, "lote": lote})
        self.assertEqual(r.status_code, 201, r.text)
        return r.json()["pcb"]

    def test_manual_y_por_codigo(self):
        p = self._manual("101")
        self.assertEqual((p["nombre"], p["estado"], p["nota"], p["firmware"]), ("TPMS-V1-0101", "RECIBIDA", "", None))
        self.assertEqual(self.c.post("/api/tpms/pcb/manual", json={"version": "V1", "serie": "0101"}).status_code, 409)
        self.assertEqual(self.c.post("/api/tpms/pcb/manual", json={"version": "x", "serie": "1"}).status_code, 422)
        self.assertEqual(self.c.post("/api/tpms/pcb/manual", json={"version": "1", "serie": "0"}).status_code, 422)
        self.assertEqual(self._manual("5", version="1.1")["nombre"], "TPMS-V1.1-0005")
        for codigo in ("101", "TPMS-V1-0101", "TPMS_V1_0101", "PCB_TQT-R3_V2_0_TIMER-0101"):
            r = self.c.get("/api/tpms/pcb/por-codigo", params={"codigo": codigo})
            self.assertEqual(r.status_code, 200, codigo)
            self.assertEqual(r.json()["pcb"]["id"], p["id"])
        self.assertEqual(r.json()["eventos"][0]["accion"], "RECIBIDA")
        self.assertEqual(r.json()["eventos"][0]["operador"], "developer@skyguardian.mx")
        self.assertEqual(self.c.get("/api/tpms/pcb/por-codigo", params={"codigo": "9999"}).status_code, 404)
        self.assertEqual(self.c.get("/api/tpms/pcb/por-codigo", params={"codigo": "basura"}).status_code, 404)

    def test_firmware_y_filtros(self):
        p = self._manual("201", lote="L-fw")
        r = self.c.put(f"/api/tpms/pcb/{p['id']}/firmware", json={"firmware": "fw-2.0.1"}).json()["pcb"]
        self.assertEqual((r["firmware"], r["estado"]), ("fw-2.0.1", "PROGRAMADA"))
        self.assertIsNotNone(r["programada_en"])
        lista = self.c.get("/api/tpms/pcb", params={"firmware": "fw-2.0.1"}).json()
        self.assertEqual([x["id"] for x in lista["items"]], [p["id"]])
        self.assertIn(p["id"], [x["id"] for x in self.c.get("/api/tpms/pcb", params={"q": "fw-2.0"}).json()["items"]])
        self.assertIn(p["id"], [x["id"] for x in self.c.get("/api/tpms/pcb", params={"estado": "PROGRAMADA"}).json()["items"]])
        r = self.c.put(f"/api/tpms/pcb/{p['id']}/firmware", json={"firmware": None}).json()["pcb"]
        self.assertEqual((r["firmware"], r["estado"], r["programada_en"]), (None, "RECIBIDA", None))
        self.assertIn(p["id"], [x["id"] for x in self.c.get("/api/tpms/pcb", params={"firmware": "__sin"}).json()["items"]])
        self.assertEqual(self.c.put("/api/tpms/pcb/999999/firmware", json={"firmware": "x"}).status_code, 404)
        acciones = [e["accion"] for e in self.c.get("/api/tpms/pcb/por-codigo", params={"codigo": "201"}).json()["eventos"]]
        self.assertEqual(acciones, ["EDITADA", "PROGRAMADA", "RECIBIDA"])

    def test_patch_y_delete(self):
        a, b = self._manual("301"), self._manual("302")
        r = self.c.patch(f"/api/tpms/pcb/{a['id']}", json={"serie": "302"})
        self.assertEqual(r.status_code, 409, r.text)
        self.assertEqual(self.c.patch(f"/api/tpms/pcb/{a['id']}", json={"estado": "ROTA"}).status_code, 422)
        r = self.c.patch(f"/api/tpms/pcb/{a['id']}", json={"version": "2", "nota": "revisar", "estado": "falla"}).json()["pcb"]
        self.assertEqual((r["nombre"], r["version"], r["nota"], r["estado"]), ("TPMS-V2-0301", "2", "revisar", "FALLA"))
        self.assertEqual(self.c.delete(f"/api/tpms/pcb/{b['id']}").json(), {"ok": True})
        self.assertEqual(self.c.delete(f"/api/tpms/pcb/{b['id']}").status_code, 404)
        with db.get_db() as c:
            self.assertEqual(c.execute("SELECT accion FROM tpms_eventos WHERE pcb_id = ? ORDER BY id DESC LIMIT 1",
                                       (b["id"],)).fetchone()[0], "ELIMINADA")

    def test_masa_todo_o_nada(self):
        ids = [self._manual(s, lote="L-masa")["id"] for s in ("401", "402", "403")]
        r = self.c.post("/api/tpms/pcb/masa", json={"ids": ids, "accion": "firmware", "valor": "fw-9"})
        self.assertEqual(r.json(), {"afectadas": 3})
        estados = {x["estado"] for x in self.c.get("/api/tpms/pcb", params={"lote": "L-masa"}).json()["items"]}
        self.assertEqual(estados, {"PROGRAMADA"})
        # Un ID inexistente: no cambia ninguna
        r = self.c.post("/api/tpms/pcb/masa", json={"ids": ids + [999999], "accion": "lote", "valor": "L-otro"})
        self.assertEqual(r.status_code, 404)
        self.assertEqual(self.c.get("/api/tpms/pcb", params={"lote": "L-masa"}).json()["total"], 3)
        # Versión que choca con una existente (TPMS-V3-0401): ninguna cambia
        self._manual("401", version="3")
        r = self.c.post("/api/tpms/pcb/masa", json={"ids": ids, "accion": "version", "valor": "3"})
        self.assertEqual(r.status_code, 409)
        versiones = {x["version"] for x in self.c.get("/api/tpms/pcb", params={"lote": "L-masa"}).json()["items"]}
        self.assertEqual(versiones, {"1"})
        self.assertEqual(self.c.post("/api/tpms/pcb/masa", json={"ids": ids, "accion": "estado", "valor": "X"}).status_code, 422)
        self.assertEqual(self.c.post("/api/tpms/pcb/masa", json={"ids": ids, "accion": "estado", "valor": "BAJA"}).json()["afectadas"], 3)
        self.assertEqual(self.c.post("/api/tpms/pcb/masa", json={"ids": ids, "accion": "eliminar"}).json()["afectadas"], 3)
        self.assertEqual(self.c.get("/api/tpms/pcb", params={"lote": "L-masa"}).json()["total"], 0)

    def test_resumen(self):
        self._manual("501", lote="L-resumen")
        r = self.c.get("/api/tpms/resumen").json()
        self.assertEqual(set(r["por_estado"]), {"RECIBIDA", "PROGRAMADA", "FALLA", "BAJA"})
        self.assertEqual(sum(r["por_estado"].values()), r["total"])
        self.assertIn({"lote": "L-resumen", "n": 1}, r["lotes"])
        self.assertGreaterEqual(r["hoy"], 1)
        self.assertTrue(any(v["version"] == "1" for v in r["versiones"]))

    # ---------------- v1.4.28b ----------------
    def test_resumen_actividad_y_ultimos(self):
        p = self._manual("801", lote="L-act")
        self.c.put(f"/api/tpms/pcb/{p['id']}/firmware", json={"firmware": "fw-a"})
        r = self.c.get("/api/tpms/resumen").json()
        self.assertEqual(len(r["actividad"]), 14)
        dias = [d["dia"] for d in r["actividad"]]
        self.assertEqual(dias, sorted(dias))
        self.assertGreaterEqual(r["actividad"][-1]["recibidas"], 1)
        self.assertGreaterEqual(r["actividad"][-1]["programadas"], 1)
        self.assertLessEqual(len(r["ultimos"]), 8)
        self.assertEqual(set(r["ultimos"][0]), {"id", "pcb_id", "nombre", "accion", "detalle", "operador", "creado_en"})
        self.assertEqual((r["ultimos"][0]["nombre"], r["ultimos"][0]["accion"]), ("TPMS-V1-0801", "PROGRAMADA"))

    def test_eventos_filtrados(self):
        p = self._manual("811", lote="L-ev")
        self.c.patch(f"/api/tpms/pcb/{p['id']}", json={"nota": "nota-ev"})
        r = self.c.get("/api/tpms/eventos", params={"q": "TPMS-V1-0811"}).json()
        self.assertEqual(r["total"], 2)
        self.assertEqual([e["accion"] for e in r["items"]], ["EDITADA", "RECIBIDA"])
        r = self.c.get("/api/tpms/eventos", params={"q": "TPMS-V1-0811", "accion": "RECIBIDA"}).json()
        self.assertEqual(r["total"], 1)
        r = self.c.get("/api/tpms/eventos", params={"q": "TPMS-V1-0811", "limit": 1, "offset": 1}).json()
        self.assertEqual((r["total"], [e["accion"] for e in r["items"]]), (2, ["RECIBIDA"]))
        self.assertEqual(self.c.get("/api/tpms/eventos", params={"q": "TPMS-V1-0811", "hasta": "2000-01-01"}).json()["total"], 0)
        self.assertEqual(self.c.get("/api/tpms/eventos", params={"desde": "ayer"}).status_code, 422)

    def test_lotes_y_renombrar(self):
        a, _ = self._manual("821", lote="L-ren"), self._manual("822", lote="L-ren")
        self.c.put(f"/api/tpms/pcb/{a['id']}/firmware", json={"firmware": "fw-l"})
        lote = next(x for x in self.c.get("/api/tpms/lotes").json()["items"] if x["lote"] == "L-ren")
        self.assertEqual((lote["placas"], lote["recibidas"], lote["programadas"], lote["falla"], lote["baja"]), (2, 1, 1, 0, 0))
        self.assertTrue(lote["primera"] and lote["ultima"])
        self.assertEqual(self.c.post("/api/tpms/lotes/renombrar", json={"de": "L-ren", "a": "L-ren"}).status_code, 422)
        self.assertEqual(self.c.post("/api/tpms/lotes/renombrar", json={"de": "L-ren", "a": "L-ren2"}).json(), {"afectadas": 2})
        nombres = {x["lote"] for x in self.c.get("/api/tpms/lotes").json()["items"]}
        self.assertIn("L-ren2", nombres)
        self.assertNotIn("L-ren", nombres)
        ev = self.c.get("/api/tpms/eventos", params={"q": "L-ren -> L-ren2", "accion": "EDITADA"}).json()
        self.assertEqual(ev["total"], 2)

    def test_exportar_xlsx(self):
        import io
        from openpyxl import load_workbook
        self._manual("831", lote="L-xlsx")
        self._manual("832", lote="L-xlsx")
        r = self.c.get("/api/tpms/exportar", params={"lote": "L-xlsx"})
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r.headers["content-type"])
        self.assertRegex(r.headers["content-disposition"], r'attachment; filename="TPMS_placas_\d{8}\.xlsx"')
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Placas", "Movimientos"])
        filas = list(wb["Placas"].iter_rows(values_only=True))
        self.assertEqual(filas[0][:3], ("Nombre", "Hardware", "Serie"))
        self.assertEqual(sorted(f[0] for f in filas[1:]), ["TPMS-V1-0831", "TPMS-V1-0832"])
        self.assertEqual(filas[1][1], "V1")
        self.assertGreaterEqual(wb["Movimientos"].max_row, 3)

    def test_no_toca_escaneos(self):
        with db.get_db() as c:
            antes = c.execute("SELECT COUNT(*) FROM escaneos").fetchone()[0]
        p = self._manual("601")
        self.c.post("/api/tpms/escanear", json={"codigo": "TPMS_V1_0602"})
        self.c.put(f"/api/tpms/pcb/{p['id']}/firmware", json={"firmware": "fw"})
        self.c.patch(f"/api/tpms/pcb/{p['id']}", json={"nota": "x"})
        self.c.post("/api/tpms/pcb/masa", json={"ids": [p["id"]], "accion": "lote", "valor": "L"})
        self.c.delete(f"/api/tpms/pcb/{p['id']}")
        with db.get_db() as c:
            self.assertEqual(c.execute("SELECT COUNT(*) FROM escaneos").fetchone()[0], antes)

    def test_migracion_desde_tabla_vieja(self):
        from app.services import tpms as svc
        with db.get_db() as c:
            c.execute("DROP TABLE IF EXISTS tpms_pcb")
            c.execute("DROP TABLE IF EXISTS tpms_eventos")
            c.execute("CREATE TABLE tpms_pcb(id INTEGER PRIMARY KEY AUTOINCREMENT, nombre TEXT NOT NULL UNIQUE,"
                      " version TEXT NOT NULL, serie TEXT NOT NULL, codigo_qr TEXT NOT NULL DEFAULT '',"
                      " lote TEXT NOT NULL DEFAULT '', operador TEXT,"
                      " recibida_en TEXT NOT NULL DEFAULT (datetime('now','localtime')))")
            c.execute("INSERT INTO tpms_pcb(nombre, version, serie) VALUES ('TPMS-V1-0701', '1', '0701')")
        svc._asegurar()
        svc._asegurar()   # idempotente
        fila = self.c.get("/api/tpms/pcb/por-codigo", params={"codigo": "701"}).json()["pcb"]
        self.assertEqual((fila["estado"], fila["nota"], fila["firmware"], fila["programada_en"]), ("RECIBIDA", "", None, None))
        self.assertEqual(fila["actualizado_en"], fila["recibida_en"])
        r = self.c.put(f"/api/tpms/pcb/{fila['id']}/firmware", json={"firmware": "fw-m"}).json()["pcb"]
        self.assertEqual(r["estado"], "PROGRAMADA")


if __name__ == "__main__":
    unittest.main()
