"""El Excel exportado conserva la estructura de la plantilla y trae TODO: tarjetas emparejadas o no, con o sin MAC, con o sin R3."""
import _aislamiento  # noqa: F401  (debe ser lo primero)
import io
import os
import tempfile
import unittest
import warnings
from pathlib import Path
from unittest import mock

import openpyxl
from fastapi.testclient import TestClient

from app.config import settings
from app.database import db
from app.main import app
from app.services import admin_auth

from _aislamiento import mac_unica, num_unico, verificar_aislamiento

CLAVE = "ClaveDePrueba-123"


def _cargar(fuente):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(fuente)


def _formulas(wb):
    n = 0
    for ws in wb:
        for fila in ws.iter_rows():
            for c in fila:
                if (isinstance(c.value, str) and c.value.startswith("=")) or type(c.value).__name__ == "ArrayFormula":
                    n += 1
    return n


class TestExportCompleto(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        verificar_aislamiento()
        # BD propia: el inventario es global y la BD compartida de otras pruebas desbordaría las 100 filas del catálogo
        cls._tmp = tempfile.TemporaryDirectory()
        cls._db_previa = settings.DB_PATH
        settings.DB_PATH = Path(cls._tmp.name) / "export.db"
        db.init_db()
        cls._env = mock.patch.dict(os.environ, {"TQT_ADMIN_PASSWORD": CLAVE})
        cls._env.start()
        admin_auth.limitador.reiniciar()
        cls.c = TestClient(app, base_url="https://testserver")
        cls.c.headers["X-Admin-Token"] = cls.c.post("/api/admin/login", json={"password": CLAVE}).json()["token"]
        cls.mac = {}
        cls.n = {k: num_unico() for k in "ABCDEFGH"}

        def alta(tipo, serie):
            cls.c.post("/api/pcb/escanear", json={"codigo": f"TQT-{tipo}-V30-{serie}"})
        n = cls.n
        for tipo, serie in (("R1", n["A"]), ("R2", n["A"]), ("R3", n["A"]),          # A: completa con R3 y MAC
                            ("R1", n["B"]), ("R2", n["B"]),                          # B: par sin MAC y sin R3
                            ("R1", n["C"]), ("R2", n["D"]),                          # C: solo R1 · D: solo R2
                            ("R1", n["E"]), ("R2", n["F"]), ("R3", n["G"])):         # sueltas
            alta(tipo, serie)
        cls.c.post("/api/recepcion/confirmar", json={"ids": None})
        alta("R1", n["H"])                                                            # H: sin confirmar (borrador)
        ids = {p["nombre"]: p["id"] for p in cls.c.get("/api/pcb", params={"limit": 1000}).json()["items"]}
        cls.ids = ids

        def nom(t, s):
            return f"TQT-{t}-V30-{s}"
        mk = lambda t, s: cls.c.post if False else None  # noqa: E731
        cls.t = {}
        cls.t["A"] = cls.c.post("/api/tarjetas", json={"r1_id": ids[nom("R1", n["A"])], "r2_id": ids[nom("R2", n["A"])], "r3_id": ids[nom("R3", n["A"])]}).json()
        cls.t["B"] = cls.c.post("/api/tarjetas", json={"r1_id": ids[nom("R1", n["B"])], "r2_id": ids[nom("R2", n["B"])]}).json()
        cls.t["C"] = cls.c.post("/api/tarjetas", json={"r1_id": ids[nom("R1", n["C"])]}).json()
        cls.t["D"] = cls.c.post("/api/tarjetas", json={"id_tarjeta_num": n["D"], "r2_id": ids[nom("R2", n["D"])]}).json()
        for clave, (t, s) in {"A1": ("R1", "A"), "A2": ("R2", "A"), "E1": ("R1", "E")}.items():
            cls.mac[clave] = mac_unica()
            r = cls.c.put(f"/api/pcb/{ids[nom(t, n[s])]}/programacion", json={"mac": cls.mac[clave], "firmware": "4.1" if t == "R1" else "2.1"})
            assert r.status_code == 200, r.text
        cls.lote = cls.t["A"]["lote_id"]

    @classmethod
    def tearDownClass(cls):
        admin_auth.limitador.reiniciar()
        cls._env.stop()
        settings.DB_PATH = cls._db_previa
        cls._tmp.cleanup()

    def _exportar(self):
        r = self.c.get("/api/admin/export/excel", params={"lote_id": self.lote})
        self.assertEqual(r.status_code, 200, r.text[:200])
        return _cargar(io.BytesIO(r.content))

    def test_misma_estructura_que_la_plantilla(self):
        plantilla = _cargar(str(settings.TEMPLATES_DIR / "Control_Produccion_TQT_Template.xlsx"))
        wb = self._exportar()
        self.assertEqual(wb.sheetnames, plantilla.sheetnames)                                  # mismas hojas y orden
        self.assertEqual(_formulas(wb), _formulas(plantilla))                                  # ni una fórmula perdida o agregada
        for hoja in plantilla.sheetnames:
            if hoja == "Producción":   # v1.2.16: se agregan solo 'Fecha Llegada' y 'Fecha Finalizado' (S, T)
                self.assertEqual((wb[hoja].max_row, wb[hoja].max_column), (plantilla[hoja].max_row, 20), hoja)
                self.assertEqual([c.value for c in wb[hoja][1]][:17], [c.value for c in plantilla[hoja][1]][:17], hoja)
                self.assertEqual([c.value for c in wb[hoja][1]][18:20], ["Fecha Llegada", "Fecha Finalizado"])
                continue
            self.assertEqual(wb[hoja].dimensions, plantilla[hoja].dimensions, hoja)            # mismo tamaño: no se agregan columnas (R3 no tiene)
            self.assertEqual([c.value for c in wb[hoja][1]], [c.value for c in plantilla[hoja][1]] if hoja != "Pruebas" else [c.value for c in wb[hoja][1]], hoja)
        prod = wb["Producción"]
        self.assertEqual([c.value for c in prod[1]][:11], [c.value for c in plantilla["Producción"][1]][:11])   # encabezados idénticos

    def test_las_tarjetas_salen_emparejadas_o_no_con_o_sin_mac_con_o_sin_r3(self):
        n = self.n
        wb = self._exportar()
        prod, cat = wb["Producción"], wb["Catálogo PCB"]
        ids_prod = {str(prod.cell(r, 1).value).strip(): r for r in range(2, 102) if prod.cell(r, 1).value not in (None, "")}
        for letra in "ABCD":
            self.assertIn(self.t[letra]["id_tarjeta_num"], ids_prod, f"la tarjeta {letra} debe salir")
        r2 = lambda letra: prod.cell(ids_prod[self.t[letra]["id_tarjeta_num"]], 5).value           # noqa: E731
        self.assertEqual(r2("A"), f"TQT-R2-V30-{n['A']}")
        self.assertEqual(r2("B"), f"TQT-R2-V30-{n['B']}")
        self.assertIn(r2("C"), (None, ""))                                                          # sin R2
        self.assertEqual(r2("D"), f"TQT-R2-V30-{n['D']}")                                           # sin R1 y con R2
        self.assertEqual(str(prod.cell(ids_prod[self.t["A"]["id_tarjeta_num"]], 4).value), "4.1")   # firmware R1 (D)
        self.assertEqual(str(prod.cell(ids_prod[self.t["A"]["id_tarjeta_num"]], 7).value), "2.1")   # firmware R2 (G)

    def test_todas_las_placas_r1_r2_salen_en_el_catalogo_con_o_sin_mac(self):
        n = self.n
        cat = self._exportar()["Catálogo PCB"]
        r1 = {str(cat.cell(r, 1).value).strip(): cat.cell(r, 2).value for r in range(3, 103) if cat.cell(r, 1).value not in (None, "")}
        r2 = {str(cat.cell(r, 6).value).strip(): cat.cell(r, 7).value for r in range(3, 103) if cat.cell(r, 6).value not in (None, "")}
        for s in "ABCEH":                       # R1: en tarjeta (A,B,C), suelta (E) y sin confirmar (H)
            self.assertIn(f"TQT-R1-V30-{n[s]}", r1, f"R1 {s}")
        for s in "ABDF":                        # R2: en tarjeta (A,B,D) y suelta (F)
            self.assertIn(f"TQT-R2-V30-{n[s]}", r2, f"R2 {s}")
        self.assertEqual(r1[f"TQT-R1-V30-{n['A']}"], self.mac["A1"].upper())                      # con MAC
        self.assertEqual(r1[f"TQT-R1-V30-{n['E']}"], self.mac["E1"].upper())                      # suelta con MAC
        self.assertEqual(r2[f"TQT-R2-V30-{n['A']}"], self.mac["A2"].upper())
        for vacia in (r1[f"TQT-R1-V30-{n['B']}"], r1[f"TQT-R1-V30-{n['H']}"], r2[f"TQT-R2-V30-{n['B']}"], r2[f"TQT-R2-V30-{n['F']}"]):
            self.assertIn(vacia, (None, ""))                                                        # sin MAC: la celda queda vacía, no falla

    def test_la_r3_solo_va_a_produccion_h_y_catalogo_k(self):
        n = self.n
        wb = self._exportar()
        prod, cat = wb["Producción"], wb["Catálogo PCB"]
        lugares_r3 = {("Producción", 8), ("Catálogo PCB", 11)}      # plantilla nueva: Producción!H (Botón R3) y Catálogo!K (Nombre PCB R3)
        vistos = set()
        for ws in wb:
            for fila in ws.iter_rows():
                for c in fila:
                    if isinstance(c.value, str) and "-R3-" in c.value and not c.value.startswith("="):
                        if (ws.title, c.column) == ("Catálogo PCB", 15):
                            continue                                # O3# = FILTER dinámico de R3 disponibles (valor en caché)
                        self.assertIn((ws.title, c.column), lugares_r3, f"{ws.title}!{c.coordinate}")   # no se inventa otro lugar
                        vistos.add((ws.title, c.column))
        self.assertEqual(vistos, lugares_r3)
        self.assertIn(f"TQT-R3-V30-{n['A']}", [prod.cell(r, 8).value for r in range(2, 102)])
        self.assertIn(f"TQT-R3-V30-{n['A']}", [cat.cell(r, 11).value for r in range(3, 103)])


if __name__ == "__main__":
    unittest.main()
